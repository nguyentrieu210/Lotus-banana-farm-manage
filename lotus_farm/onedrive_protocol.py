"""OneDrive production rules, translated from inspected Excel formulas.

Selection plans use EDATE(original_anchor, 6*n), not 180 days or a chain of
previous rounded dates. A Chinese first cut is an equivalent selection DATE,
not proof of planting, selected sucker count, or a Lotus harvest receipt.
"""
import calendar
import hashlib
import re
import math
from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from zipfile import ZipFile


SELECTION_INTERVAL_MONTHS = 6
SELECTION_SOURCE = {
    "file": "Báo cáo Chọn chồi.xlsx",
    "sheet": "Tổng hợp",
    "range": "K6:U6",
    "clear_suckers_formula": "EDATE(K6,4)",
    "selection_formulas": [f"EDATE(K6,{n})" for n in range(6, 43, 6)],
    "actual_sheet": "NL_Chọn chồi",
    "actual_date_column": "C",
    "actual_cycle_column": "B",
    "actual_quantity_column": "Y",
    "url": "https://onedrive.live.com/personal/af58eed931a9a3a8/_layouts/15/doc.aspx?resid=f9cdde84-d94d-4d74-a04d-dbd1ec9bf15d&cid=af58eed931a9a3a8",
}

PROCESS_STEPS = (
    ("handover", "Bàn giao theo khu / đợt", "Lotus Handover Cohort", "process", "Nhật ký chặt cây cũ; giữ ngày và số cây nguồn"),
    ("land", "Vườn, lô, thửa, vụ cây", "Crop Cycle", "farms", "Datalo, TrongMoi, ttKetVu"),
    ("selection", "Chọn chồi vụ tiếp theo", "Sucker Selection Batch", "farms", "NL_Chọn chồi; ngày làm và số chồi thật"),
    ("planning", "Kế hoạch công việc, nhân công, vật tư", "Daily Work Plan", "plans", "KH Chọn chồi, KHSL, định mức theo công đoạn"),
    ("care", "Chăm sóc vườn, BVTV, bệnh", "Daily Work Result", "work", "NL_Chăm sóc vườn, BVTV, kiểm kê Fusarium"),
    ("bunch", "Trổ, chích bắp, cắt bắp, chăm sóc buồng", "Bunch Operation", "work", "ChichBap, CatBap, chăm sóc buồng"),
    ("measurement", "Đo trái và loại buồng", "Fruit Measurement", "measurements", "FM đo trái; nguyên nhân hủy buồng"),
    ("forecast", "Dự thu theo nguồn và kỳ", "Harvest Forecast", "forecast", "KHSL tuần/tháng, tsKHTH, tsPC"),
    ("harvest", "Thu hoạch và nhận buồng", "Harvest Actual", "harvest", "ThuHoach; số buồng và kg cân thực tế"),
    ("quality", "QC đầu vào", "Incoming Quality Inspection", "incoming-qc", "Sơ chế, phân loại, hủy tại xưởng"),
    ("packing", "Đóng gói, phẩm cấp và phụ phẩm", "Packout Run", "packhouse", "FM xưởng đóng gói; A456/A789/B456/B789/CL/CP/NB"),
    ("stock", "Thành phẩm, kho và xuất bán", "Stock Entry", "kho", "Thùng, kg, container và sản lượng xuất bán"),
    ("day_close", "Đối soát và chốt ngày", "Daily Production Close", "daily-close", "Kế hoạch / kết quả / công / vật tư / sản lượng"),
    ("cycle_close", "Kết vụ và hiệu chuẩn vụ sau", "Cycle Closing Report", "cycle-close", "ttKetVu; số cây, buồng, kg, diện tích và năng suất"),
)


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except (ValueError, TypeError):
        raise ValueError("Ngày phải có định dạng YYYY-MM-DD.") from None


def edate(origin, months):
    origin = as_date(origin)
    if isinstance(months, bool) or int(months) != months:
        raise ValueError("Số tháng phải là số nguyên.")
    year, month = divmod(origin.year * 12 + origin.month - 1 + int(months), 12)
    return date(year, month + 1, min(origin.day, calendar.monthrange(year, month + 1)[1]))


def selection_schedule(origin, *, basis="First Cut Equivalent", count=7, first_cycle=2):
    """Dates only. No inferred date becomes an observed operation."""
    origin = as_date(origin)
    if basis not in ("Planting", "Selection", "First Cut Equivalent"):
        raise ValueError("Loại ngày mốc không được hỗ trợ.")
    if not 1 <= int(count) <= 24 or int(first_cycle) < 2:
        raise ValueError("Số kỳ hoặc số vụ không hợp lệ.")
    rows = []
    for index in range(int(count)):
        offset = SELECTION_INTERVAL_MONTHS * (index + (1 if basis == "Planting" else 0))
        rows.append({
            "selection_date": str(edate(origin, offset)),
            "cycle_no": int(first_cycle) + index,
            "offset_months": offset,
            "date_kind": "Equivalent" if index == 0 and basis == "First Cut Equivalent" else "Planned",
            "source_date": str(origin),
            "source_basis": basis,
            "observed": False,
        })
    return rows


def harvest_curve(origin, *, basis="Selection", cycles=1):
    """tsRaHoa calendar ages; first selected sucker belongs to ratoon crop 2.

    Month buckets follow PowerQuery Date.AddMonths(anchor, age-(day<=2)).
    A selection at planting+6 therefore harvests at selection+9..12, not +2..5.
    The date is a planning bucket, never a promised physical harvest day.
    """
    origin = as_date(origin)
    if basis not in ("Planting", "Selection", "First Cut Equivalent"):
        raise ValueError("Loại ngày mốc không được hỗ trợ.")
    if not 1 <= int(cycles) <= 12:
        raise ValueError("Chỉ dự phóng từ 1 đến 12 vụ.")
    rows = []
    for cycle in range(int(cycles)):
        mother = basis == "Planting" and cycle == 0
        start = (8 if mother else 15 + (cycle - 1) * 6) if basis == "Planting" else 9 + cycle * 6
        shares = (0.2, 0.4, 0.3, 0.1) if mother else (0.17, 0.25, 0.4, 0.18)
        for index, share in enumerate(shares):
            offset = start + index
            bucket = edate(origin, offset - (1 if origin.day <= 2 else 0))
            rows.append({"month": bucket.strftime("%Y-%m"), "harvest_age_months": offset,
                         "cut_age_months": offset - 2, "share": share,
                         "cycle_offset": cycle, "date_kind": "Planned"})
    return rows


def forecast_quantities(quantity, params, share=1):
    """Keep harvested gross mass and processed mass on distinct bases."""
    names = ("cut_rate", "harvest_rate", "processed_rate", "gross_bunch_kg", "processed_bunch_kg")
    values = [float(params.get(key) or 0) for key in names]
    quantity, share = float(quantity), float(share)
    if any(not math.isfinite(v) for v in [quantity, share, *values]):
        raise ValueError("Hệ số và số lượng phải hữu hạn.")
    cut_rate, harvest_rate, processed_rate, gross_kg, processed_kg = values
    if quantity < 0 or not 0 <= share <= 1 or any(not 0 <= v <= 1 for v in values[:3]):
        raise ValueError("Số lượng hoặc tỷ lệ không hợp lệ.")
    if gross_kg <= 0 or processed_kg <= 0 or processed_kg > gross_kg:
        raise ValueError("Khối lượng buồng phải dương; kg sơ chế không vượt kg nguyên buồng.")
    cut = quantity * share * cut_rate
    harvest = cut * harvest_rate
    processed = harvest * processed_rate
    return {"cut_bunches": round(cut, 3), "harvest_bunches": round(harvest, 3),
            "processed_bunches": round(processed, 3), "gross_kg": round(harvest * gross_kg, 3),
            "processed_kg": round(processed * processed_kg, 3)}


def color_periods(anchor, anchor_code, colors, start, end, period_days=5):
    """Extend a versioned program; generated periods remain plans, not tags used."""
    anchor, start, end = map(as_date, (anchor, start, end))
    if not colors or anchor_code not in [c["code"] for c in colors] or not 1 <= int(period_days) <= 31:
        raise ValueError("Chương trình màu không hợp lệ.")
    if end < start or (end - start).days > 3653:
        raise ValueError("Khoảng chương trình màu tối đa 10 năm.")
    step = int(period_days)
    index = (start - anchor).days // step
    origin = next(n for n, row in enumerate(colors) if row["code"] == anchor_code)
    rows = []
    while (when := anchor + timedelta(days=index * step)) <= end:
        color = colors[(origin + index) % len(colors)]
        rows.append({"valid_from": str(when), "valid_to": str(when + timedelta(days=step - 1)),
                     "color_code": color["code"], "color_label": color["label"],
                     "tag_code": f"{when:%d.%m.%y}_{color['code']}", "date_kind": "Planned"})
        index += 1
    return rows


def scheduled_work_due(anchor, on_date, *, offset_days=0, offset_months=0, repeat_days=0, repeat_months=0):
    """Month recurrences always refer to the original date, including month ends."""
    anchor, on_date = as_date(anchor), as_date(on_date)
    if repeat_days and repeat_months:
        raise ValueError("Chỉ chọn lặp theo ngày hoặc theo tháng.")
    first = edate(anchor, int(offset_months)) + timedelta(days=int(offset_days))
    if on_date < first:
        return False
    if repeat_months:
        span = (on_date.year-anchor.year)*12 + on_date.month-anchor.month-int(offset_months)
        # Offsets can carry dates across a month boundary; inspect adjacent cycles.
        approx = max(0,span//int(repeat_months))
        return any(edate(anchor,int(offset_months)+n*int(repeat_months))+timedelta(days=int(offset_days)) == on_date
                   for n in range(max(0,approx-2),approx+3))
    return ((on_date-first).days % int(repeat_days)==0) if repeat_days else on_date==first


def parse_handover_workbook(content):
    """Read daily Chinese handover history; ignore monthly subtotal rows.

    The six-column zone sheets are authoritative for zone identity. Their
    daily totals must reconcile to the separate summary before any import.
    """
    from openpyxl import load_workbook
    if not isinstance(content, bytes) or len(content) > 40 * 1024 * 1024:
        raise ValueError("Nhật ký phải là file Excel không quá 40 MiB.")
    with ZipFile(BytesIO(content)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 200 * 1024 * 1024:
            raise ValueError("Nhật ký có dữ liệu giải nén quá lớn.")
    book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    digest = hashlib.sha256(content).hexdigest()
    groups = defaultdict(lambda: {"quantity": 0, "cells": []})
    totals = defaultdict(int)
    warnings = []
    try:
        for sheet in book:
            match = re.match(r"^Khu\s+([A-E])\s*-\s*Chặt chuối TH1$", sheet.title)
            if not match:
                continue
            zone = match[1]
            if sheet.max_row > 10000:
                raise ValueError("Mỗi khu được nhập tối đa 10.000 dòng.")
            heading = str(sheet.cell(4, 1).value or "")
            if heading and f"Khu {zone}" not in heading:
                warnings.append(f"{sheet.title}!A4 ghi '{heading}'; lấy khu từ tên sheet.")
            if [sheet.cell(6, n).value for n in (1, 2)] != ["Ngày", "Số lượng"]:
                raise ValueError(f"{sheet.title}: thiếu hai cột Ngày / Số lượng tại dòng 6.")
            for number, row in enumerate(sheet.iter_rows(min_row=7, max_col=2, values_only=True), 7):
                when, quantity = row
                if not isinstance(when, (date, datetime)):
                    continue
                when = as_date(when)
                if isinstance(quantity, bool) or not isinstance(quantity, (int, float)) or quantity <= 0 or not float(quantity).is_integer():
                    raise ValueError(f"{sheet.title}!B{number}: số cây đã chặt phải là số nguyên dương.")
                key = (zone, str(when))
                groups[key]["quantity"] += int(quantity)
                groups[key]["cells"].append({"sheet": sheet.title, "row": number, "date_cell": f"A{number}", "quantity_cell": f"B{number}"})
                totals[zone] += int(quantity)
        if set(totals) != set("ABCDE"):
            raise ValueError("Nhật ký phải có đủ các sheet khu A–E với dữ liệu ngày.")
        if "Tổng hợp" not in book.sheetnames:
            raise ValueError("Thiếu sheet Tổng hợp để đối chiếu số cây.")
        control = {}
        for row in book["Tổng hợp"].iter_rows(min_row=4, max_row=8, min_col=2, max_col=3, values_only=True):
            zone, quantity = row
            match = re.fullmatch(r"Khu ([A-E])", str(zone or ""))
            if match:
                control[match[1]] = quantity
        if any(control.get(zone) != totals[zone] for zone in "ABCDE"):
            raise ValueError(f"Số ngày không khớp Tổng hợp: chi tiết {dict(totals)}, tổng hợp {control}.")
        rows = [{"farm_zone": zone, "source_event_date": when,
                 "equivalent_selection_date": when, "source_cut_qty": info["quantity"],
                 "source_cells": info["cells"], "source_sha256": digest,
                 "source_event_type": "First Cut", "source_basis": "First Cut Equivalent"}
                for (zone, when), info in sorted(groups.items())]
        return {"rows": rows, "zone_totals": dict(totals), "total_cut_qty": sum(totals.values()),
                "source_sha256": digest, "warnings": warnings}
    finally:
        book.close()
