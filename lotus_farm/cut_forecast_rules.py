"""Dự thu theo nhật ký chặt mẹ lần 2 — pure rules (no Frappe), mirroring the Lotus forecast workbook
("LOTUS FARM 222 DỰ THU 2026.2027.xlsx": sheets NHẬT KÝ GỐC, THÔNG SỐ, DỰ BÁO).

Each journal row (khu, ngày chặt mẹ lần 2, số cây) splits into 4 shooting waves:
  chồi      = round(cây × chồi/cây)
  trổ đợt k = round(chồi × tỷ lệ k); đợt cuối = phần còn lại (so the waves always add up to the suckers)
  trổ từ    = ngày chặt + days_to_first_shoot + (k-1) × wave_step_days; trổ đến = trổ từ + window_days - 1
  đại diện  = trổ từ + representative_day
  chích bắp = đại diện + shoot_to_inject_days        (chích bắp before the flower opens)
  cắt bắp   = chích + inject_to_budcut_days
  thu       = cắt bắp + budcut_to_harvest_days
  buồng cắt bắp = round(trổ × cut_rate); thu = round(× harvest_rate); đóng gói = round(× pack_rate); kg = đóng gói × kg_per_bunch
Rounding is Excel's ROUND (half away from zero), so totals match the workbook exactly.
"""
import datetime as dt
import hashlib
import re
from collections import OrderedDict
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

# (field, label, kind, default) in the order of sheet THÔNG SỐ; kind 'days' = whole days, 'rate' = 0..1.
PARAMS = (
    ('days_to_first_shoot', 'Chặt mẹ lần 2 → bắt đầu đợt trổ 1 (ngày)', 'days', 120),
    ('wave_step_days', 'Bước giữa 2 đợt trổ (ngày)', 'days', 30),
    ('window_days', 'Số ngày trong mỗi cửa sổ trổ', 'days', 30),
    ('representative_day', 'Ngày đại diện trong cửa sổ', 'days', 14),
    ('shoot_to_inject_days', 'Trổ → chích bắp (ngày)', 'days', 5),
    ('inject_to_budcut_days', 'Chích bắp → cắt bắp (ngày)', 'days', 15),
    ('budcut_to_harvest_days', 'Cắt bắp → thu hoạch (ngày)', 'days', 65),
    ('cut_rate', 'Tỷ lệ cắt bắp / số cây trổ', 'rate', 0.9),
    ('harvest_rate', 'Tỷ lệ thu hoạch / cắt bắp', 'rate', 0.97),
    ('pack_rate', 'Tỷ lệ đóng gói / thu hoạch', 'rate', 0.98),
    ('kg_per_bunch', 'Kg sơ chế / buồng', 'kg', 20),
    ('sucker_per_tree', 'Chồi hữu hiệu / cây mẹ chặt', 'ratio', 1),
    ('wave_1_share', 'Tỷ lệ trổ đợt 1', 'rate', 0.17),
    ('wave_2_share', 'Tỷ lệ trổ đợt 2', 'rate', 0.40),
    ('wave_3_share', 'Tỷ lệ trổ đợt 3', 'rate', 0.25),
    ('wave_4_share', 'Tỷ lệ trổ đợt 4', 'rate', 0.18),
)
LABELS = {f: label for f, label, _, _ in PARAMS}
DEFAULTS = {f: d for f, _, _, d in PARAMS}
WAVES = ('wave_1_share', 'wave_2_share', 'wave_3_share', 'wave_4_share')
STAGES = (('shoot', 'Trổ buồng'), ('inject', 'Chích bắp'), ('budcut', 'Cắt bắp'), ('harvest', 'Thu hoạch'), ('pack', 'Đóng gói'))
# Column A labels of sheet THÔNG SỐ (accents and spacing are normalised before matching).
SHEET_LABELS = {
    'chat me lan 2 bat dau dot tro 1': 'days_to_first_shoot',
    'buoc giua 2 dot tro': 'wave_step_days',
    'so ngay trong moi cua so tro': 'window_days',
    'ngay dai dien trong cua so': 'representative_day',
    'tro chich bap': 'shoot_to_inject_days',
    'chich cat bap': 'inject_to_budcut_days',
    'cat bap thu hoach': 'budcut_to_harvest_days',
    'ty le cat bap so cay': 'cut_rate',
    'ty le thu hoach cat bap': 'harvest_rate',
    'ty le dong goi thu hoach': 'pack_rate',
    'kg so che buong': 'kg_per_bunch',
    'choi huu hieu cay me chat': 'sucker_per_tree',
}


def excel_round(value, digits=0):
    """Excel ROUND: half away from zero (Python's round() is half-to-even)."""
    q = Decimal(1).scaleb(-digits)
    return float(Decimal(str(value)).quantize(q, rounding=ROUND_HALF_UP)) if digits else int(Decimal(str(value)).quantize(q, rounding=ROUND_HALF_UP))


def _plain(text):
    import unicodedata
    text = unicodedata.normalize('NFD', str(text or '').lower().replace('đ', 'd'))
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', re.sub(r'\(.*?\)', ' ', text)).split())


def as_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value or '').strip()
    for pattern in ('%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d', '%d-%m-%Y'):
        try:
            return dt.datetime.strptime(text[:10], pattern).date()
        except ValueError:
            continue
    raise ValueError('ngày không đúng định dạng')


def validate_profile(values):
    """Clean numbers for every parameter -> (profile, errors). Wave shares must add up to 100%."""
    profile, errors = {}, []
    for field, label, kind, default in PARAMS:
        raw = values.get(field, default)
        try:
            number = float(raw)
        except (TypeError, ValueError):
            errors.append('%s: cần một số' % label)
            continue
        if kind == 'days':
            if number < 0 or number != int(number):
                errors.append('%s: cần số ngày nguyên không âm' % label)
                continue
            number = int(number)
        elif kind == 'rate' and not 0 <= number <= 1:
            errors.append('%s: cần từ 0 đến 1 (0–100%%)' % label)
            continue
        elif kind in ('kg', 'ratio') and number <= 0:
            errors.append('%s: cần lớn hơn 0' % label)
            continue
        profile[field] = number
    if not errors:
        if profile['window_days'] < 1:
            errors.append('%s: ít nhất 1 ngày' % LABELS['window_days'])
        elif profile['representative_day'] >= profile['window_days']:
            errors.append('Ngày đại diện phải nằm trong cửa sổ trổ (nhỏ hơn %d)' % profile['window_days'])
        total = sum(profile[w] for w in WAVES)
        if abs(total - 1) > 1e-6:
            errors.append('Tổng tỷ lệ 4 đợt trổ phải bằng 100%% (đang là %s%%)' % ('%.2f' % (total * 100)).rstrip('0').rstrip('.'))
    return profile, errors


def entry_key(farm, zone, cut_date, source_ref):
    return hashlib.sha1('|'.join([str(farm), str(zone), str(cut_date), str(source_ref)]).encode()).hexdigest()


def expand(entries, profile):
    """entries: dicts with zone, cut_date, tree_qty (and optionally name, source_ref) -> one row per wave."""
    add = dt.timedelta
    rows = []
    for e in entries:
        cut = as_date(e['cut_date'])
        suckers = excel_round(float(e['tree_qty']) * profile['sucker_per_tree'])
        given = 0
        for k, wave in enumerate(WAVES, start=1):
            shoot = max(0, suckers - given) if k == len(WAVES) else excel_round(suckers * profile[wave])
            given += shoot
            start = cut + add(days=profile['days_to_first_shoot'] + (k - 1) * profile['wave_step_days'])
            shoot_date = start + add(days=profile['representative_day'])
            inject = shoot_date + add(days=profile['shoot_to_inject_days'])
            budcut = inject + add(days=profile['inject_to_budcut_days'])
            harvest = budcut + add(days=profile['budcut_to_harvest_days'])
            budcut_qty = excel_round(shoot * profile['cut_rate'])
            harvest_qty = excel_round(budcut_qty * profile['harvest_rate'])
            pack_qty = excel_round(harvest_qty * profile['pack_rate'])
            rows.append(dict(entry=e.get('name') or e.get('entry_key'), zone=e['zone'], cut_date=cut, tree_qty=int(e['tree_qty']), source_ref=e.get('source_ref') or '',
                             farm_block=e.get('farm_block'), crop_cycle=e.get('crop_cycle'), plan_cohort=e.get('plan_cohort'),
                             entry_modified=e.get('modified'),
                             wave=k, share=profile[wave], suckers=suckers, shoot_qty=shoot, shoot_from=start, shoot_to=start + add(days=profile['window_days'] - 1),
                             shoot_date=shoot_date, inject_date=inject, inject_qty=shoot, budcut_date=budcut, budcut_qty=budcut_qty,
                             harvest_date=harvest, harvest_qty=harvest_qty, pack_qty=pack_qty, kg=pack_qty * profile['kg_per_bunch']))
    return rows


def _month(day):
    return '%04d-%02d' % (day.year, day.month)


def _months(first, last):
    y, m = int(first[:4]), int(first[5:])
    out = []
    while '%04d-%02d' % (y, m) <= last:
        out.append('%04d-%02d' % (y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


# (stage, quantity field, date field) — đóng gói and kg are counted in the harvest month, as in the workbook.
STAGE_FIELDS = (('shoot', 'shoot_qty', 'shoot_date'), ('inject', 'inject_qty', 'inject_date'), ('budcut', 'budcut_qty', 'budcut_date'),
                ('harvest', 'harvest_qty', 'harvest_date'), ('pack', 'pack_qty', 'harvest_date'), ('kg', 'kg', 'harvest_date'))


def monthly(rows):
    """Every month from the first stage date to the last (none dropped) with each stage's quantity."""
    if not rows:
        return []
    found = [_month(r[d]) for r in rows for _, _, d in STAGE_FIELDS]
    table = OrderedDict((m, dict(month=m, **{s: 0 for s, _, _ in STAGE_FIELDS})) for m in _months(min(found), max(found)))
    for r in rows:
        for stage, qty, day in STAGE_FIELDS:
            table[_month(r[day])][stage] += r[qty]
    for row in table.values():
        row['tonnes'] = round(row['kg'] / 1000, 2)
    return list(table.values())


def by_zone(rows, entries):
    trees = OrderedDict()
    for e in sorted(entries, key=lambda e: str(e['zone'])):
        trees[e['zone']] = trees.get(e['zone'], 0) + int(e['tree_qty'])
    out = OrderedDict((z, dict(zone=z, trees=t, entries=0, first_cut=None, last_cut=None, suckers=0, **{s: 0 for s, _, _ in STAGE_FIELDS})) for z, t in trees.items())
    for e in entries:
        z, day = out[e['zone']], as_date(e['cut_date'])
        z['entries'] += 1
        z['first_cut'] = min(z['first_cut'] or day, day)
        z['last_cut'] = max(z['last_cut'] or day, day)
    for r in rows:
        z = out[r['zone']]
        if r['wave'] == 1:
            z['suckers'] += r['suckers']
        for stage, qty, _ in STAGE_FIELDS:
            z[stage] += r[qty]
    for z in out.values():
        z['tonnes'] = round(z['kg'] / 1000, 2)
    return list(out.values())


def zone_months(rows):
    """Tonnes by zone and harvest month: {zone: {month: tonnes}}."""
    out = {}
    for r in rows:
        cell = out.setdefault(r['zone'], {})
        month = _month(r['harvest_date'])
        cell[month] = cell.get(month, 0) + r['kg']
    return {z: {m: round(kg / 1000, 2) for m, kg in months.items()} for z, months in out.items()}


def totals(rows, entries, months):
    t = dict(trees=sum(int(e['tree_qty']) for e in entries), entries=len(entries), suckers=sum(r['suckers'] for r in rows if r['wave'] == 1),
             **{s: sum(r[q] for r in rows) for s, q, _ in STAGE_FIELDS})
    t['tonnes'] = round(t['kg'] / 1000, 2)
    t['first_cut'] = min((as_date(e['cut_date']) for e in entries), default=None)
    t['last_cut'] = max((as_date(e['cut_date']) for e in entries), default=None)
    t['first_harvest'] = min((r['harvest_date'] for r in rows), default=None)
    t['last_harvest'] = max((r['harvest_date'] for r in rows), default=None)
    peak = max(months, key=lambda m: m['kg'], default=None)
    t['peak_month'], t['peak_tonnes'] = (peak['month'], peak['tonnes']) if peak and peak['kg'] else (None, 0)
    t['zones'] = sorted({e['zone'] for e in entries})
    return t


def compute(entries, profile):
    rows = expand(entries, profile)
    months = monthly(rows)
    return dict(rows=rows, months=months, zones=by_zone(rows, entries), zone_months=zone_months(rows), totals=totals(rows, entries, months))


# Workbook ------------------------------------------------------------------------------------------------

def _sheet(book, *names):
    wanted = {_plain(n) for n in names}
    for sheet in book.worksheets:
        if _plain(sheet.title) in wanted:
            return sheet
    return None


def parse_workbook(content):
    """The forecast workbook -> dict(entries, entry_errors, params, param_errors, found_params).

    NHẬT KÝ GỐC: header row 'Khu | Ngày chặt mẹ lần 2 | Số cây | Dòng trên file nguồn', data below it.
    THÔNG SỐ: parameters by their column-A label, wave shares from the 'Đợt trổ' table (1..4 in column A).
    """
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(content), data_only=True)
    journal = _sheet(book, 'NHẬT KÝ GỐC', 'NHAT KY GOC')
    if journal is None:
        raise ValueError('Không thấy sheet "NHẬT KÝ GỐC" trong file.')
    entries, entry_errors, header = [], [], None
    for row in journal.iter_rows():
        cells = [c.value for c in row[:4]] + [None] * (4 - len(row[:4]))
        line = row[0].row
        if header is None:
            if _plain(cells[0]) == 'khu' and 'ngay' in _plain(cells[1]):
                header = line
            continue
        if all(v in (None, '') for v in cells[:3]):
            continue
        zone = str(cells[0] or '').strip().upper()
        problems = []
        if not re.fullmatch(r'[A-Z][A-Z0-9]{0,3}', zone):
            problems.append('khu "%s" không hợp lệ' % (cells[0] or ''))
        try:
            cut = as_date(cells[1])
        except ValueError:
            cut, problems = None, problems + ['ngày chặt "%s" không đúng định dạng' % (cells[1] or '')]
        try:
            qty = float(cells[2])
            if qty <= 0 or qty != int(qty):
                raise ValueError
            qty = int(qty)
        except (TypeError, ValueError):
            qty, problems = None, problems + ['số cây "%s" phải là số nguyên dương' % (cells[2] if cells[2] is not None else '')]
        source = str(cells[3] or '').strip() or 'NHẬT KÝ GỐC | dòng %d' % line
        if problems:
            entry_errors.append(dict(line=line, errors=problems))
            continue
        entries.append(dict(line=line, zone=zone, cut_date=cut, tree_qty=qty, source_ref=source))
    if header is None:
        raise ValueError('Sheet "NHẬT KÝ GỐC" cần dòng tiêu đề bắt đầu bằng "Khu | Ngày chặt mẹ lần 2 | Số cây".')

    params, found, in_waves = {}, [], False
    settings = _sheet(book, 'THÔNG SỐ', 'THONG SO')
    if settings is not None:
        for row in settings.iter_rows(max_col=2):
            label, value = row[0].value, row[1].value if len(row) > 1 else None
            key = _plain(label)
            if key.startswith('dot tro'):
                in_waves = True
                continue
            if in_waves and str(label).strip() in ('1', '2', '3', '4'):
                field = WAVES[int(str(label).strip()) - 1]
            else:
                field = next((f for text, f in SHEET_LABELS.items() if key == text), None)
            if field and value not in (None, ''):
                params[field] = value
                found.append(field)
    profile, param_errors = validate_profile({**DEFAULTS, **params}) if params else ({}, [])
    return dict(entries=entries, entry_errors=entry_errors, params=profile, param_errors=param_errors, found_params=found)
