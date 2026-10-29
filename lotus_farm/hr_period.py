"""Timekeeping calendar, month close and the monthly payroll sheet handed to accounting.

Closing a month freezes Attendance, crew sheets and labour allocation for its dates and stores the
payroll sheet as a snapshot. Amounts are gross before social insurance and personal income tax;
contractor crew piecework is listed as a payable to the crew, never as an employee wage.
"""
import hashlib
import json
from datetime import date

import frappe
from frappe.utils import cint, flt, getdate, now_datetime, nowdate

from lotus_farm.payroll_rules import HOURS_BASE, STANDARD_DAYS, empty_summary, hour_summary, hourly_pay, params_on, tax_method, pay_base, salary_segments, segmented_pay
from lotus_farm.units import label as unit_label
from lotus_farm.timekeeping_rules import WEEKDAYS, calendar_days, close_blockers, missing_attendance, month_range

PERIOD = 'Lotus Timekeeping Period'
GUARDED = {'Attendance': 'attendance_date', 'Lotus Crew Attendance': 'posting_date', 'Labor Allocation': 'posting_date'}


def _hr_manager():
    if frappe.session.user == 'Guest':
        frappe.throw('Vui lòng đăng nhập.', frappe.AuthenticationError)
    if not (frappe.session.user == 'Administrator' or {'HR Manager', 'System Manager'} & set(frappe.get_roles())):
        frappe.throw('Cần vai trò HR Manager để xử lý kỳ công và bảng lương.', frappe.PermissionError)


def _company(name):
    from lotus_farm.timekeeping import _company as check
    return check(name)


def _month(month):
    try:
        return month_range(month)
    except ValueError as exc:
        frappe.throw(str(exc))


def _list_name(company, year):
    abbr = frappe.get_cached_value('Company', company, 'abbr') or company
    return 'LOTUS-' + abbr + '-' + str(year)


def holiday_map(company, begin, end):
    """The Holiday List covering [begin, end] for the company, with {date: weekly_off}; None when not set up."""
    for name in (frappe.get_cached_value('Company', company, 'default_holiday_list'), _list_name(company, begin.year)):
        if not name:
            continue
        span = frappe.db.get_value('Holiday List', name, ['from_date', 'to_date'], as_dict=True)
        if span and getdate(span.from_date) <= begin and getdate(span.to_date) >= end:
            rows = frappe.get_all('Holiday', filters={'parent': name, 'parenttype': 'Holiday List',
                                                      'holiday_date': ['between', [begin, end]]},
                                  fields=['holiday_date', 'weekly_off'], limit_page_length=0)
            return frappe._dict(name=name, days={getdate(r.holiday_date): cint(r.weekly_off) for r in rows})
    return None


def _ym(value):
    return getdate(value).strftime('%Y-%m')


def is_closed(company, day):
    return bool(company and day and frappe.db.exists(PERIOD, {'company': company, 'month': _ym(day), 'status': 'Closed'}))


def period_guard(doc, method=None):
    """doc_events hook: no Attendance, crew sheet or labour allocation changes inside a closed month."""
    field = GUARDED.get(doc.doctype)
    if not field or doc.flags.get('lotus_period_reopen'):
        return
    company = doc.get('company') or (frappe.db.get_value('Employee', doc.get('employee'), 'company') if doc.get('employee') else None)
    if company:
        from lotus_farm.controls import lock
        lock('Company',company)
    before = doc.get_doc_before_save() if method == 'validate' else None
    for day in {doc.get(field), before.get(field) if before else None} - {None, ''}:
        if is_closed(company, day):
            frappe.throw('Tháng ' + _ym(day) + ' đã khóa công. HR Manager mở lại kỳ trước khi sửa.')


# Calendar ---------------------------------------------------------------------------------------

@frappe.whitelist()
def calendar_info(company, year=None):
    _company(company)
    year = cint(year) or getdate(nowdate()).year
    name = _list_name(company, year)
    rows = frappe.get_all('Holiday', filters={'parent': name, 'parenttype': 'Holiday List'},
                          fields=['holiday_date', 'description', 'weekly_off'], order_by='holiday_date',
                          limit_page_length=0) if frappe.db.exists('Holiday List', name) else []
    return dict(year=year, name=name if rows else None,
                is_default=frappe.get_cached_value('Company', company, 'default_holiday_list') == name,
                holidays=[dict(date=str(r.holiday_date), description=frappe.utils.strip_html(r.description or ''),
                               weekly_off=cint(r.weekly_off)) for r in rows],
                weekdays=list(WEEKDAYS))


@frappe.whitelist(methods=['POST'])
def setup_calendar(company, year, weekly_off=None, extra=None):
    """Create or rebuild the company Holiday List of one year: weekly off days, fixed holidays, manual extras."""
    _hr_manager()
    _company(company)
    year = cint(year)
    if not 2020 <= year <= 2100:
        frappe.throw('Năm không hợp lệ.')
    weekly = frappe.parse_json(weekly_off) if isinstance(weekly_off, str) else (weekly_off or ['Sunday'])
    extras = frappe.parse_json(extra) if isinstance(extra, str) else (extra or [])
    if not isinstance(weekly, list) or any(day not in WEEKDAYS for day in weekly) or len(weekly) > 2:
        frappe.throw('Chọn tối đa 2 ngày nghỉ hằng tuần.')
    if not isinstance(extras, list) or len(extras) > 40:
        frappe.throw('Danh sách ngày lễ thêm không hợp lệ.')
    try:
        rows = calendar_days(year, weekly, [(item.get('date'), item.get('description')) for item in extras])
    except (ValueError, TypeError, AttributeError) as exc:
        frappe.throw('Ngày lễ không hợp lệ: ' + str(exc))
    closed = frappe.get_all(PERIOD, filters={'company': company, 'status': 'Closed', 'month': ['like', str(year) + '-%']}, pluck='month')
    if closed:
        frappe.throw('Mở lại các kỳ đã khóa trước khi sửa lịch năm: ' + ', '.join(sorted(closed)))
    from lotus_farm.controls import lock
    lock('Company', company)
    name = _list_name(company, year)
    doc = frappe.get_doc('Holiday List', name) if frappe.db.exists('Holiday List', name) else frappe.new_doc('Holiday List')
    doc.holiday_list_name = name
    doc.from_date, doc.to_date = date(year, 1, 1), date(year, 12, 31)
    doc.set('holidays', [dict(holiday_date=day, description=label, weekly_off=flag) for day, label, flag in rows])
    doc.save()
    # HRMS 16 resolves holidays through Holiday List Assignment (salary slips, leave), not the Company field.
    if frappe.db.exists('DocType', 'Holiday List Assignment') and not frappe.db.exists(
            'Holiday List Assignment', {'applicable_for': 'Company', 'assigned_to': company, 'holiday_list': name, 'docstatus': 1}):
        frappe.get_doc(dict(doctype='Holiday List Assignment', applicable_for='Company', assigned_to=company, holiday_list=name,
                            from_date=date(year, 1, 1))).insert().submit()
    current = frappe.get_cached_value('Company', company, 'default_holiday_list')
    if not current or year == getdate(nowdate()).year:
        # HRMS reads the company default list for leave and attendance; HR Manager may lack Company write.
        frappe.db.set_value('Company', company, 'default_holiday_list', name)
        frappe.clear_document_cache('Company', company)
    return dict(name=name, days=len(rows), weekly_off=sum(1 for r in rows if r[2]))


# Month close ------------------------------------------------------------------------------------

PEOPLE_FIELDS = ['name', 'employee_name', 'status', 'department', 'designation', 'lotus_work_team',
                 'lotus_worker_category', 'lotus_payment_basis', 'lotus_base_salary', 'lotus_daily_rate',
                 'lotus_tax_code', 'lotus_social_insurance_no', 'date_of_joining', 'relieving_date', 'contract_end_date',
                 'lotus_onboarding_state', 'lotus_dependents', 'lotus_pit_commitment']


def _people(company, begin, end):
    rows = frappe.get_all('Employee', filters={'company': company, 'date_of_joining': ['<=', end]},
                          fields=PEOPLE_FIELDS, order_by='employee_name', limit_page_length=0)
    # People who left before the month are out; Inactive means onboarding is not finished.
    return [r for r in rows if r.status == 'Active' or (r.status == 'Left' and r.relieving_date and getdate(r.relieving_date) >= begin)]


def blockers(company, begin, end):
    holidays = holiday_map(company, begin, end)
    count = lambda dt, filters: frappe.db.count(dt, filters)
    drafts = count('Attendance', {'company': company, 'attendance_date': ['between', [begin, end]], 'docstatus': 0})
    crew = count('Lotus Crew Attendance', {'company': company, 'posting_date': ['between', [begin, end]], 'docstatus': 0})
    punches = count('Lotus Attendance Punch', {'company': company, 'state': ['not in', ['Synced', 'Ignored']],
                                               'punched_at': ['between', [str(begin) + ' 00:00:00', str(end) + ' 23:59:59']]})
    people = [p for p in _people(company, begin, end) if p.lotus_payment_basis != 'Piecework']
    records = frappe.get_all('Attendance', filters={'company': company, 'attendance_date': ['between', [begin, end]],
                                                    'docstatus': ['<', 2]},
                             fields=['employee', 'attendance_date'], limit_page_length=0)
    gaps = missing_attendance(people, records, begin, end, holidays.days if holidays else {}, getdate(nowdate()))
    names = {p.name: p.employee_name for p in people}
    rows = close_blockers(drafts, crew, punches, gaps, has_calendar=bool(holidays))
    if frappe.db.exists('DocType', 'Lotus Field Report'):
        pending = frappe.get_all('Lotus Field Report', filters={'company': company,
            'posting_date': ['between', [begin, end]], 'state': ['!=', 'Confirmed']}, pluck='name', limit_page_length=0)
        pieces = frappe.get_all('Lotus Field Report Line', filters={'parent': ['in', pending],
            'parenttype': 'Lotus Field Report', 'pay_kind': 'Piece'}, pluck='parent', limit_page_length=0) if pending else []
        if pieces:
            rows.append(dict(code='draft_piece_reports', label='Báo cáo ngày có công khoán chưa xác nhận', count=len(set(pieces))))
    for row in rows:
        if row.get('employees'):
            row['employees'] = [dict(name=n, employee_name=names.get(n, n), days=[str(d) for d in gaps[n]][:31]) for n in row['employees']]
    return rows


@frappe.whitelist()
def period_status(company, month):
    _hr_manager()
    _company(company)
    begin, end = _month(month)
    doc = frappe.db.get_value(PERIOD, {'company': company, 'month': month},
                              ['name', 'status', 'closed_by', 'closed_at', 'reopen_reason'], as_dict=True)
    period = frappe.get_doc(PERIOD, doc.name) if doc else None
    signing_info = approval(period)
    from urllib.parse import urlencode
    return dict(month=month, status=doc.status if doc else 'Open', closed_by=doc.closed_by if doc else None,
                closed_at=str(doc.closed_at) if doc and doc.closed_at else None,
                reopen_reason=doc.reopen_reason if doc else None, signing=signing_info,
                can_send=not (doc and doc.status == 'Closed') and (not signing_info or signing_info['returned'] or signing_info['done']),
                print_url='/lotus/in-luong?' + urlencode(dict(company=company, month=month)),
                blockers=[] if doc and doc.status == 'Closed' else blockers(company, begin, end),
                future=end >= getdate(nowdate()))


@frappe.whitelist(methods=['POST'])
def close_period(company, month):
    """Close without signatures: Administrator only (recovery). Normally send_period → sign_period closes."""
    if frappe.session.user != 'Administrator':
        frappe.throw('Kỳ công được khóa khi Giám đốc ký duyệt. Dùng Gửi duyệt khóa kỳ.', frappe.PermissionError)
    _company(company)
    begin, end = _month(month)
    if end >= getdate(nowdate()):
        frappe.throw('Chỉ khóa kỳ khi tháng đã kết thúc.')
    from lotus_farm.controls import lock
    lock('Company', company)
    found = blockers(company, begin, end)
    if found:
        frappe.throw('Chưa khóa được kỳ: ' + '; '.join(r['label'] + ' (' + str(r['count']) + ')' for r in found))
    sheet = _sheet(company, month, begin, end)
    _check_sheet_ready(sheet)
    snapshot = json.dumps(sheet, sort_keys=True, default=str, ensure_ascii=False)
    name = frappe.db.get_value(PERIOD, {'company': company, 'month': month})
    doc = frappe.get_doc(PERIOD, name) if name else frappe.new_doc(PERIOD)
    if doc.status == 'Closed':
        frappe.throw('Kỳ đã khóa.')
    doc.update(dict(company=company, month=month, status='Closed', closed_by=frappe.session.user, closed_at=now_datetime(),
                    snapshot=snapshot, snapshot_hash=hashlib.sha256(snapshot.encode()).hexdigest()))
    doc.flags.lotus_period_action = True
    doc.save()
    return dict(name=doc.name, status=doc.status)


@frappe.whitelist(methods=['POST'])
def reopen_period(company, month, reason):
    _hr_manager()
    _company(company)
    _month(month)
    reason = str(reason or '').strip()
    if len(reason) < 5:
        frappe.throw('Ghi lý do mở lại kỳ (ít nhất 5 ký tự).')
    from lotus_farm.controls import lock
    lock('Company', company)
    name = frappe.db.get_value(PERIOD, {'company': company, 'month': month, 'status': 'Closed'})
    if not name:
        frappe.throw('Kỳ chưa khóa.')
    begin, end = _month(month)
    from lotus_farm.finance.payroll import assert_reopenable
    assert_reopenable(name)
    # Posted pay must be reversed in HRMS first, so accounting and attendance never disagree.
    posted = frappe.db.count('Salary Slip', {'company': company, 'start_date': begin, 'end_date': end, 'docstatus': 1})
    if posted:
        frappe.throw('Tháng đã có ' + str(posted) + ' phiếu lương đã duyệt. Hủy phiếu lương và bút toán trong HRMS trước khi mở lại kỳ.')
    doc = frappe.get_doc(PERIOD, name)
    stamp = str(now_datetime())[:16] + ' · ' + frappe.session.user + ': ' + reason[:500]
    doc.update(dict(status='Open', reopen_reason=((doc.reopen_reason or '') + '\n' + stamp).strip(), snapshot=None, snapshot_hash=None))
    doc.flags.lotus_period_action = True
    doc.save()
    return dict(name=doc.name, status=doc.status)


# Approval: the preparer sends the finished month to HR and the director by name. The sheet is frozen
# in a snapshot when sent; the director's signature closes the period if the data still matches it.

def _period(company, month):
    name = frappe.db.get_value(PERIOD, {'company': company, 'month': month})
    return frappe.get_doc(PERIOD, name) if name else None


def _check_sheet_ready(sheet):
    if sheet.get('issues'):
        frappe.throw('Hoàn thiện thông tin lương trước khi gửi/ký khóa kỳ: '+ '; '.join(
            str(r.get('employee_name') or r.get('employee') or '')+': '+str(r['issue']) for r in sheet['issues'][:10]))


def _frozen(company, month, begin, end):
    sheet=_sheet(company,month,begin,end)
    _check_sheet_ready(sheet)
    snapshot = json.dumps(sheet, sort_keys=True, default=str, ensure_ascii=False)
    return snapshot, hashlib.sha256(snapshot.encode()).hexdigest()


def is_signer(doc, user=None):
    """Named signer of the current round (any step); they may view the sheet they are asked to sign."""
    from lotus_farm import signing
    user = user or frappe.session.user
    rows, round_ = signing._rows(doc)
    return bool(round_) and any(r.assignee == user for r in rows)


def _signer_or_hr(doc):
    if frappe.session.user == 'Guest':
        frappe.throw('Vui lòng đăng nhập.', frappe.AuthenticationError)
    hr = frappe.session.user == 'Administrator' or {'HR Manager', 'System Manager'} & set(frappe.get_roles())
    if not (hr or (doc and is_signer(doc))):
        frappe.throw('Chỉ nhân sự hoặc người được chọn ký mới xem được bảng lương này.', frappe.PermissionError)


@frappe.whitelist()
def period_signers(company):
    _hr_manager()
    _company(company)
    from lotus_farm import signing
    return signing.candidates('timekeeping_period', company)


def approval(doc):
    from lotus_farm import signing
    if not doc:
        return None
    info = signing.status(doc, 'timekeeping_period')
    return info if info['round'] else None


@frappe.whitelist(methods=['POST'])
def send_period(company, month, signers):
    """Freeze the month's sheet and send it to the chosen HR person and director, signing as preparer."""
    _hr_manager()
    _company(company)
    begin, end = _month(month)
    if end >= getdate(nowdate()):
        frappe.throw('Chỉ gửi duyệt khi tháng đã kết thúc.')
    from lotus_farm.controls import lock
    lock('Company', company)
    found = blockers(company, begin, end)
    if found:
        frappe.throw('Chưa gửi duyệt được: ' + '; '.join(r['label'] + ' (' + str(r['count']) + ')' for r in found))
    doc = _period(company, month) or frappe.new_doc(PERIOD)
    if doc.status == 'Closed':
        frappe.throw('Kỳ đã khóa.')
    snapshot, digest = _frozen(company, month, begin, end)
    doc.update(dict(company=company, month=month, status='Open', snapshot=snapshot, snapshot_hash=digest))
    doc.flags.lotus_period_action = True
    doc.save()
    from lotus_farm import signing
    signing.start(doc, signers)
    return period_status(company, month)


@frappe.whitelist(methods=['POST'])
def sign_period(company, month):
    """The named signer signs; the last signature closes the period with the snapshot that was sent."""
    _company(company)
    begin, end = _month(month)
    from lotus_farm.controls import lock
    lock('Company', company)
    doc = _period(company, month)
    if not doc or doc.status == 'Closed' or not approval(doc):
        frappe.throw('Kỳ này không có bảng công đang chờ ký.')
    found=blockers(company,begin,end)
    if found:
        frappe.throw('Chưa ký khóa được kỳ: '+ '; '.join(r['label']+' ('+str(r['count'])+')' for r in found))
    if _frozen(company, month, begin, end)[1] != doc.snapshot_hash:
        frappe.throw('Chấm công hoặc bảng lương đã thay đổi sau khi gửi duyệt. Trả lại để người lập gửi duyệt lại.')
    from lotus_farm import signing
    info = signing.sign(doc)
    if info['done']:
        doc.reload()
        doc.update(dict(status='Closed', closed_by=frappe.session.user, closed_at=now_datetime()))
        doc.flags.lotus_period_action = True
        doc.flags.ignore_permissions = True  # the director signs as the named last signer, not as HR
        doc.save()
    return signing.status(doc, 'timekeeping_period')


def punch_period_guard(doc, method=None):
    """Raw late punches remain ingestible; serialize them with the final period check."""
    company=doc.get('company') or (frappe.db.get_value('Employee',doc.employee,'company') if doc.get('employee') else None)
    if company:
        from lotus_farm.controls import lock
        lock('Company',company)


def salary_rate_guard(doc,method=None):
    punch_period_guard(doc,method)
    if doc.get('company') and doc.get('from_date') and is_closed(doc.company,doc.from_date):
        frappe.throw('Kỳ tại ngày hiệu lực lương đã khóa. Mở lại kỳ trước khi sửa mức lương.')


@frappe.whitelist(methods=['POST'])
def return_period(company, month, reason):
    _company(company)
    from lotus_farm.controls import lock
    lock('Company', company)
    doc = _period(company, month)
    if not doc or doc.status == 'Closed' or not approval(doc):
        frappe.throw('Kỳ này không có bảng công đang chờ ký.')
    from lotus_farm import signing
    return signing.send_back(doc, reason)


def signed_sheet(company, month):
    """Sheet for the print / sign page: HR, or a named signer of the period. Frozen snapshot while signing."""
    _company(company)
    begin, end = _month(month)
    doc = _period(company, month)
    _signer_or_hr(doc)
    if doc and doc.snapshot and (doc.status == 'Closed' or approval(doc)):
        if hashlib.sha256(doc.snapshot.encode()).hexdigest() != doc.snapshot_hash:
            frappe.throw('Dữ liệu kỳ không khớp mã kiểm tra. Mở lại và gửi duyệt lại kỳ.')
        return doc, dict(json.loads(doc.snapshot), draft=doc.status != 'Closed', closed_at=str(doc.closed_at or '') or None)
    return doc, dict(_sheet(company, month, begin, end), draft=True, closed_at=None)


# Payroll sheet ----------------------------------------------------------------------------------

def month_hours(company, begin, end, people):
    """Approved hours per employee for [begin, end]; shared by the payroll sheet and the HRMS salary slip hook."""
    holidays = holiday_map(company, begin, end)
    days = holidays.days if holidays else {}
    names = [p.name for p in people] or ['']
    records = frappe.get_all('Attendance', filters={'company': company, 'employee': ['in', names],
                                                    'attendance_date': ['between', [begin, end]], 'docstatus': 1},
                             fields=['employee', 'attendance_date', 'status', 'leave_type', 'working_hours', 'shift',
                                     'in_time', 'out_time', 'lotus_overtime_hours as overtime_hours', 'lotus_shift_start', 'lotus_shift_end'],
                             limit_page_length=0)
    shifts = {s.name: s for s in frappe.get_all('Shift Type', filters={'name': ['in', list({r.shift for r in records if r.shift}) or ['']]},
                                                fields=['name', 'start_time', 'end_time'])}
    # Days recorded with the shift times they were worked under use those, not the shift as edited later.
    for row in records:
        if row.get('lotus_shift_start') is not None and row.get('lotus_shift_end') is not None:
            key = 'worked:%s-%s' % (row.lotus_shift_start, row.lotus_shift_end)
            shifts.setdefault(key, frappe._dict(name=key, start_time=row.lotus_shift_start, end_time=row.lotus_shift_end))
            row.shift = key
    employment = {p.name: (max(begin, getdate(p.date_of_joining)) if p.date_of_joining else begin,
                           min(end, getdate(p.relieving_date)) if p.relieving_date else end) for p in people}
    unpaid = frappe.get_all('Leave Type', filters={'is_lwp': 1}, pluck='name')
    return hour_summary(records, days, unpaid, shifts, employment), holidays


def _sheet(company, month, begin, end):
    people = _people(company, begin, end)
    records = frappe.get_all('Attendance', filters={'company': company, 'attendance_date': ['between', [begin, end]], 'docstatus': 1},
                             fields=['employee'], limit_page_length=0)
    worked = {r.employee for r in records}
    known = {p.name for p in people}
    # An employee with approved attendance stays on the sheet even if the profile changed later.
    if worked - known:
        people += frappe.get_all('Employee', filters={'name': ['in', list(worked - known)]}, fields=PEOPLE_FIELDS)
    summaries, holidays = month_hours(company, begin, end, people)
    params = params_on(end)
    region = frappe.db.get_single_value('Lotus Finance Settings', 'minimum_wage_region') \
        if frappe.db.exists('DocType', 'Lotus Finance Settings') and frappe.get_meta('Lotus Finance Settings').has_field('minimum_wage_region') else None
    employees = []
    rates={}
    for rate in frappe.get_all('Salary Structure Assignment',filters={'company':company,'employee':['in',[p.name for p in people] or ['']],'docstatus':1},
            fields=['name','employee','base','from_date','lotus_payment_basis'],order_by='from_date asc, name',limit_page_length=0):
        rates.setdefault(rate.employee,[]).append(rate)
    for person in people:
        summary = summaries.get(person.name, empty_summary())
        months = None
        if person.contract_end_date and person.date_of_joining:
            span = getdate(person.contract_end_date) - getdate(person.date_of_joining)
            months = span.days / 30.4
        tax = dict(method=tax_method(person.lotus_worker_category, months), dependents=cint(person.lotus_dependents),
                   commitment=bool(cint(person.lotus_pit_commitment)))
        fallback=dict(name='Employee:'+person.name,from_date=person.date_of_joining or begin,
            base=pay_base(person.lotus_payment_basis,person.lotus_base_salary,person.lotus_daily_rate),lotus_payment_basis=person.lotus_payment_basis)
        rate_begin=max(begin,getdate(person.date_of_joining)) if person.date_of_joining else begin
        rate_end=min(end,getdate(person.relieving_date)) if person.relieving_date else end
        intervals=salary_segments(rates.get(person.name,[]),rate_begin,rate_end,fallback)
        for interval in intervals:
            interval['summary']=summary if len(intervals)==1 else month_hours(company,interval['start'],interval['end'],[person])[0].get(person.name,empty_summary())
        pay=segmented_pay(intervals,insured=bool(person.lotus_social_insurance_no),params=params,region=region,tax=tax)
        employees.append(dict(employee=person.name, employee_name=person.employee_name, department=person.department,
                              designation=person.designation, work_team=person.lotus_work_team,
                              worker_category=person.lotus_worker_category, payment_basis=intervals[-1]['basis'],
                              base_salary=flt(pay['base']) if intervals[-1]['basis']!='Daywork' else 0,
                              daily_rate=flt(pay['base'])/STANDARD_DAYS if intervals[-1]['basis']=='Daywork' else 0,
                              tax_code=person.lotus_tax_code, social_insurance_no=person.lotus_social_insurance_no,
                              dependents=tax['dependents'], pit_commitment=tax['commitment'],
                              onboarding=person.lotus_onboarding_state, status=person.status, **summary, **pay))
    # Phụ cấp công việc (not taxed, outside the insurance base) is added after the tax; tạm ứng is deducted from it.
    from lotus_farm.work_allowance import month_advances, month_allowance
    names = [e['employee'] for e in employees]
    allowances, advances = month_allowance(company, begin, end, names), month_advances(company, month)
    team_names = {t.name: t.team_name for t in frappe.get_all('Work Team', filters={'name': ['in', list({e['work_team'] for e in employees if e['work_team']}) or ['']]},
                                                               fields=['name', 'team_name'])}
    for e in employees:
        extra = allowances.get(e['employee']) or {}
        e['wage'] = flt(e['gross'])
        e['allowance_days'], e['allowance'] = round(flt(extra.get('days')), 2), round(flt(extra.get('amount')))
        e['gross'] = flt(e['gross']) + e['allowance']
        e['net'] = flt(e['net']) + e['allowance']
        e['advance'] = round(flt(advances.get(e['employee'])))
        e['take_home'] = e['net'] - e['advance']
        e['group'] = team_names.get(e['work_team']) or e['work_team'] or (e['department'] or '').rsplit(' - ', 1)[0] or 'Chưa xếp tổ'
    crews = {}
    for row in frappe.get_all('Lotus Crew Attendance', filters={'company': company, 'posting_date': ['between', [begin, end]], 'docstatus': 1},
                              fields=['work_team', 'work_type', 'uom', 'unit_rate', 'completed_qty', 'workdays', 'present', 'amount'],
                              limit_page_length=0):
        key = (row.work_team, row.work_type or '', row.uom or '', flt(row.unit_rate))
        item = crews.setdefault(key, dict(work_team=row.work_team, work_type=row.work_type, uom=row.uom, unit_rate=flt(row.unit_rate),
                                          completed_qty=0.0, workdays=0.0, person_days=0, amount=0.0))
        item['completed_qty'] += flt(row.completed_qty)
        item['workdays'] += flt(row.workdays)
        item['person_days'] += cint(row.present)
        item['amount'] += flt(row.amount)
    # Công khoán confirmed in the đội daily reports (báo cáo ngày) is owed to the tổ the same way.
    from lotus_farm.field_report import confirmed_piece_lines
    for row in confirmed_piece_lines(company, begin, end):
        key = (row.work_team, row.work_type or '', row.uom or '', flt(row.rate))
        item = crews.setdefault(key, dict(work_team=row.work_team, work_type=row.work_type, uom=row.uom, unit_rate=flt(row.rate),
                                          completed_qty=0.0, workdays=0.0, person_days=0, amount=0.0))
        item['completed_qty'] += flt(row.qty)
        item['workdays'] += flt(row.workdays)
        item['amount'] += flt(row.amount)
    teams = {t.name: t for t in frappe.get_all('Work Team', filters={'name': ['in', list({k[0] for k in crews}) or ['']]},
                                                fields=['name', 'team_name', 'lotus_foreman', 'lotus_foreman_phone'])}
    works = {w.name: w.work_name for w in frappe.get_all('Work Type', filters={'name': ['in', list({k[1] for k in crews}) or ['']]},
                                                          fields=['name', 'work_name'])}
    crew_rows = []
    for key in sorted(crews):
        item = crews[key]
        team = teams.get(item['work_team']) or frappe._dict()
        crew_rows.append(dict(item, team_name=team.team_name or item['work_team'], foreman=team.lotus_foreman,
                              foreman_phone=team.lotus_foreman_phone, work_name=works.get(item['work_type']) or item['work_type'] or '',
                              amount=round(item['amount'])))
    issues = [dict(employee=e['employee'], employee_name=e['employee_name'], issue=text) for e in employees for text in e['issues']]
    issues += [dict(employee=e['employee'], employee_name=e['employee_name'], issue='Đang tính công nhưng hồ sơ chưa duyệt tiếp nhận')
               for e in employees if e['status'] == 'Active' and e['onboarding'] != 'Ready']
    return dict(month=month, company=company, currency=frappe.get_cached_value('Company', company, 'default_currency'),
                standard_days=STANDARD_DAYS, hours_base=HOURS_BASE, holiday_list=holidays.name if holidays else None,
                employees=employees, crews=crew_rows, issues=issues,
                totals=dict(employees=len(employees), gross=sum(e['gross'] for e in employees), net=sum(e['net'] for e in employees),
                            employee_insurance=sum(e['employee_insurance'] for e in employees),
                            employer_insurance=sum(e['employer_insurance'] for e in employees),
                            pit=sum(e.get('pit') or 0 for e in employees), union_fee=sum(e.get('union_fee') or 0 for e in employees),
                            crew_amount=sum(c['amount'] for c in crew_rows), workdays=sum(e['workdays'] for e in employees),
                            allowance=sum(e['allowance'] for e in employees), advance=sum(e['advance'] for e in employees),
                            take_home=sum(e['take_home'] for e in employees)))


@frappe.whitelist()
def payroll_sheet(company, month):
    _hr_manager()
    _company(company)
    begin, end = _month(month)
    period = frappe.db.get_value(PERIOD, {'company': company, 'month': month}, ['status', 'snapshot', 'snapshot_hash', 'closed_at'], as_dict=True)
    if period and period.status == 'Closed' and period.snapshot:
        if hashlib.sha256(period.snapshot.encode()).hexdigest() != period.snapshot_hash:
            frappe.throw('Dữ liệu kỳ đã khóa không khớp mã kiểm tra. Mở lại và khóa lại kỳ.')
        return dict(json.loads(period.snapshot), draft=False, closed_at=str(period.closed_at))
    return dict(_sheet(company, month, begin, end), draft=True, closed_at=None)


def _cell(value):
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + value
    return value


@frappe.whitelist()
def export_payroll(company, month):
    data = payroll_sheet(company, month)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    banner = 'BẢN NHÁP — kỳ chưa khóa, số liệu còn thay đổi' if data['draft'] else 'Kỳ đã khóa lúc ' + str(data['closed_at'])[:16]
    labels = {'Monthly': 'Lương tháng', 'Daywork': 'Công nhật', 'Piecework': 'Khoán sản lượng', 'Mixed': 'Hỗn hợp'}

    def sheet(ws, title, header, rows, note='Tiền công = công đã duyệt × đơn giá công. Phụ cấp công việc không tính thuế TNCN, '
              'không tính vào lương đóng BHXH. Thực lĩnh = tiền công + phụ cấp − BH người LĐ − thuế TNCN − tạm ứng đã chi.'):
        ws.title = title
        ws.append(['Bảng lương ' + str(month) + ' · ' + str(company) + ' · ' + banner])
        ws.append([note])
        ws.append(header)
        for cell in ws[3]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_cell(v) for v in row])
        ws.freeze_panes = 'C4'

    from lotus_farm.www.lotus_payroll_print import layout
    plan = layout(data)
    sheet(book.active, 'Công nhân',
          ['Tổ', 'Mã NV', 'Họ tên', 'Công', 'Đơn giá công', 'Tiền công', 'Ngày phụ cấp', 'Phụ cấp', 'Thuế TNCN', 'Tạm ứng', 'Thực lĩnh', 'Ký nhận'],
          [[g.name, e['employee'], e['employee_name'], e['workdays'], e['daily_rate'], e['wage'], e['allowance_days'], e['allowance'],
            e['pit'], e['advance'], e['take_home'], ''] for g in plan['worker_groups'] for e in g.rows])
    sheet(book.create_sheet(), 'Lương tháng',
          ['Bộ phận', 'Mã NV', 'Họ tên', 'Lương CB (26 công)', 'Công', 'Lương theo công', 'Phụ cấp', 'BH người LĐ', 'Thuế TNCN', 'Tạm ứng', 'Thực lĩnh', 'Ký nhận'],
          [[e['group'], e['employee'], e['employee_name'], e['base'], e['workdays'], e['wage'], e['allowance'], e['employee_insurance'],
            e['pit'], e['advance'], e['take_home'], ''] for e in plan['staff_rows']])
    sheet(book.create_sheet(), 'Chi tiết (kế toán)',
          ['Mã NV', 'Họ tên', 'Bộ phận', 'Tổ/đội', 'Loại lao động', 'Cơ chế lương', 'Lương CB (26 công)', 'Đơn giá giờ',
           'Công', 'Giờ công thường', 'Giờ đêm', 'Giờ TC ngày thường', 'Giờ làm ngày nghỉ tuần', 'Giờ làm ngày lễ',
           'Phép có lương (ngày)', 'Lễ hưởng lương (ngày)', 'Nghỉ không lương', 'Vắng',
           'Lương giờ công', 'Lương lễ, phép', 'Phụ cấp ca đêm', 'Tiền tăng ca', 'Phụ cấp công việc', 'Tổng thu nhập',
           'BHXH+BHYT+BHTN NLĐ', 'Người phụ thuộc', 'Thu nhập tính thuế', 'Thuế TNCN', 'Thực lĩnh',
           'BH doanh nghiệp đóng', 'KPCĐ 2% (DN)', 'MST', 'Số BHXH', 'Ghi chú'],
          [[e['employee'], e['employee_name'], e['department'], e['work_team'], e['worker_category'],
            labels.get(e['payment_basis'], e['payment_basis']), e['base'], e['hour_rate'],
            e['workdays'], e['regular_hours'], e['night_hours'], e['ot_regular'], e['ot_weekly_off'], e['ot_holiday'],
            e['paid_leave_days'], e['paid_holiday_days'], e['unpaid_leave_days'], e['absent_days'],
            e['regular_pay'], e['paid_pay'], e['night_pay'], e['overtime_pay'], e.get('allowance') or 0, e['gross'],
            e['employee_insurance'], e.get('dependents') or 0, e.get('taxable_income') or 0, e.get('pit') or 0, e['net'],
            e['employer_insurance'], e.get('union_fee') or 0, e['tax_code'], e['social_insurance_no'],
            'Trả qua đội khoán' if e['via_crew'] else '; '.join(e['issues'])] for e in data['employees']])
    sheet(book.create_sheet(), 'Đội khoán',
          ['Đội', 'Đầu câu', 'Điện thoại', 'Công việc', 'Khối lượng', 'Đơn vị', 'Đơn giá', 'Thành tiền', 'Tổng công', 'Người-ngày'],
          [[c['team_name'], c['foreman'], c['foreman_phone'], c['work_name'], c['completed_qty'], unit_label(c['uom']), c['unit_rate'],
            c['amount'], c['workdays'], c['person_days']] for c in data['crews']])
    sheet(book.create_sheet(), 'Kiểm tra', ['Mã NV', 'Họ tên', 'Vấn đề'],
          [[i['employee'], i['employee_name'], i['issue']] for i in data['issues']])
    from io import BytesIO
    output = BytesIO()
    book.save(output)
    frappe.local.response.filename = 'bang-luong-' + str(month) + '.xlsx'
    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.type = 'download'
