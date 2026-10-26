"""Dự thu theo nhật ký chặt mẹ lần 2: the journal of second mother-plant cuts by zone, the parameter
profile of each farm (versioned, an approved version is frozen), the forecast computed from both, and the
Excel upload (preview every row, then write the new ones and keep the file as evidence).

Everyone who may read Harvest Forecast sees the forecast of the farms they are permitted; Administrator,
System Manager and Lotus Farm Manager import, undo an import and change parameters. Formulas live in
lotus_farm.cut_forecast_rules and match the Lotus forecast workbook.
"""
import base64
import hashlib
import json
from io import BytesIO

import frappe
from frappe.utils import cint, now_datetime

from lotus_farm import cut_forecast_rules as rules
from lotus_farm import forecast_book2_reference as book2
from lotus_farm.farm_land import can_edit

ENTRY = 'Lotus Cut Journal Entry'
IMPORT = 'Lotus Cut Journal Import'
PROFILE = 'Lotus Cut Forecast Profile'
PARAM_FIELDS = [f for f, _, _, _ in rules.PARAMS]


def _login():
    if frappe.session.user == 'Guest':
        frappe.throw('Vui lòng đăng nhập.', frappe.AuthenticationError)


def _farm(farm):
    doc = frappe.get_doc('Farm', farm)
    doc.check_permission('read')
    return doc


def _reader(farm):
    _login()
    if not (can_edit() or frappe.has_permission('Harvest Forecast', 'read')):
        frappe.throw('Bạn chưa có quyền xem dự thu.', frappe.PermissionError)
    return _farm(farm)


def _editor(farm):
    _login()
    if not can_edit():
        frappe.throw('Chỉ Quản lý vườn, Quản trị hệ thống hoặc Administrator được nhập nhật ký và sửa thông số dự thu.', frappe.PermissionError)
    return _farm(farm)


def _entries(farm, zone=None):
    filters = {'farm': farm}
    if zone:
        filters['farm_zone'] = zone
    rows = frappe.get_all(ENTRY, filters=filters, fields=['name', 'farm_zone', 'farm_block', 'crop_cycle', 'plan_cohort', 'cut_date', 'tree_qty', 'source_ref', 'entry_key', 'import_batch', 'modified'],
                          order_by='farm_zone, cut_date, source_ref', limit_page_length=0)
    return [dict(name=r.name, zone=r.farm_zone, farm_block=r.farm_block, crop_cycle=r.crop_cycle,
                 plan_cohort=r.plan_cohort, modified=r.modified,
                 cut_date=r.cut_date, tree_qty=r.tree_qty, source_ref=r.source_ref or '',
                 entry_key=r.entry_key, import_batch=r.import_batch)
            for r in rows]


def _profiles(farm):
    return frappe.get_all(PROFILE, filters={'farm': farm}, fields=['name', 'version', 'approval_state', 'source', 'source_import', 'approved_by', 'approved_at',
                                                                   'notes', 'modified', *PARAM_FIELDS], order_by='version desc', limit_page_length=0)


def _active(profiles):
    return next((p for p in profiles if p.approval_state == 'Approved'), None)


def _profile_values(doc):
    return {f: doc.get(f) for f in PARAM_FIELDS}


def _jsonable(value):
    return json.loads(json.dumps(value, default=str))


@frappe.whitelist()
def board(farm=None, zone=None):
    """Farms the user may see; for one farm: the active profile, totals, months, zones, the journal with each wave's dates."""
    _login()
    farms = frappe.get_list('Farm', fields=['name', 'farm_name'], order_by='name', limit_page_length=0)
    with_journal = set(frappe.get_all(ENTRY, filters={'farm': ['in', [f.name for f in farms] or ['']]}, pluck='farm', distinct=True))
    for f in farms:
        f['has_journal'] = f.name in with_journal
    if not farm:
        farm = next((f.name for f in farms if f.has_journal), None) or (farms[0].name if farms else None)
    if not farm:
        return dict(farms=[], farm=None, can_edit=can_edit(), params=rules.PARAMS)
    _reader(farm)
    entries = _entries(farm)
    zone_list = sorted({e['zone'] for e in entries})
    shown = [e for e in entries if not zone or e['zone'] == zone]
    profiles = _profiles(farm)
    active = _active(profiles)
    # Expose only permission-filtered options; legacy zone-only rows stay unresolved.
    blocks = frappe.get_list('Farm Block', filters={'farm': farm}, fields=['name','block_code','block_name','farm_zone'],
                             order_by='farm_zone,block_code', limit_page_length=0) if frappe.has_permission('Farm Block', 'read') else []
    cycles = frappe.get_list('Crop Cycle', filters={'farm': farm},
                             fields=['name','farm_block','cycle_no','status'], limit_page_length=0) if frappe.has_permission('Crop Cycle', 'read') else []
    cohorts = frappe.get_list('Plan Cohort', filters={'farm': farm},
                              fields=['name','farm_block','crop_cycle','cohort_code','cohort_name','approval_state'],
                              limit_page_length=0) if frappe.db.exists('DocType','Plan Cohort') and frappe.has_permission('Plan Cohort','read') else []
    issues = []
    if not entries:
        issues.append('Vườn này chưa có nhật ký chặt mẹ lần 2. Bấm "Nhập Excel" để tải file dự thu lên.')
    if not active:
        issues.append('Chưa có bộ thông số đã duyệt nên chưa tính được dự thu.')
    unmapped_count = sum(not bool(e.get('farm_block')) for e in shown)
    if unmapped_count:
        issues.append(f'{unmapped_count} dòng nhật ký chưa gắn lô; không được tự phân bổ cây từ khu sang lô.')
    result = rules.compute(shown, _profile_values(active)) if active and shown else dict(rows=[], months=[], zones=[], zone_months={}, totals=None)
    imports = frappe.get_all(IMPORT, filters={'farm': farm}, fields=['name', 'file_name', 'source_file', 'sha256', 'new_rows', 'existing_rows', 'conflict_rows',
                                                                      'error_rows', 'trees', 'profile', 'imported_by', 'imported_at', 'undone_at'],
                             order_by='imported_at desc', limit_page_length=20)
    return _jsonable(dict(farms=farms, farm=farm, zone=zone or '', zone_list=zone_list, can_edit=can_edit(), params=rules.PARAMS, stages=rules.STAGES,
                          profile=active, profiles=profiles, issues=issues, imports=imports,
                           cut_entries=shown, blocks=blocks, cycles=cycles, cohorts=cohorts,
                           unmapped_count=unmapped_count,
                           book2_reference=book2.reference(_profile_values(active) if active else None), **result))


# Excel upload ----------------------------------------------------------------------------------------

def _cache_key(token):
    return 'lotus_cut_import:' + frappe.session.user + ':' + str(token)


@frappe.whitelist(methods=['POST'])
def preview_import(farm, content, file_name=None):
    """content: base64 of the forecast workbook -> new / already imported / conflicting rows and parameter changes; nothing is saved."""
    _editor(farm)
    raw = base64.b64decode(content.split(',', 1)[-1])
    try:
        parsed = rules.parse_workbook(raw)
    except ValueError as exc:
        frappe.throw(str(exc))
    except Exception:
        frappe.throw('Không đọc được file. Hãy chọn file Excel (.xlsx) dự thu có sheet "NHẬT KÝ GỐC".')
    known_zones = set(frappe.get_all('Farm Block', filters={'farm': farm}, pluck='farm_zone', distinct=True)) - {None, ''}
    existing = {e['entry_key']: e for e in _entries(farm)}
    rows, new = [], []
    for e in parsed['entries']:
        key = rules.entry_key(farm, e['zone'], e['cut_date'], e['source_ref'])
        row = dict(line=e['line'], zone=e['zone'], cut_date=str(e['cut_date']), tree_qty=e['tree_qty'], source_ref=e['source_ref'], entry_key=key, errors=[])
        if known_zones and e['zone'] not in known_zones:
            row.update(state='error', errors=['Khu %s không có trong vườn' % e['zone']])
        elif key in existing:
            old = existing[key]['tree_qty']
            row.update(state='same' if old == e['tree_qty'] else 'conflict', old_qty=old)
        else:
            row['state'] = 'new'
            new.append(row)
        rows.append(row)
    for bad in parsed['entry_errors']:
        rows.append(dict(line=bad['line'], state='error', errors=bad['errors']))
    rows.sort(key=lambda r: r['line'])
    active = _active(_profiles(farm))
    current = _profile_values(active) if active else {}
    param_changes = [dict(field=f, label=rules.LABELS[f], old=current.get(f), new=v) for f, v in parsed['params'].items()
                     if not active or abs(float(current.get(f) or 0) - float(v)) > 1e-9]
    preview_result = rules.compute([dict(e, zone=e['zone']) for e in parsed['entries']], parsed['params']) if parsed['params'] and not parsed['param_errors'] else None
    token = frappe.generate_hash(length=20)
    frappe.cache.set_value(_cache_key(token), json.dumps(dict(farm=farm, file_name=file_name or 'du-thu.xlsx', content=base64.b64encode(raw).decode(),
                                                             sha256=hashlib.sha256(raw).hexdigest(), rows=new, params=parsed['params'] if not parsed['param_errors'] else None,
                                                             counts=dict(existing=sum(1 for r in rows if r['state'] == 'same'), conflict=sum(1 for r in rows if r['state'] == 'conflict'),
                                                                         error=sum(1 for r in rows if r['state'] == 'error'))), default=str), expires_in_sec=1800)
    count = lambda state: sum(1 for r in rows if r['state'] == state)
    return _jsonable(dict(token=token, rows=rows, new=count('new'), same=count('same'), conflict=count('conflict'), errors=count('error'),
                          new_trees=sum(r['tree_qty'] for r in new), param_changes=param_changes, param_errors=parsed['param_errors'],
                          has_profile=bool(active), file_totals=preview_result['totals'] if preview_result else None))


@frappe.whitelist(methods=['POST'])
def commit_import(farm, token, apply_params=1):
    """Write the new rows of the preview, keep the uploaded file and, if asked, approve its parameters as the next version."""
    _editor(farm)
    raw = frappe.cache.get_value(_cache_key(token))
    if not raw:
        frappe.throw('Bản xem trước đã hết hạn. Tải file lên lại.')
    plan = json.loads(raw)
    if plan['farm'] != farm:
        frappe.throw('Bản xem trước thuộc vườn khác.')
    from lotus_farm.controls import lock
    lock('Farm', farm)
    batch = frappe.get_doc(dict(doctype=IMPORT, farm=farm, file_name=plan['file_name'], sha256=plan['sha256'], new_rows=0, existing_rows=plan['counts']['existing'],
                                conflict_rows=plan['counts']['conflict'], error_rows=plan['counts']['error'], trees=0, imported_by=frappe.session.user,
                                imported_at=now_datetime()))
    batch.flags.ignore_permissions = True
    batch.insert()
    file = frappe.get_doc(dict(doctype='File', file_name=plan['file_name'], attached_to_doctype=IMPORT, attached_to_name=batch.name, is_private=1,
                               content=base64.b64decode(plan['content'])))
    file.flags.ignore_permissions = True
    file.insert()
    taken = set(frappe.get_all(ENTRY, filters={'farm': farm}, pluck='entry_key'))
    saved = trees = 0
    for row in plan['rows']:
        if row['entry_key'] in taken:
            continue
        doc = frappe.get_doc(dict(doctype=ENTRY, farm=farm, farm_zone=row['zone'], cut_date=row['cut_date'], tree_qty=row['tree_qty'],
                                  source_ref=row['source_ref'], import_batch=batch.name))
        doc.flags.ignore_permissions = True
        doc.insert()
        taken.add(doc.entry_key)
        saved += 1
        trees += row['tree_qty']
    profile = None
    if cint(apply_params) and plan.get('params'):
        active = _active(_profiles(farm))
        if not active or any(abs(float(active.get(f) or 0) - float(v)) > 1e-9 for f, v in plan['params'].items()):
            profile = _new_profile(farm, plan['params'], source='Excel', source_import=batch.name, notes='Lấy từ file ' + plan['file_name'])
    batch.db_set(dict(source_file=file.file_url, new_rows=saved, trees=trees, profile=profile.name if profile else None))
    frappe.cache.delete_value(_cache_key(token))
    return dict(saved=saved, trees=trees, profile=profile.version if profile else None, batch=batch.name)


@frappe.whitelist(methods=['POST'])
def undo_import(farm, batch):
    """Remove the journal rows one upload added (the upload record and its file stay as evidence)."""
    _editor(farm)
    doc = frappe.get_doc(IMPORT, batch)
    if doc.farm != farm:
        frappe.throw('Lần nhập thuộc vườn khác.')
    names = frappe.get_all(ENTRY, filters={'import_batch': batch}, pluck='name')
    for name in names:
        frappe.delete_doc(ENTRY, name, ignore_permissions=True)
    doc.db_set(dict(undone_at=now_datetime(), undone_by=frappe.session.user))
    return dict(removed=len(names))


@frappe.whitelist(methods=['POST'])
def assign_entry_scope(farm, entry, farm_block=None, crop_cycle=None, plan_cohort=None, modified=None):
    """Manually reconcile ONE legacy cut row. Never infer a block from the zone/area."""
    _editor(farm)
    if not modified:
        frappe.throw("Thiếu phiên bản nhật ký để đối chiếu thay đổi.")
    from lotus_farm.controls import lock
    lock('Farm', farm)
    doc = frappe.get_doc(ENTRY, entry, for_update=True)
    if doc.farm != farm:
        frappe.throw("Nhật ký không thuộc vườn này.", frappe.PermissionError)
    if str(doc.modified) != str(modified):
        frappe.throw("Nhật ký đã được người khác sửa. Tải lại và kiểm tra trước khi ghi.")
    block = (farm_block or '').strip() or None
    cycle = (crop_cycle or '').strip() or None
    cohort = (plan_cohort or '').strip() or None
    if (cycle or cohort) and not block:
        frappe.throw("Chọn lô lớn trước khi gắn vụ/đợt cây.")
    doc.farm_block = block
    doc.crop_cycle = cycle
    doc.plan_cohort = cohort
    # DocType.validate rejects cross-farm/zone/cycle links; key stays source-based for dedupe.
    doc.save(ignore_permissions=True)
    return dict(entry=doc.name, farm_block=doc.farm_block, crop_cycle=doc.crop_cycle,
                plan_cohort=doc.plan_cohort, modified=str(doc.modified))


# Care-plan preview and explicitly requested DRAFT creation ---------------------------------------

def _care_plan_preview(farm, from_date, to_date, farm_block=None):
    """Shared source for read preview and POST, without any implicit write."""
    from datetime import date as _date
    from lotus_farm.cut_work_schedule import work_calendar
    from lotus_farm.cut_forecast_rules import as_date

    _reader(farm)
    start, end = as_date(from_date), as_date(to_date)
    if end < start or (end - start).days > 45:
        frappe.throw("Chỉ xem lịch trong một khoảng từ 1 đến 46 ngày.")
    if farm_block:
        block = frappe.get_doc("Farm Block", farm_block)
        block.check_permission("read")
        if block.farm != farm:
            frappe.throw("Lô không thuộc vườn.")
    active = _active(_profiles(farm))
    if not active:
        frappe.throw("Chưa có bộ thông số dự thu đã duyệt.")
    all_entries = _entries(farm)
    rows, issues, unmapped = work_calendar(
        all_entries, _profile_values(active), from_date=start, to_date=end,
        farm_block=farm_block)
    # Work Type mapping is an explicit code match; fuzzy name lookup is forbidden.
    works = frappe.get_list("Work Type", fields=[
        "name", "work_code", "work_name", "active", "uom",
        "productivity_per_workday", "allow_daywork"
    ], filters={"active": 1}, limit_page_length=0)
    by_code = {}
    for w in works:
        key = (w.work_code or "").strip().upper()
        if not key:
            continue
        by_code.setdefault(key, []).append(w)
    ready_count = 0
    for row in rows:
        matches = by_code.get(row["work_code"], [])
        work = matches[0] if len(matches) == 1 else None
        if not work:
            row["blocking_issue"] = ("Trùng mã công việc trong ERP." if matches else "Chưa có mã công việc trong ERP.")
        elif not work.allow_daywork:
            row["blocking_issue"] = "Chưa cho phép công nhật; chọn hình thức trả công trong danh mục trước."
        elif float(work.productivity_per_workday or 0) <= 0:
            row["blocking_issue"] = "Chưa xác nhận năng suất/công."
        elif (work.uom or "").strip().casefold() != row["uom"].casefold():
            row["blocking_issue"] = "Đơn vị của công việc chưa khớp Book2; cần đối chiếu quy đổi."
        else:
            row["blocking_issue"] = None
            ready_count += 1
        row["work_type"] = work.name if work else None
        row["work_name"] = work.work_name if work else None
        row["required_workdays"] = round(row["planned_qty"] / float(work.productivity_per_workday), 4) if not row["blocking_issue"] else None
    signatures = sorted([{
        "block": r["farm_block"], "day": r["date"], "code": r["work_code"],
        "qty": r["planned_qty"], "key": r["schedule_key"],
        "work_type": r["work_type"], "issue": r["blocking_issue"],
    } for r in rows], key=lambda x:(x["day"],x["block"],x["code"]))
    fingerprint = hashlib.sha256(json.dumps({
        "farm":farm, "from":str(start), "to":str(end),
        "block":farm_block, "profile":active.name,
        "entries":[(e["entry_key"],str(e["modified"]),e.get("farm_block")) for e in all_entries],
        "lines":signatures,
    }, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return dict(farm=farm, from_date=str(start), to_date=str(end),
                farm_block=farm_block, profile=active.name, profile_version=active.version,
                unmapped_count=unmapped, ready_count=ready_count,
                total_count=len(rows), warnings=issues, rows=rows, fingerprint=fingerprint)


@frappe.whitelist()
def preview_care_plan(farm, from_date, to_date, farm_block=None):
    """Only read: time-window and mapped-block care plan proposed from four waves."""
    return _jsonable(_care_plan_preview(farm, from_date, to_date, farm_block))


@frappe.whitelist(methods=['POST'])
def create_care_plan_draft(farm, farm_block, posting_date, fingerprint):
    """Create at most one day of DRAFT work, never submit, post payroll or issue stock."""
    from lotus_farm.cut_forecast_rules import as_date
    from lotus_farm.controls import lock

    _editor(farm)
    if not frappe.has_permission("Daily Work Plan", "create"):
        frappe.throw("Không có quyền tạo kế hoạch công việc.", frappe.PermissionError)
    if not farm_block:
        frappe.throw("Phải chọn một lô lớn để tạo kế hoạch nháp.")
    date = str(as_date(posting_date))
    lock("Farm", farm)
    proposal = _care_plan_preview(farm, date, date, farm_block)
    if proposal["fingerprint"] != fingerprint:
        frappe.throw("Lịch dự kiến hoặc nguồn vừa thay đổi, cần tải lại và xác nhận.")
    if not proposal["rows"]:
        frappe.throw("Không có lịch việc ngày này cho lô đã chọn.")
    missing = [r["work_code"] for r in proposal["rows"] if r["blocking_issue"]]
    if missing:
        frappe.throw("Chưa đủ danh mục/năng suất/đơn vị để tạo kế hoạch an toàn: " + ", ".join(sorted(set(missing))))
    if len(proposal["rows"]) > 100:
        frappe.throw("Quá nhiều công việc trong một ngày. Kiểm tra dữ liệu chặt mẹ.")
    # Do not duplicate a work type already assigned in any non-cancelled plan.
    assigned = set()
    existing = frappe.get_list("Daily Work Plan",
        filters={"farm":farm, "posting_date":date,"docstatus":["<",2]},
        fields=["name"], limit_page_length=0)
    for p in existing:
        doc = frappe.get_doc("Daily Work Plan", p.name)
        doc.check_permission("read")
        assigned.update((line.farm_block, line.work_type)
                        for line in doc.lines)
    lines = []
    for row in proposal["rows"]:
        key = (row["farm_block"], row["work_type"])
        if key in assigned:
            frappe.throw("Đã có kế hoạch cùng lô và công việc trong ngày. Không tạo trùng: " + row["work_code"])
        lines.append(dict(farm_block=row["farm_block"], work_type=row["work_type"],
                          planned_qty=row["planned_qty"], payment_mode="Daywork",
                          schedule_key=row["schedule_key"],
                          lotus_forecast_sources=json.dumps({
                              "profile":proposal["profile"], "farm_block":row["farm_block"],
                              "crop_cycle":row["crop_cycle"], "plan_cohort":row["plan_cohort"],
                              "cut_entries":row["sources"]}, ensure_ascii=False),
                          execution_state="Planned"))
    company = frappe.db.get_value("Farm", farm, "company")
    doc = frappe.get_doc(dict(doctype="Daily Work Plan",
                              company=company, farm=farm, posting_date=date,
                              lines=lines))
    doc.insert()
    # Explicit DRAFT, submission stays a separate approved action.
    return dict(plan=doc.name, docstatus=doc.docstatus, lines=len(lines), farm_block=farm_block, posting_date=date)


# Parameters ------------------------------------------------------------------------------------------

def _new_profile(farm, values, source='Web', source_import=None, notes=None):
    version = (frappe.get_all(PROFILE, filters={'farm': farm}, pluck='version', order_by='version desc', limit_page_length=1) or [0])[0] + 1
    doc = frappe.get_doc(dict(doctype=PROFILE, farm=farm, version=version, approval_state='Approved', source=source, source_import=source_import,
                              approved_by=frappe.session.user, approved_at=now_datetime(), notes=notes, **values))
    doc.flags.ignore_permissions = True
    doc.insert()
    return doc


@frappe.whitelist(methods=['POST'])
def save_profile(farm, values, notes=None):
    """values: the parameters -> a new approved version (the previous versions stay for the record)."""
    _editor(farm)
    values = json.loads(values) if isinstance(values, str) else dict(values)
    active = _active(_profiles(farm))
    merged = {**(_profile_values(active) if active else rules.DEFAULTS), **{k: v for k, v in values.items() if k in PARAM_FIELDS}}
    clean, errors = rules.validate_profile(merged)
    if errors:
        frappe.throw('<br>'.join(errors))
    if active and all(abs(float(active.get(f) or 0) - float(clean[f])) < 1e-9 for f in PARAM_FIELDS):
        frappe.throw('Thông số không thay đổi so với phiên bản đang dùng.')
    return dict(version=_new_profile(farm, clean, notes=notes).version)


# Export ----------------------------------------------------------------------------------------------

@frappe.whitelist()
def export_forecast(farm, zone=None):
    """The forecast as .xlsx: months, zones, zone x month, every wave of every journal row, and the parameters used."""
    data = board(farm, zone)
    if not data.get('profile'):
        frappe.throw('Chưa có bộ thông số đã duyệt.')
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    stage_cols = [('shoot', 'Trổ buồng (cây)'), ('inject', 'Chích bắp (buồng)'), ('budcut', 'Cắt bắp (buồng)'), ('harvest', 'Thu hoạch (buồng)'),
                  ('pack', 'Đóng gói (buồng)'), ('tonnes', 'Tấn sơ chế')]

    def sheet(title, header, rows, widths):
        ws = book.active if book.active.title == 'Sheet' else book.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = 'B2'
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w

    month_label = lambda m: m[5:7] + '/' + m[0:4]
    sheet('Theo tháng', ['Tháng', *[l for _, l in stage_cols]], [[month_label(m['month']), *[m[k] for k, _ in stage_cols]] for m in data['months']], [10] + [16] * 6)
    sheet('Theo khu', ['Khu', 'Cây mẹ chặt lần 2', *[l for _, l in stage_cols]], [[z['zone'], z['trees'], *[z[k] for k, _ in stage_cols]] for z in data['zones']], [8] + [16] * 7)
    months = [m['month'] for m in data['months']]
    sheet('Khu x tháng (tấn)', ['Khu', *[month_label(m) for m in months]], [[z, *[data['zone_months'].get(z, {}).get(m, 0) for m in months]] for z in sorted(data['zone_months'])], [8] + [10] * len(months))
    vn = lambda d: d[8:10] + '/' + d[5:7] + '/' + d[0:4] if d else ''
    sheet('Chi tiết', ['Khu', 'Ngày chặt mẹ 2', 'Cây mẹ', 'Đợt trổ', 'Tỷ lệ', 'Số cây trổ', 'Trổ từ', 'Trổ đến', 'Ngày đại diện trổ', 'Ngày chích bắp', 'Ngày cắt bắp',
                       'Buồng cắt bắp', 'Ngày thu', 'Buồng thu', 'Buồng đóng gói', 'Kg sơ chế', 'Dòng nguồn'],
          [[r['zone'], vn(r['cut_date']), r['tree_qty'], r['wave'], r['share'], r['shoot_qty'], vn(r['shoot_from']), vn(r['shoot_to']), vn(r['shoot_date']),
            vn(r['inject_date']), vn(r['budcut_date']), r['budcut_qty'], vn(r['harvest_date']), r['harvest_qty'], r['pack_qty'], r['kg'], r['source_ref']] for r in data['rows']],
          [6, 13, 9, 8, 7, 10, 11, 11, 13, 13, 13, 12, 11, 10, 12, 10, 22])
    p = data['profile']
    sheet('Thông số', ['Thông số', 'Giá trị'], [['Phiên bản', p['version']], *[[label, p[f]] for f, label, _, _ in rules.PARAMS]], [44, 12])
    out = BytesIO()
    book.save(out)
    frappe.local.response.filename = 'du-thu-chat-me-2-' + str(farm) + ('-khu-' + zone if zone else '') + '.xlsx'
    frappe.local.response.filecontent = out.getvalue()
    frappe.local.response.type = 'download'
