"""Workbook profiling with explicit headers; formulas are evidence, never executed."""
from io import BytesIO
import json
from zipfile import ZipFile
from lotus_farm.source_data import source_text, source_number, source_date, staged_row
from lotus_farm.rules import RuleError, canonical_hash

MAX_BYTES=40*1024*1024
MAX_ROWS=10000


def read_workbook(data, mapping):
    from openpyxl import load_workbook
    if not isinstance(data,bytes) or len(data)>MAX_BYTES:raise RuleError('Sổ Excel vượt giới hạn 40 MiB.')
    with ZipFile(BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist())>200*1024*1024:raise RuleError('Sổ Excel sau khi giải nén vượt 200 MiB.')
    book=load_workbook(BytesIO(data),data_only=False,read_only=False,keep_links=False)
    try:
        if mapping['sheet_name'] not in book.sheetnames:raise RuleError('Trang tính đã duyệt không tồn tại.')
        sheet=book[mapping['sheet_name']]
        if sheet.max_row>MAX_ROWS:raise RuleError('Tách sổ Excel thành các đợt tối đa 10.000 dòng.')
        headers=[source_text(c.value) for c in sheet[int(mapping['header_row'])]]
        used=[h for h in headers if h]
        if len(used)!=len(set(used)):raise RuleError('Dòng tiêu đề có tên cột trùng; chọn dòng tiêu đề rõ ràng.')
        columns={h:i for i,h in enumerate(headers) if h}
        required=[mapping['legacy_org_column']]+[f['source_column'] for f in mapping['fields']]
        if any(c not in columns for c in required):raise RuleError('Dòng tiêu đề đã chọn thiếu các cột của ánh xạ đã duyệt.')
        keys=json.loads(mapping.get('business_key_columns') or '[]')
        samples=json.loads(mapping.get('sample_columns_json') or '[]')
        if any(c not in columns for c in keys) or any(s['column'] not in columns for s in samples):
            raise RuleError('Thiếu cột khóa nghiệp vụ hoặc cột mẫu đo.')
        rows=[];seen_keys=set()
        for number,cells in enumerate(sheet.iter_rows(min_row=int(mapping['first_data_row'])),int(mapping['first_data_row'])):
            if all(c.value is None for c in cells):continue
            raw={h:cells[i].value for h,i in columns.items()}
            legacy=source_text(raw.get(mapping['legacy_org_column'])) or '__MISSING_ORG__'
            row=staged_row(mapping['source_key'],sheet.title,number,raw,legacy)
            errors=[];normalized={}
            if legacy=='__MISSING_ORG__':errors.append('Missing organization mapping key.')
            for f in mapping['fields']:
                value=raw.get(f['source_column'])
                try:
                    if isinstance(value,str) and value.startswith('='):raise RuleError('Ô công thức cần phiên bản nguồn chỉ chứa giá trị đã kiểm tra.')
                    kind=f['value_type']
                    if kind=='Date':value=source_date(value)
                    elif kind in ('Number','Integer','Percent'):
                        value=source_number(value,mapping['decimal_mark'],mapping.get('grouping_mark') or '',percent=kind=='Percent')
                        if kind=='Percent' and isinstance(raw.get(f['source_column']),str) and raw[f['source_column']].strip().endswith('%'):value*=100
                        if kind=='Integer' and value is not None:
                            if not float(value).is_integer():raise RuleError('Cần một số nguyên.')
                            value=int(value)
                    elif kind=='JSON':value=json.loads(value) if value not in (None,'') else None
                    else:value=source_text(value)
                    if f.get('value_map_json') and value is not None:
                        lookup=json.loads(f['value_map_json'])
                        if str(value) not in lookup:raise RuleError('Giá trị chưa có ánh xạ danh mục đã duyệt.')
                        value=lookup[str(value)]
                    if f.get('required') and value in (None,''):raise RuleError('Giá trị bắt buộc trong nguồn đang để trống.')
                    normalized[f['target_field']]=value
                except (RuleError,ValueError,TypeError) as exc:errors.append(f['source_column']+': '+str(exc))
            if samples:
                normalized['samples']=[]
                for sample in samples:
                    try:
                        value=raw.get(sample['column'])
                        if value in (None,'','-'):continue
                        if isinstance(value,str) and value.startswith('='):raise RuleError('Công thức mẫu cần giá trị đã kiểm tra.')
                        value=source_number(value,mapping['decimal_mark'],mapping.get('grouping_mark') or '')
                        normalized['samples'].append(dict(sample_no=sample.get('sample_no') or sample['column'],lane=sample.get('lane') or sample['column'],dimension=sample.get('dimension') or 'Length',value=value,uom=sample['uom']))
                    except (RuleError,ValueError,TypeError) as exc:errors.append(sample['column']+': '+str(exc))
            if keys:
                key_values=[source_text(raw.get(c)) for c in keys]
                if any(v in (None,'') for v in key_values):errors.append('Business identity is blank.')
                row['business_key']=canonical_hash([legacy,key_values])
                if row['business_key'] in seen_keys:errors.append('Duplicate business identity within this source.')
                seen_keys.add(row['business_key'])
                row['row_key']=canonical_hash([mapping.get('source_family') or mapping['source_key'],sheet.title,row['business_key']])
            row['normalized_hash']=canonical_hash(normalized)
            row['normalized_json']=json.dumps(normalized,ensure_ascii=False,default=str,sort_keys=True)
            row['validation_error']='; '.join(errors)
            rows.append(row)
        totals={f['target_field']:sum(float(json.loads(r['normalized_json']).get(f['target_field']) or 0) for r in rows if not r['validation_error']) for f in mapping['fields'] if f['value_type'] in ('Number','Integer')}
        return dict(rows=rows,control_totals=totals,profile=dict(sheets=book.sheetnames,selected_sheet=sheet.title,rows=sheet.max_row,columns=sheet.max_column,
            merged_headers=[str(r) for r in sheet.merged_cells.ranges],formula_cells=[c.coordinate for row in sheet for c in row if c.data_type=='f'][:1000],
            external_links=[n for n in ZipFile(BytesIO(data)).namelist() if n.startswith('xl/externalLinks/')],named_ranges=list(book.defined_names)))
    finally:book.close()
