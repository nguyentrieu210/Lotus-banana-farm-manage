"""Tài chính – Kế toán: settings, chart of accounts, periods and opening balances (chief accountant)."""
import base64
import json
from io import BytesIO

import frappe
from frappe.utils import cint, flt, getdate

from lotus_farm.finance import posting, rules, setup


@frappe.whitelist()
def settings():
    setup.require_view()
    cfg = setup.settings()
    company = setup.company()
    accounts = frappe.get_all('Account', filters={'company': company, 'is_group': 0, 'disabled': 0},
                              fields=['name', 'account_number', 'account_name', 'root_type', 'account_type'], order_by='account_number, name')
    return dict(enabled=bool(cfg.finance_enabled), company=cfg.company or company, cutover_date=str(cfg.cutover_date or ''),
                coa_version=cfg.coa_version or '', coa_applied_at=str(cfg.coa_applied_at or ''), coa_applied_by=cfg.coa_applied_by or '',
                mapping=setup.mapping(cfg), keys=[dict(key=k, label=v[0], roots=v[1], account_type=v[2]) for k, v in rules.KEYS.items()],
                errors=setup.mapping_errors(cfg), accounts=accounts, periods=posting.periods(company), is_chief=setup.is_chief())


@frappe.whitelist(methods=['POST'])
def save_settings(values):
    """values: {enabled, cutover_date, mapping: {key: account}} -> saved (validation explains what is missing)."""
    setup.require_chief()
    values = frappe.parse_json(values) if isinstance(values, str) else values
    doc = setup.settings()
    doc.company = doc.company or setup.company()
    if 'cutover_date' in values:
        doc.cutover_date = values.get('cutover_date') or None
    for key, account in (values.get('mapping') or {}).items():
        if key in rules.KEYS:
            doc.set('acct_' + key, account or None)
    if 'enabled' in values:
        doc.finance_enabled = 1 if values.get('enabled') else 0
    doc.flags.ignore_permissions = True
    doc.save()
    return settings()


@frappe.whitelist()
def coa_preview(on_date=None, warehouse_accounts=None):
    setup.require_view()
    return setup.coa_preview(on_date, frappe.parse_json(warehouse_accounts) if isinstance(warehouse_accounts, str) else warehouse_accounts)


@frappe.whitelist(methods=['POST'])
def coa_apply(on_date=None, warehouse_accounts=None):
    return setup.coa_apply(on_date, frappe.parse_json(warehouse_accounts) if isinstance(warehouse_accounts, str) else warehouse_accounts)


# Opening balances ------------------------------------------------------------------------------------

OPENING_COLUMNS = ('Số tài khoản', 'Mã nhà cung cấp (TK 331)', 'Dư Nợ', 'Dư Có', 'Mã lô (TK 154)', 'Ghi chú')


@frappe.whitelist()
def opening_template():
    setup.require_view()
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    ws = book.active
    ws.title = 'Số dư đầu kỳ'
    ws.append(list(OPENING_COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(['1111', '', 50000000, '', '', 'Tiền mặt tồn quỹ'])
    ws.append(['331', 'NCC-0001', '', 12000000, '', 'Nợ nhà cung cấp'])
    ws.append(['4111', '', '', 38000000, '', 'Vốn góp'])
    out = BytesIO()
    book.save(out)
    frappe.local.response.filename = 'so-du-dau-ky.xlsx'
    frappe.local.response.filecontent = out.getvalue()
    frappe.local.response.type = 'download'


def _opening_rows(content, company):
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(content), read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows or str(rows[0][0] or '').strip() != OPENING_COLUMNS[0]:
        frappe.throw('Dùng file mẫu số dư đầu kỳ (cột đầu là "Số tài khoản").')
    by_number = {a.account_number: a for a in frappe.get_all('Account', filters={'company': company, 'account_number': ['is', 'set']},
                                                               fields=['name', 'account_number', 'is_group', 'account_type', 'root_type'])}
    out = []
    for line_no, row in enumerate(rows[1:], start=2):
        row = list(row) + [None] * 6
        number, supplier, debit, credit, block, note = (str(row[0] or '').strip(), str(row[1] or '').strip(), flt(row[2]), flt(row[3]), str(row[4] or '').strip(), row[5])
        if not number and not debit and not credit:
            continue
        errors = []
        a = by_number.get(number)
        if not a:
            errors.append('Số tài khoản %s không có' % number)
        elif a.is_group:
            errors.append('%s là tài khoản tổng hợp' % number)
        if bool(round(debit)) == bool(round(credit)):
            errors.append('Ghi một bên Dư Nợ hoặc Dư Có')
        if a and a.account_type == 'Payable' and not frappe.db.exists('Supplier', supplier):
            errors.append('TK %s cần mã nhà cung cấp có trong danh mục' % number)
        if a and a.root_type in ('Income', 'Expense'):
            errors.append('Số dư đầu kỳ chỉ cho tài khoản tài sản, nợ phải trả, vốn')
        dims = {}
        if block:
            b = frappe.db.get_value('Farm Block', {'block_code': block}, ['name', 'farm'], as_dict=True)
            if not b:
                errors.append('Mã lô %s không có' % block)
            else:
                dims = dict(lotus_farm=b.farm, lotus_farm_block=b.name)
        out.append(dict(line=line_no, number=number, account=a.name if a else None, supplier=supplier or None, debit=round(debit), credit=round(credit),
                        dims=dims, note=note, errors=errors))
    return out


@frappe.whitelist(methods=['POST'])
def opening_preview(content):
    setup.require_chief()
    cfg = setup.settings()
    if not cfg.cutover_date:
        frappe.throw('Chọn ngày bắt đầu ghi sổ trước khi nhập số dư đầu kỳ.')
    rows = _opening_rows(base64.b64decode(content.split(',', 1)[-1]), setup.company())
    good = [r for r in rows if not r['errors']]
    debit, credit = sum(r['debit'] for r in good), sum(r['credit'] for r in good)
    token = frappe.generate_hash(length=20)
    frappe.cache.set_value('lotus_opening:' + frappe.session.user + ':' + token, json.dumps(good, default=str), expires_in_sec=1800)
    return dict(token=token, rows=rows, debit=debit, credit=credit, balanced=debit == credit and debit > 0, errors=sum(1 for r in rows if r['errors']),
                on_date=str(cfg.cutover_date), existing=bool(posting.active('Lotus Finance Settings', setup.SETTINGS, 'Opening Balance')))


@frappe.whitelist(methods=['POST'])
def opening_commit(token):
    """One opening Journal Entry dated the cutover day (Nợ = Có required)."""
    setup.require_chief()
    raw = frappe.cache.get_value('lotus_opening:' + frappe.session.user + ':' + str(token))
    if not raw:
        frappe.throw('Bản xem trước đã hết hạn. Tải file lên lại.')
    rows = json.loads(raw)
    lines = [rules.line(r['account'], debit=r['debit'], credit=r['credit'], dims=r['dims'], note=r.get('note') or 'Số dư đầu kỳ') for r in rows]
    errors = rules.check_entry(lines)
    if errors:
        frappe.throw('<br>'.join(errors))
    if posting.active('Lotus Finance Settings', setup.SETTINGS, 'Opening Balance'):
        frappe.throw('Đã có bút toán số dư đầu kỳ. Đảo bút toán đó trước nếu cần nhập lại.')
    cfg = setup.settings()
    company = setup.company()
    posting.check_open(company, cfg.cutover_date)
    cost_center = frappe.get_cached_value('Company', company, 'cost_center')
    accounts = []
    for r, ln in zip(rows, lines):
        row = dict(account=ln['account'], debit_in_account_currency=ln['debit'], credit_in_account_currency=ln['credit'], cost_center=cost_center,
                   user_remark=ln['note'] or None, **r['dims'])
        if r.get('supplier'):
            row.update(party_type='Supplier', party=r['supplier'])
        accounts.append(row)
    je = frappe.get_doc(dict(doctype='Journal Entry', voucher_type='Opening Entry', is_opening='Yes', company=company, posting_date=cfg.cutover_date,
                             user_remark='Số dư đầu kỳ tại ngày %s' % getdate(cfg.cutover_date).strftime('%d/%m/%Y'), accounts=accounts))
    je.flags.ignore_permissions = True
    je.insert()
    je.submit()
    posting.record('Lotus Finance Settings', setup.SETTINGS, 'Opening Balance', je, amount=sum(ln['debit'] for ln in lines), source_hash=str(token))
    frappe.cache.delete_value('lotus_opening:' + frappe.session.user + ':' + str(token))
    return dict(voucher=je.name, amount=sum(ln['debit'] for ln in lines))
