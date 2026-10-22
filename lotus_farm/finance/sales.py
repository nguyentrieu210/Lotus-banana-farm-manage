"""Tài chính – Kế toán: customers, sales invoices, receipts and VAT listings.

A sales invoice ships and bills at once (ERPNext Sales Invoice with update_stock): Nợ 632 / Có 155 at the stock value
(the costing run adjusts 632 by lot at month end) and Nợ 131 / Có 511, 33311 at the selling price, one VAT rate per
invoice (data/vat_vn.json; the farm's own bananas are not subject to VAT). The electronic invoice number is typed in
from the e-invoice provider. A receipt is a Payment Entry (Nợ 111 / 112, Có 131) allocated to open invoices; money
coming in needs no signatures. The VAT listings and the monthly summary are read from the submitted invoices and
checked against 1331 / 33311 in the ledger.
"""
from io import BytesIO

import frappe
from frappe.utils import add_days, flt, get_first_day, get_last_day, getdate, nowdate

from lotus_farm.finance import payments, posting, rules, setup

SI, PE, PI = 'Sales Invoice', 'Payment Entry', 'Purchase Invoice'


def _json(value):
    return frappe.parse_json(value) if isinstance(value, str) else (value or [])


def _leaf(doctype, root_name):
    leaf = frappe.get_all(doctype, filters={'is_group': 0}, pluck='name', limit_page_length=1)
    if leaf:
        return leaf[0]
    root = frappe.get_all(doctype, filters={'is_group': 1}, pluck='name', order_by='lft', limit_page_length=1)
    field = {'Customer Group': 'customer_group_name', 'Territory': 'territory_name'}[doctype]
    parent = {'Customer Group': 'parent_customer_group', 'Territory': 'parent_territory'}[doctype]
    doc = frappe.get_doc({'doctype': doctype, field: root_name, parent: root[0] if root else None, 'is_group': 0})
    doc.flags.ignore_permissions = True
    return doc.insert().name


# Customers --------------------------------------------------------------------------------------------

@frappe.whitelist(methods=['POST'])
def save_customer(values):
    """values: {name?, customer_name, tax_id, address, credit_days}"""
    setup.require_post()
    values = _json(values)
    title = str(values.get('customer_name') or '').strip()
    if len(title) < 2:
        frappe.throw('Nhập tên khách hàng.')
    tax_id = str(values.get('tax_id') or '').strip()
    if tax_id and not tax_id.replace('-', '').isdigit():
        frappe.throw('Mã số thuế chỉ gồm chữ số (và dấu - cho chi nhánh).')
    if tax_id and frappe.db.exists('Customer', {'tax_id': tax_id, 'name': ['!=', values.get('name') or '']}):
        frappe.throw('Mã số thuế %s đã có khách hàng khác.' % tax_id)
    doc = frappe.get_doc('Customer', values['name']) if values.get('name') else frappe.new_doc('Customer')
    doc.customer_name = title
    doc.customer_type = 'Company' if tax_id else 'Individual'
    doc.tax_id = tax_id or None
    doc.customer_group = doc.customer_group or _leaf('Customer Group', 'Khách hàng')
    doc.territory = doc.territory or _leaf('Territory', 'Việt Nam')
    if doc.meta.has_field('customer_details'):
        doc.customer_details = str(values.get('address') or '')[:1000] or None
    if doc.meta.has_field('lotus_credit_days'):
        doc.lotus_credit_days = max(0, int(flt(values.get('credit_days'))))
    doc.flags.ignore_permissions = True
    doc.save()
    return dict(name=doc.name, customer_name=doc.customer_name)


@frappe.whitelist()
def receivables():
    """Per customer: invoiced, received, still owed, overdue; open invoices with aging; 131 in the ledger."""
    setup.require_view()
    company = setup.company()
    m = setup.mapping()
    today = getdate(nowdate())
    meta = frappe.get_meta('Customer')
    customers = {c.name: c for c in frappe.get_all('Customer', filters={'disabled': 0},
                                                   fields=['name', 'customer_name', 'tax_id'] + (['customer_details'] if meta.has_field('customer_details') else [])
                                                   + (['lotus_credit_days'] if meta.has_field('lotus_credit_days') else []), limit_page_length=0)}
    invoices = frappe.get_all(SI, filters={'docstatus': 1, 'company': company, 'is_return': 0},
                              fields=['name', 'customer', 'customer_name', 'posting_date', 'due_date', 'grand_total', 'outstanding_amount', 'net_total',
                                      'total_taxes_and_charges', 'lotus_einvoice_no', 'lotus_einvoice_series', 'lotus_vat_code'],
                              order_by='posting_date desc', limit_page_length=2000) if frappe.get_meta(SI).has_field('lotus_vat_code') else []
    summary = {}
    for inv in invoices:
        s = summary.setdefault(inv.customer, dict(customer=inv.customer, customer_name=inv.customer_name, invoiced=0, owed=0, overdue=0, open_invoices=0))
        s['invoiced'] += flt(inv.grand_total)
        s['owed'] += flt(inv.outstanding_amount)
        if flt(inv.outstanding_amount) > 0.5:
            s['open_invoices'] += 1
            if inv.due_date and getdate(inv.due_date) < today:
                s['overdue'] += flt(inv.outstanding_amount)
    rows = []
    for name, c in customers.items():
        s = summary.get(name) or dict(customer=name, customer_name=c.customer_name, invoiced=0, owed=0, overdue=0, open_invoices=0)
        rows.append(dict(s, received=s['invoiced'] - s['owed'], tax_id=c.tax_id or '', address=c.get('customer_details') or '',
                         credit_days=c.get('lotus_credit_days') or 0))
    rows.sort(key=lambda r: (-r['owed'], r['customer_name'] or ''))
    open_rows = []
    for inv in invoices:
        if flt(inv.outstanding_amount) <= 0.5:
            continue
        open_rows.append(dict(inv, posting_date=str(inv.posting_date), due_date=str(inv.due_date or ''),
                              aging=rules.aging(getdate(inv.due_date) if inv.due_date else None, today)))
    gl = -flt(frappe.db.sql("""select sum(credit-debit) from `tabGL Entry` where company=%s and account=%s and is_cancelled=0""",
                            (company, m['receivable_customer']))[0][0]) if m.get('receivable_customer') else 0
    return dict(customers=rows, invoices=open_rows, recent=[dict(i, posting_date=str(i.posting_date), due_date=str(i.due_date or '')) for i in invoices[:50]],
                gl_total=gl)


# Sales invoice ----------------------------------------------------------------------------------------

@frappe.whitelist()
def stock_for_sale():
    """Items in finished-goods warehouses (account 155) with quantity, stock value per unit and the farm of the warehouse."""
    setup.require_view()
    m = setup.mapping()
    from lotus_farm.finance.stock import _warehouse_farm
    rows = frappe.db.sql("""select b.item_code, i.item_name, i.stock_uom, i.has_batch_no, b.warehouse, b.actual_qty, b.valuation_rate
        from tabBin b join tabItem i on i.name=b.item_code join tabWarehouse w on w.name=b.warehouse
        where b.actual_qty > 0 and w.company=%s and w.account=%s
        order by i.item_name, b.warehouse""", (setup.company(), m['finished_goods']), as_dict=True)
    for r in rows:
        r['farm'] = _warehouse_farm(r.warehouse)
        if r.has_batch_no:
            r['batches'] = frappe.db.sql("""select b.name, sum(e.qty) qty from `tabSerial and Batch Entry` e join `tabSerial and Batch Bundle` sb on sb.name=e.parent
                join tabBatch b on b.name=e.batch_no where sb.item_code=%s and sb.warehouse=%s and sb.docstatus=1 and sb.is_cancelled=0
                group by b.name having sum(e.qty) > 0 order by b.name""", (r.item_code, r.warehouse), as_dict=True)
    return rows


@frappe.whitelist(methods=['POST'])
def make_sales_invoice(customer, posting_date, lines, vat_code='KCT', due_date=None, einvoice_series=None, einvoice_no=None, remarks=None):
    """lines: [{item_code, warehouse, qty, rate, batch_no?, lotus_farm?, lotus_farm_block?}] -> submitted Sales Invoice that also
    takes the goods out of stock."""
    setup.require_post()
    cfg = setup.require_enabled()
    company = setup.company()
    m = setup.mapping(cfg)
    lines = _json(lines)
    if not frappe.db.exists('Customer', customer):
        frappe.throw('Chọn khách hàng.')
    if not lines:
        frappe.throw('Nhập ít nhất một dòng hàng.')
    einvoice_no = str(einvoice_no or '').strip()
    if einvoice_no and frappe.db.exists(SI, {'lotus_einvoice_no': einvoice_no, 'lotus_einvoice_series': einvoice_series or '', 'docstatus': ['<', 2]}):
        frappe.throw('Số hóa đơn điện tử %s %s đã được dùng.' % (einvoice_series or '', einvoice_no))
    setup.check_cutover(posting_date)
    posting.check_open(company, posting_date, SI)
    credit_days = frappe.db.get_value('Customer', customer, 'lotus_credit_days') if frappe.get_meta('Customer').has_field('lotus_credit_days') else 0
    due = getdate(due_date) if due_date else add_days(getdate(posting_date), int(credit_days or 0))
    if due < getdate(posting_date):
        frappe.throw('Hạn thanh toán không được trước ngày hóa đơn.')
    from lotus_farm.finance.stock import _warehouse_farm
    cost_center = frappe.get_cached_value('Company', company, 'cost_center')
    items = []
    for i, ln in enumerate(lines, start=1):
        qty, rate = flt(ln.get('qty')), flt(ln.get('rate'))
        if qty <= 0 or rate < 0:
            frappe.throw('Dòng %d: số lượng phải lớn hơn 0, đơn giá không âm.' % i)
        if not ln.get('item_code') or not ln.get('warehouse'):
            frappe.throw('Dòng %d: chọn hàng và kho xuất.' % i)
        dims = payments._dims(dict(ln, lotus_farm=ln.get('lotus_farm') or _warehouse_farm(ln['warehouse'])))
        row = dict(item_code=ln['item_code'], qty=qty, rate=rate, warehouse=ln['warehouse'], income_account=m['sales_revenue'],
                   expense_account=m['cost_of_goods_sold'], cost_center=cost_center, **dims)
        if ln.get('batch_no'):
            row.update(batch_no=ln['batch_no'], use_serial_batch_fields=1)
        items.append(row)
    si = frappe.get_doc(dict(doctype=SI, company=company, customer=customer, posting_date=posting_date, set_posting_time=1, due_date=due,
                             debit_to=m['receivable_customer'], update_stock=1, items=items, remarks=remarks or None,
                             lotus_einvoice_series=einvoice_series or None, lotus_einvoice_no=einvoice_no or None))
    payments.apply_vat(si, vat_code, posting_date, m['output_vat'])
    si.set('payment_schedule', [])
    si.flags.ignore_permissions = True
    si.insert()
    si.submit()
    posting.record(SI, si.name, 'Sales Invoice', si, amount=si.grand_total, source_hash=einvoice_no or si.name)
    return dict(invoice=si.name, amount=si.grand_total, tax=si.total_taxes_and_charges, cost=_cost(si.name, m['cost_of_goods_sold']))


def _cost(voucher, account):
    return flt(frappe.db.sql("""select sum(debit-credit) from `tabGL Entry` where voucher_no=%s and account=%s and is_cancelled=0""", (voucher, account))[0][0])


@frappe.whitelist(methods=['POST'])
def receive_payment(customer, allocations, mode, posting_date, reference_no=None, remark=None):
    """allocations: [{invoice, amount}] -> submitted Payment Entry (Nợ tiền / Có 131)."""
    setup.require_post()
    setup.require_enabled()
    company = setup.company()
    m = setup.mapping()
    allocations = [a for a in _json(allocations) if flt(a.get('amount')) > 0]
    if not allocations:
        frappe.throw('Chọn hóa đơn và số tiền thu.')
    if mode == 'bank' and not reference_no:
        frappe.throw('Tiền về tài khoản cần số giao dịch / số báo có.')
    setup.check_cutover(posting_date)
    posting.check_open(company, posting_date, PE)
    refs, total = [], 0
    for a in allocations:
        inv = frappe.db.get_value(SI, a['invoice'], ['customer', 'docstatus', 'outstanding_amount', 'grand_total', 'due_date'], as_dict=True)
        if not inv or inv.docstatus != 1 or inv.customer != customer:
            frappe.throw('Hóa đơn %s không thuộc khách hàng này hoặc chưa ghi sổ.' % a['invoice'])
        amount = flt(a['amount'])
        if amount > flt(inv.outstanding_amount) + 0.5:
            frappe.throw('Hóa đơn %s chỉ còn phải thu %s.' % (a['invoice'], format(round(flt(inv.outstanding_amount)), ',').replace(',', '.')))
        refs.append(dict(reference_doctype=SI, reference_name=a['invoice'], total_amount=inv.grand_total, outstanding_amount=inv.outstanding_amount,
                         allocated_amount=amount, due_date=inv.due_date))
        total += amount
    money = m['bank'] if mode == 'bank' else m['cash']
    currency = frappe.get_cached_value('Company', company, 'default_currency')
    pe = frappe.get_doc(dict(doctype=PE, payment_type='Receive', company=company, posting_date=posting_date, party_type='Customer', party=customer,
                             paid_from=m['receivable_customer'], paid_to=money, paid_from_account_currency=currency, paid_to_account_currency=currency,
                             paid_amount=total, received_amount=total, source_exchange_rate=1, target_exchange_rate=1,
                             reference_no=reference_no or None, reference_date=posting_date if reference_no else None,
                             remarks=remark or 'Thu tiền khách hàng', references=refs))
    pe.flags.ignore_permissions = True
    pe.insert()
    pe.submit()
    posting.record(PE, pe.name, 'Customer Receipt', pe, amount=total, source_hash=','.join(a['invoice'] for a in allocations))
    return dict(receipt=pe.name, amount=total)


# VAT --------------------------------------------------------------------------------------------------

@frappe.whitelist()
def vat_options(on_date=None):
    setup.require_view()
    data = rules.load_vat()
    return dict(note=data.get('note'), rates=rules.vat_rates(on_date or nowdate(), data))


def _month(month):
    start = get_first_day(str(month)[:7] + '-01')
    return start, get_last_day(start)


def vat_listing(month):
    """Purchase and sales invoices of the month with net, rate code and VAT (bảng kê hóa đơn mua vào / bán ra)."""
    company = setup.company()
    start, end = _month(month)
    has = frappe.get_meta(PI).has_field('lotus_vat_code')
    purchases = frappe.db.sql("""select p.name, p.bill_no, p.bill_date, p.posting_date, p.supplier, s.supplier_name, s.tax_id, p.net_total net,
            p.total_taxes_and_charges tax, %s code from `tabPurchase Invoice` p left join tabSupplier s on s.name=p.supplier
        where p.company=%%s and p.docstatus=1 and p.posting_date between %%s and %%s order by p.posting_date, p.name""" % ("p.lotus_vat_code" if has else "'KCT'"),
                              (company, start, end), as_dict=True)
    has = frappe.get_meta(SI).has_field('lotus_vat_code')
    sales = frappe.db.sql("""select i.name, %s, i.posting_date, i.customer, c.customer_name, c.tax_id, i.net_total net, i.total_taxes_and_charges tax, %s code
        from `tabSales Invoice` i left join tabCustomer c on c.name=i.customer
        where i.company=%%s and i.docstatus=1 and i.posting_date between %%s and %%s order by i.posting_date, i.name""" % (
        "i.lotus_einvoice_series series, i.lotus_einvoice_no einvoice_no" if has else "'' series, '' einvoice_no", "i.lotus_vat_code" if has else "'KCT'"),
                          (company, start, end), as_dict=True)
    for r in purchases + sales:
        for f in ('bill_date', 'posting_date'):
            if r.get(f):
                r[f] = str(r[f])
        r['code'] = r.get('code') or 'KCT'
    return purchases, sales, start, end


@frappe.whitelist()
def vat_report(month):
    setup.require_view()
    m = setup.mapping()
    company = setup.company()
    purchases, sales, start, end = vat_listing(month)
    summary = rules.vat_summary(purchases, sales)

    def moved(account, side):
        if not account:
            return 0
        return flt(frappe.db.sql("""select sum(%s) from `tabGL Entry` where company=%%s and account=%%s and is_cancelled=0 and posting_date between %%s and %%s
            and voucher_type in ('Purchase Invoice', 'Sales Invoice')""" % side, (company, account, start, end))[0][0])
    checks = [dict(label='Thuế GTGT đầu vào: bảng kê so với TK 1331', listing=summary['input_tax'], ledger=moved(m.get('input_vat'), 'debit-credit')),
              dict(label='Thuế GTGT đầu ra: bảng kê so với TK 33311', listing=summary['output_tax'], ledger=moved(m.get('output_vat'), 'credit-debit'))]
    for c in checks:
        c['ok'] = abs(round(c['listing']) - round(c['ledger'])) < 1
    return dict(month=str(month)[:7], from_date=str(start), to_date=str(end), purchases=purchases, sales=sales, summary=summary, checks=checks)


@frappe.whitelist()
def export_vat(month):
    setup.require_view()
    data = vat_report(month)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()

    def sheet(ws, title, header, rows):
        ws.title = title
        ws.append([title + ' tháng ' + data['month'][5:] + '/' + data['month'][:4]])
        ws.append(header)
        for cell in ws[2]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(r)
    sheet(book.active, 'Bảng kê mua vào', ['STT', 'Số hóa đơn', 'Ngày hóa đơn', 'Người bán', 'MST người bán', 'Giá trị chưa thuế', 'Thuế suất', 'Thuế GTGT'],
          [[i, p['bill_no'], p['bill_date'], p['supplier_name'], p['tax_id'], p['net'], p['code'], p['tax']] for i, p in enumerate(data['purchases'], start=1)])
    sheet(book.create_sheet(), 'Bảng kê bán ra', ['STT', 'Ký hiệu', 'Số hóa đơn', 'Ngày', 'Người mua', 'MST người mua', 'Doanh thu chưa thuế', 'Thuế suất', 'Thuế GTGT'],
          [[i, s['series'], s['einvoice_no'] or s['name'], s['posting_date'], s['customer_name'], s['tax_id'], s['net'], s['code'], s['tax']]
           for i, s in enumerate(data['sales'], start=1)])
    t = data['summary']
    sheet(book.create_sheet(), 'Tổng hợp 01-GTGT', ['Chỉ tiêu', 'Giá trị'],
          [['Hàng hóa, dịch vụ mua vào: giá trị', sum(r['net'] for r in t['purchases'])], ['Thuế GTGT đầu vào', t['input_tax']],
           ['Doanh thu không chịu thuế', t['exempt_sales']], ['Doanh thu chịu thuế', t['taxable_sales']], ['Thuế GTGT đầu ra', t['output_tax']],
           ['Thuế GTGT phải nộp (âm = còn được khấu trừ chuyển kỳ sau)', t['payable']],
           ['Số liệu tham khảo để lập tờ khai 01/GTGT trên phần mềm HTKK / eTax; chưa trừ số còn được khấu trừ kỳ trước.', '']])
    output = BytesIO()
    book.save(output)
    frappe.local.response.filename = 'thue-gtgt-' + data['month'] + '.xlsx'
    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.type = 'download'
