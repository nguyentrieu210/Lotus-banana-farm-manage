"""Tài chính – Kế toán: báo cáo tài chính (Thông tư 99) and the books.

B01 (tình hình tài chính) takes every account's balance at a date, receivables / payables split by party so debit
balances are assets and credit balances liabilities; B02 (kết quả kinh doanh) takes the movements of the period
without the result-closing vouchers (911); B03 (lưu chuyển tiền tệ, gián tiếp) starts from the profit before tax and
adds the change of every other balance sheet account, so its closing cash agrees with 111 + 112 + 113 by
construction. Which accounts feed which line is data (data/bctc_tt99.json). The books — nhật ký chung, sổ cái with
contra accounts, sổ quỹ / tiền gửi, công nợ by party, nhập xuất tồn — and the bank statement match read the same ledger.
"""
import base64
from io import BytesIO

import frappe
from frappe.utils import add_days, flt, get_first_day, getdate, nowdate

from lotus_farm.finance import rules, setup


def _excluded(company):
    from lotus_farm.finance.closing import result_vouchers
    return tuple(result_vouchers(company)) or ('',)


def _parts(company, to_date=None, from_date=None, without_result=False):
    """Account (and party) balances: up to to_date, or movements between from_date and to_date."""
    where, args = ['g.company=%s', 'g.is_cancelled=0', "ifnull(a.account_number, '')!=''"], [company]
    if from_date:
        where.append('g.posting_date between %s and %s')
        args += [from_date, to_date]
    else:
        where.append('g.posting_date <= %s')
        args.append(to_date)
    if without_result:
        where.append("g.voucher_no not in %s and g.voucher_type != 'Period Closing Voucher'")
        args.append(_excluded(company))
    rows = frappe.db.sql("""select a.account_number number, a.account_name name, ifnull(g.party, '') party, sum(g.debit) debit, sum(g.credit) credit
        from `tabGL Entry` g join tabAccount a on a.name=g.account where %s group by a.account_number, a.account_name, ifnull(g.party, '')""" % ' and '.join(where),
                         tuple(args), as_dict=True)
    return [dict(number=r.number, name=r.name, party=r.party, debit=flt(r.debit), credit=flt(r.credit), net=flt(r.debit) - flt(r.credit)) for r in rows]


def _rows(lines, *columns):
    return [dict(code=ln['code'], name=ln['name'], level=ln.get('level', 1 if 'sum' in ln else 2), total='sum' in ln, detail=bool(ln.get('detail')),
                 **{k: col.get(ln['code'], 0) for k, col in columns}) for ln in lines]


def _year_start(on_date):
    return getdate(str(getdate(on_date).year) + '-01-01')


def _year_before(day):
    day = getdate(day)
    try:
        return day.replace(year=day.year - 1)
    except ValueError:  # 29/02
        return day.replace(year=day.year - 1, day=28)


@frappe.whitelist()
def b01(on_date=None, compare_date=None):
    setup.require_view()
    company = setup.company()
    data = rules.load_statements()
    on_date = getdate(on_date or nowdate())
    compare_date = getdate(compare_date) if compare_date else add_days(_year_start(on_date), -1)
    closing, opening = _parts(company, on_date), _parts(company, compare_date)
    end = rules.evaluate_statement(data['b01'], lambda t: rules.term_amount(t, closing))
    start = rules.evaluate_statement(data['b01'], lambda t: rules.term_amount(t, opening))
    uncovered, double = rules.coverage(data['b01'], closing)
    checks = [dict(label='Tổng cộng tài sản (280) = Tổng cộng nguồn vốn (440), cuối kỳ', ok=end['280'] == end['440'], difference=end['280'] - end['440']),
              dict(label='Tổng cộng tài sản (280) = Tổng cộng nguồn vốn (440), đầu kỳ', ok=start['280'] == start['440'], difference=start['280'] - start['440'])]
    return dict(form='B01-DN', title='BÁO CÁO TÌNH HÌNH TÀI CHÍNH', on_date=str(on_date), compare_date=str(compare_date), company=company,
                rows=_rows(data['b01'], ('closing', end), ('opening', start)), checks=checks, source=data['source'], notes=data['notes'],
                uncovered=[dict(number=p['number'], name=p['name'], party=p['party'], net=round(p['net'])) for p in uncovered],
                double=[dict(number=p['number'], lines=p['lines'], net=round(p['net'])) for p in double])


def _b02(company, from_date, to_date, data):
    moves = _parts(company, to_date, from_date, without_result=True)
    values = rules.evaluate_statement(data['b02'], lambda t: rules.term_amount(t, moves))
    profit_all = -sum(p['net'] for p in moves if p['number'][:1] in '5678' and not p['number'].startswith('821'))
    return values, moves, round(profit_all)


@frappe.whitelist()
def b02(from_date=None, to_date=None, compare_from=None, compare_to=None):
    setup.require_view()
    company = setup.company()
    data = rules.load_statements()
    to_date = getdate(to_date or nowdate())
    from_date = getdate(from_date or _year_start(to_date))
    now, _, pbt_all = _b02(company, from_date, to_date, data)
    compare_from = getdate(compare_from) if compare_from else _year_before(from_date)
    compare_to = getdate(compare_to) if compare_to else _year_before(to_date)
    before, _, _ = _b02(company, compare_from, compare_to, data)
    checks = [dict(label='Lợi nhuận trước thuế (50) bằng toàn bộ doanh thu trừ chi phí loại 5–8 của kỳ', ok=now['50'] == pbt_all, difference=now['50'] - pbt_all)]
    return dict(form='B02-DN', title='BÁO CÁO KẾT QUẢ HOẠT ĐỘNG KINH DOANH', from_date=str(from_date), to_date=str(to_date), compare_from=str(compare_from),
                compare_to=str(compare_to), company=company, rows=_rows(data['b02'], ('current', now), ('previous', before)), checks=checks, source=data['source'],
                notes=['Không gồm bút toán kết chuyển 911. Chênh lệch ở dòng kiểm tra thường do chi phí 62x chưa kết chuyển sang 154.'])


@frappe.whitelist()
def b03(from_date=None, to_date=None):
    setup.require_view()
    company = setup.company()
    data = rules.load_statements()
    to_date = getdate(to_date or nowdate())
    from_date = getdate(from_date or _year_start(to_date))
    b02_values, moves, pbt_all = _b02(company, from_date, to_date, data)
    opening = _parts(company, add_days(from_date, -1))
    lines = data['b03']
    cash = data['cash']
    balance_sheet = [p for p in moves if p['number'][:1] in '12349' and not any(p['number'].startswith(c) for c in cash)]
    uncovered, double = rules.coverage([ln for ln in lines if ln['code'] not in ('60',)], balance_sheet)
    fallback = next(ln['code'] for ln in lines if ln.get('fallback'))
    adjust = next(ln['code'] for ln in lines if ln.get('adjust'))

    def amount(term):
        if term.get('ref'):
            return b02_values[term['ref'].split(':')[1]]
        return rules.term_amount(term, opening if term.get('at') == 'opening' else moves)
    extra = {fallback: -sum(p['net'] for p in uncovered), adjust: pbt_all - b02_values['50']}
    values = rules.evaluate_statement([dict(ln, terms=(ln.get('terms') or []) + ([{'value': extra[ln['code']]}] if ln['code'] in extra else []))
                                       for ln in lines], lambda t: t['value'] if 'value' in t else amount(t))
    cash_end = round(sum(p['net'] for p in _parts(company, to_date) if any(p['number'].startswith(c) for c in cash)))
    checks = [dict(label='Tiền cuối kỳ (70) bằng số dư 111 + 112 + 113', ok=values['70'] == cash_end, difference=values['70'] - cash_end)]
    return dict(form='B03-DN', title='BÁO CÁO LƯU CHUYỂN TIỀN TỆ (theo phương pháp gián tiếp)', from_date=str(from_date), to_date=str(to_date), company=company,
                rows=_rows(lines, ('current', values)), checks=checks, source=data['source'],
                uncovered=[dict(number=p['number'], name=p['name'], net=round(p['net'])) for p in uncovered], double=[dict(number=p['number'], lines=p['lines']) for p in double],
                notes=['Biến động tài khoản không thuộc dòng nào được đưa vào dòng %s và liệt kê bên dưới.' % fallback,
                       'Dòng %s gồm doanh thu, chi phí chưa vào B02 (chi phí 62x chưa kết chuyển).' % adjust])


# Books ------------------------------------------------------------------------------------------------

def _accounts(company):
    return {a.name: a for a in frappe.get_all('Account', filters={'company': company}, fields=['name', 'account_number', 'account_name'])}


@frappe.whitelist()
def general_journal(from_date=None, to_date=None):
    """Nhật ký chung: every ledger line of the period in date order."""
    setup.require_view()
    company = setup.company()
    to_date = getdate(to_date or nowdate())
    from_date = getdate(from_date or get_first_day(to_date))
    accounts = _accounts(company)
    rows = frappe.db.sql("""select name, posting_date, voucher_type, voucher_no, account, party, debit, credit, remarks from `tabGL Entry`
        where company=%s and is_cancelled=0 and posting_date between %s and %s order by posting_date, creation, name""", (company, from_date, to_date), as_dict=True)
    out = []
    for i, r in enumerate(rows, start=1):
        a = accounts.get(r.account)
        out.append(dict(line=i, posting_date=str(r.posting_date), voucher_type=r.voucher_type, voucher_no=r.voucher_no, number=a.account_number if a else '',
                        account=r.account, party=r.party or '', debit=flt(r.debit), credit=flt(r.credit), remarks=str(r.remarks or '').split('\n')[0][:200]))
    return dict(from_date=str(from_date), to_date=str(to_date), rows=out, debit=sum(r['debit'] for r in out), credit=sum(r['credit'] for r in out))


@frappe.whitelist()
def general_ledger(account, from_date=None, to_date=None, party=None):
    """Sổ cái / sổ chi tiết of one account (optionally one party) with the contra accounts of each line and the running balance.
    Cash (111) reads as sổ quỹ tiền mặt, 112 as sổ tiền gửi ngân hàng."""
    setup.require_view()
    company = setup.company()
    to_date = getdate(to_date or nowdate())
    from_date = getdate(from_date or get_first_day(to_date))
    accounts = _accounts(company)
    if account not in accounts:
        frappe.throw('Tài khoản không thuộc công ty.')
    party_clause, args = ('and party=%s', (party,)) if party else ('', ())
    opening = flt(frappe.db.sql("""select sum(debit-credit) from `tabGL Entry` where company=%%s and account=%%s and is_cancelled=0 and posting_date<%%s %s"""
                                % party_clause, (company, account, from_date) + args)[0][0])
    rows = frappe.db.sql("""select posting_date, voucher_type, voucher_no, party, debit, credit, remarks from `tabGL Entry` where company=%%s and account=%%s and is_cancelled=0
        and posting_date between %%s and %%s %s order by posting_date, creation""" % party_clause, (company, account, from_date, to_date) + args, as_dict=True)
    vouchers = tuple({r.voucher_no for r in rows}) or ('',)
    others = {}
    for o in frappe.db.sql("""select voucher_no, account, sum(debit) debit, sum(credit) credit from `tabGL Entry` where company=%s and is_cancelled=0 and voucher_no in %s
            and account != %s group by voucher_no, account""", (company, vouchers, account), as_dict=True):
        others.setdefault(o.voucher_no, []).append(o)
    balance, out = opening, []
    for r in rows:
        balance += flt(r.debit) - flt(r.credit)
        side = 'credit' if flt(r.debit) else 'debit'
        contra = sorted({accounts[o.account].account_number or o.account for o in others.get(r.voucher_no, []) if flt(o[side]) and o.account in accounts})
        out.append(dict(posting_date=str(r.posting_date), voucher_type=r.voucher_type, voucher_no=r.voucher_no, party=r.party or '', debit=flt(r.debit), credit=flt(r.credit),
                        contra=', '.join(contra), balance=balance, remarks=str(r.remarks or '').split('\n')[0][:200]))
    a = accounts[account]
    return dict(account=account, label=(a.account_number + ' - ' if a.account_number else '') + a.account_name, party=party or '', from_date=str(from_date),
                to_date=str(to_date), opening=opening, closing=balance, debit=sum(r['debit'] for r in out), credit=sum(r['credit'] for r in out), lines=out)


@frappe.whitelist()
def party_summary(account, from_date=None, to_date=None):
    """Bảng tổng hợp chi tiết công nợ (131 / 331): each party's opening, movements and closing balance."""
    setup.require_view()
    company = setup.company()
    to_date = getdate(to_date or nowdate())
    from_date = getdate(from_date or get_first_day(to_date))
    rows = frappe.db.sql("""select party_type, party, sum(case when posting_date < %s then debit-credit else 0 end) opening,
            sum(case when posting_date >= %s then debit else 0 end) debit, sum(case when posting_date >= %s then credit else 0 end) credit
        from `tabGL Entry` where company=%s and account=%s and is_cancelled=0 and posting_date <= %s and ifnull(party, '')!=''
        group by party_type, party""", (from_date, from_date, from_date, company, account, to_date), as_dict=True)
    names = {}
    for kind, field in (('Customer', 'customer_name'), ('Supplier', 'supplier_name'), ('Employee', 'employee_name')):
        parties = [r.party for r in rows if r.party_type == kind]
        if parties:
            names.update({p.name: p.get(field) for p in frappe.get_all(kind, filters={'name': ['in', parties]}, fields=['name', field])})
    out = []
    for r in rows:
        closing = flt(r.opening) + flt(r.debit) - flt(r.credit)
        if not any(abs(v) > 0.5 for v in (flt(r.opening), flt(r.debit), flt(r.credit), closing)):
            continue
        out.append(dict(party=r.party, party_name=names.get(r.party) or r.party, opening_debit=max(flt(r.opening), 0), opening_credit=max(-flt(r.opening), 0),
                        debit=flt(r.debit), credit=flt(r.credit), closing_debit=max(closing, 0), closing_credit=max(-closing, 0)))
    out.sort(key=lambda r: r['party_name'] or '')
    totals = {k: sum(r[k] for r in out) for k in ('opening_debit', 'opening_credit', 'debit', 'credit', 'closing_debit', 'closing_credit')}
    return dict(account=account, from_date=str(from_date), to_date=str(to_date), rows=out, totals=totals)


@frappe.whitelist()
def stock_summary(from_date=None, to_date=None, warehouse=None):
    """Bảng tổng hợp nhập – xuất – tồn by item and warehouse (quantities and values from the stock ledger)."""
    setup.require_view()
    company = setup.company()
    to_date = getdate(to_date or nowdate())
    from_date = getdate(from_date or get_first_day(to_date))
    clause, args = ('and sle.warehouse=%s', (warehouse,)) if warehouse else ('', ())
    rows = frappe.db.sql("""select sle.item_code, i.item_name, i.stock_uom, sle.warehouse,
            sum(case when sle.posting_date < %%s then sle.actual_qty else 0 end) opening_qty,
            sum(case when sle.posting_date < %%s then sle.stock_value_difference else 0 end) opening_value,
            sum(case when sle.posting_date >= %%s and sle.actual_qty > 0 then sle.actual_qty else 0 end) in_qty,
            sum(case when sle.posting_date >= %%s and sle.actual_qty > 0 then sle.stock_value_difference else 0 end) in_value,
            sum(case when sle.posting_date >= %%s and sle.actual_qty < 0 then -sle.actual_qty else 0 end) out_qty,
            sum(case when sle.posting_date >= %%s and sle.actual_qty < 0 then -sle.stock_value_difference else 0 end) out_value
        from `tabStock Ledger Entry` sle join tabItem i on i.name=sle.item_code
        where sle.company=%%s and sle.is_cancelled=0 and sle.posting_date <= %%s %s
        group by sle.item_code, i.item_name, i.stock_uom, sle.warehouse order by i.item_name, sle.warehouse""" % clause,
                         (from_date,) * 6 + (company, to_date) + args, as_dict=True)
    out = []
    for r in rows:
        r = {k: (flt(v) if k.endswith(('_qty', '_value')) else v) for k, v in r.items()}
        r['closing_qty'] = r['opening_qty'] + r['in_qty'] - r['out_qty']
        r['closing_value'] = r['opening_value'] + r['in_value'] - r['out_value']
        if any(abs(r[k]) > 0.0005 for k in ('opening_qty', 'in_qty', 'out_qty', 'closing_qty', 'closing_value')):
            out.append(r)
    totals = {k: sum(r[k] for r in out) for k in ('opening_value', 'in_value', 'out_value', 'closing_value')}
    return dict(from_date=str(from_date), to_date=str(to_date), rows=out, totals=totals)


# Bank statement match -------------------------------------------------------------------------------------

def _statement_rows(content):
    """Excel statement: a header row naming Ngày, Số tiền vào / ra (or one signed Số tiền), Số tham chiếu, Nội dung."""
    from openpyxl import load_workbook
    raw = base64.b64decode(content.split(',', 1)[1] if ',' in str(content)[:100] else content)
    sheet = load_workbook(BytesIO(raw), read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    norm = lambda v: str(v or '').strip().lower()
    head_index = next((i for i, r in enumerate(rows[:20]) if any('ngày' in norm(c) for c in r) and any('tiền' in norm(c) or 'ghi' in norm(c) for c in r)), None)
    if head_index is None:
        frappe.throw('Không tìm thấy dòng tiêu đề (cần cột "Ngày" và "Số tiền" / "Ghi có" / "Ghi nợ").')
    head = [norm(c) for c in rows[head_index]]

    def col(*names):
        return next((i for i, h in enumerate(head) if any(n in h for n in names)), None)
    c_date, c_in, c_out = col('ngày'), col('ghi có', 'tiền vào', 'phát sinh có', 'credit'), col('ghi nợ', 'tiền ra', 'phát sinh nợ', 'debit')
    c_amount, c_ref, c_desc = col('số tiền'), col('tham chiếu', 'số ct', 'mã giao dịch', 'reference'), col('nội dung', 'diễn giải', 'description')
    out = []
    for r in rows[head_index + 1:]:
        if c_date is None or r[c_date] in (None, ''):
            continue
        try:
            when = getdate(r[c_date]) if not isinstance(r[c_date], str) else getdate('-'.join(reversed(r[c_date].strip()[:10].replace('.', '/').split('/'))))
        except Exception:
            continue
        number = lambda v: flt(str(v).replace('.', '').replace(',', '')) if isinstance(v, str) else flt(v)
        amount = number(r[c_in]) - number(r[c_out]) if c_in is not None or c_out is not None else number(r[c_amount]) if c_amount is not None else 0
        if not amount:
            continue
        out.append(dict(date=str(when), amount=round(amount), reference=str(r[c_ref] or '') if c_ref is not None else '',
                        description=str(r[c_desc] or '')[:200] if c_desc is not None else ''))
    return out


@frappe.whitelist(methods=['POST'])
def bank_match(content, account=None, days=3):
    """Upload a bank statement (Excel) and match it to the ledger of the bank account over the statement's dates."""
    setup.require_view()
    company = setup.company()
    account = account or setup.mapping().get('bank')
    statement = _statement_rows(content)
    if not statement:
        frappe.throw('Sao kê không có dòng giao dịch.')
    start, end = min(s['date'] for s in statement), max(s['date'] for s in statement)
    ledger = [dict(date=str(r.posting_date), amount=round(flt(r.amount)), voucher=r.voucher_no, voucher_type=r.voucher_type, reference=r.reference or '', remarks=str(r.remarks or '')[:200])
              for r in frappe.db.sql("""select g.posting_date, g.voucher_no, g.voucher_type, sum(g.debit-g.credit) amount, max(g.remarks) remarks,
                    coalesce(max(pe.reference_no), max(je.cheque_no), '') reference
                from `tabGL Entry` g left join `tabPayment Entry` pe on pe.name=g.voucher_no left join `tabJournal Entry` je on je.name=g.voucher_no
                where g.company=%s and g.account=%s and g.is_cancelled=0 and g.posting_date between %s and %s
                group by g.posting_date, g.voucher_no, g.voucher_type""", (company, account, add_days(start, -int(days)), add_days(end, int(days))), as_dict=True)]
    result = rules.match_bank(statement, ledger, int(days))
    return dict(account=account, from_date=start, to_date=end, statement=len(statement), **result,
                statement_total=sum(s['amount'] for s in statement), ledger_total=sum(l['amount'] for l in ledger if start <= l['date'] <= end))


# Excel ------------------------------------------------------------------------------------------------

@frappe.whitelist()
def export(kind, from_date=None, to_date=None, account=None, party=None):
    setup.require_view()
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    ws = book.active
    company = setup.company()
    bold = Font(bold=True)

    def header(title, form, period):
        ws.append(['Đơn vị báo cáo: ' + company, '', '', 'Mẫu số ' + form if form else ''])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=13)
        ws.append([period])
        ws.append([])

    def signatures():
        ws.append([])
        ws.append(['NGƯỜI LẬP', '', 'KẾ TOÁN TRƯỞNG', '', 'NGƯỜI ĐẠI DIỆN THEO PHÁP LUẬT'])
        ws.append(['(Ký, họ tên)', '', '(Ký, họ tên)', '', '(Ký, họ tên, đóng dấu)'])
    if kind in ('b01', 'b02', 'b03'):
        data = {'b01': lambda: b01(to_date), 'b02': lambda: b02(from_date, to_date), 'b03': lambda: b03(from_date, to_date)}[kind]()
        ws.title = data['form']
        header(data['title'], data['form'] + ' (Thông tư 99/2025/TT-BTC)',
               ('Tại ngày ' + data['on_date']) if kind == 'b01' else 'Kỳ kế toán từ ngày %s đến ngày %s' % (data['from_date'], data['to_date']))
        cols = {'b01': ['TÀI SẢN / NGUỒN VỐN', 'Mã số', 'Số cuối kỳ', 'Số đầu năm'], 'b02': ['Chỉ tiêu', 'Mã số', 'Kỳ này', 'Kỳ trước'],
                'b03': ['Chỉ tiêu', 'Mã số', 'Kỳ này']}[kind]
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font = bold
        keys = {'b01': ('closing', 'opening'), 'b02': ('current', 'previous'), 'b03': ('current',)}[kind]
        for r in data['rows']:
            ws.append(['   ' * max(0, r['level'] - 1) + r['name'], r['code']] + [r[k] for k in keys])
            if r['total']:
                for c in ws[ws.max_row]:
                    c.font = bold
        signatures()
    elif kind == 'general_journal':
        data = general_journal(from_date, to_date)
        ws.title = 'Nhật ký chung'
        header('SỔ NHẬT KÝ CHUNG', 'S03a-DN', 'Từ ngày %s đến ngày %s' % (data['from_date'], data['to_date']))
        ws.append(['STT', 'Ngày ghi sổ', 'Số chứng từ', 'Diễn giải', 'Số hiệu TK', 'Đối tượng', 'Nợ', 'Có'])
        for r in data['rows']:
            ws.append([r['line'], r['posting_date'], r['voucher_no'], r['remarks'], r['number'], r['party'], r['debit'], r['credit']])
        ws.append(['', '', '', 'Cộng', '', '', data['debit'], data['credit']])
        signatures()
    elif kind == 'general_ledger':
        data = general_ledger(account, from_date, to_date, party)
        ws.title = 'Sổ cái'
        header('SỔ CÁI / SỔ CHI TIẾT TÀI KHOẢN ' + data['label'] + ((' – ' + data['party']) if data['party'] else ''), 'S03b-DN',
               'Từ ngày %s đến ngày %s' % (data['from_date'], data['to_date']))
        ws.append(['Ngày', 'Số chứng từ', 'Diễn giải', 'TK đối ứng', 'Đối tượng', 'Nợ', 'Có', 'Số dư'])
        ws.append(['', '', 'Số dư đầu kỳ', '', '', '', '', data['opening']])
        for r in data['lines']:
            ws.append([r['posting_date'], r['voucher_no'], r['remarks'], r['contra'], r['party'], r['debit'], r['credit'], r['balance']])
        ws.append(['', '', 'Cộng phát sinh', '', '', data['debit'], data['credit'], data['closing']])
        signatures()
    elif kind == 'party_summary':
        data = party_summary(account, from_date, to_date)
        ws.title = 'Tổng hợp công nợ'
        header('BẢNG TỔNG HỢP CHI TIẾT CÔNG NỢ TÀI KHOẢN ' + str(account), '', 'Từ ngày %s đến ngày %s' % (data['from_date'], data['to_date']))
        ws.append(['Mã', 'Tên đối tượng', 'Dư Nợ đầu kỳ', 'Dư Có đầu kỳ', 'Phát sinh Nợ', 'Phát sinh Có', 'Dư Nợ cuối kỳ', 'Dư Có cuối kỳ'])
        for r in data['rows']:
            ws.append([r['party'], r['party_name'], r['opening_debit'], r['opening_credit'], r['debit'], r['credit'], r['closing_debit'], r['closing_credit']])
        t = data['totals']
        ws.append(['', 'Cộng', t['opening_debit'], t['opening_credit'], t['debit'], t['credit'], t['closing_debit'], t['closing_credit']])
    elif kind == 'stock_summary':
        data = stock_summary(from_date, to_date)
        ws.title = 'Nhập xuất tồn'
        header('BẢNG TỔNG HỢP NHẬP – XUẤT – TỒN', '', 'Từ ngày %s đến ngày %s' % (data['from_date'], data['to_date']))
        ws.append(['Mã hàng', 'Tên hàng', 'ĐVT', 'Kho', 'Tồn đầu SL', 'Tồn đầu GT', 'Nhập SL', 'Nhập GT', 'Xuất SL', 'Xuất GT', 'Tồn cuối SL', 'Tồn cuối GT'])
        for r in data['rows']:
            ws.append([r['item_code'], r['item_name'], r['stock_uom'], r['warehouse'], r['opening_qty'], r['opening_value'], r['in_qty'], r['in_value'],
                       r['out_qty'], r['out_value'], r['closing_qty'], r['closing_value']])
    else:
        frappe.throw('Loại báo cáo không hợp lệ.')
    output = BytesIO()
    book.save(output)
    frappe.local.response.filename = '%s-%s.xlsx' % (kind, to_date or nowdate())
    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.type = 'download'
