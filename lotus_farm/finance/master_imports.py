"""Previewed, additive Excel imports for accounts and asset categories; never post a balance."""
import base64
from io import BytesIO
from zipfile import BadZipFile, ZipFile

import frappe
from frappe.utils import cint

from lotus_farm.finance import asset_rules, assets, setup

ACCOUNT_FIELDS = ('account_number', 'account_name', 'parent_account', 'is_group', 'account_type')
ACCOUNT_LABELS = ('Số tài khoản', 'Tên tài khoản', 'Tài khoản cha', 'Là nhóm (0/1)', 'Loại tài khoản')
CATEGORY_ACCOUNTS = (
    ('fixed_asset_account', 'Tài khoản nguyên giá', 'Asset'),
    ('accumulated_depreciation_account', 'Tài khoản hao mòn', 'Asset'),
    ('depreciation_expense_account', 'Tài khoản khấu hao', 'Expense'),
    ('capital_work_in_progress_account', 'Tài khoản mua sắm chờ ghi tăng', 'Asset'),
)
CATEGORY_FIELDS = ('asset_category_name', 'non_depreciable_category') + tuple(r[0] for r in CATEGORY_ACCOUNTS)
CATEGORY_LABELS = ('Tên nhóm tài sản', 'Không khấu hao (0/1)') + tuple(r[1] for r in CATEGORY_ACCOUNTS)
KINDS = {'accounts': (ACCOUNT_FIELDS, ACCOUNT_LABELS, 'he-thong-tai-khoan.xlsx'),
         'categories': (CATEGORY_FIELDS, CATEGORY_LABELS, 'nhom-tai-san.xlsx')}


def _kind(kind):
    if kind not in KINDS:
        frappe.throw('Chọn nhập tài khoản hoặc nhóm tài sản.')
    return KINDS[kind]


def _text(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() if value is not None else ''


def _flag(value, field, errors):
    value = _text(value) or '0'
    if value not in ('0', '1'):
        errors[field] = 'Chỉ nhận 0 hoặc 1.'
    return int(value) if value in ('0', '1') else 0


def _accounts(lock=False):
    # Locking reads see the latest committed state, including a concurrent Desk edit.
    return frappe.db.sql('''select name, account_number, account_name, parent_account, is_group,
        root_type, account_type, disabled, modified from `tabAccount` where company=%s
        order by name''' + (' for update' if lock else ''), setup.company(), as_dict=True)


def _resolve(value, accounts):
    value = _text(value)
    matches = [a for a in accounts if a.name == value or (a.account_number and a.account_number == value)] if value else []
    return matches[0] if len(matches) == 1 else None


def _read(content, fields, labels):
    from openpyxl import load_workbook
    encoded = str(content).split(',', 1)[-1]
    if len(encoded) > 6_700_000:
        frappe.throw('File vượt quá 5 MB.')
    try:
        raw = base64.b64decode(encoded, validate=True)
        if len(raw) > 5_000_000:
            frappe.throw('File vượt quá 5 MB.')
        with ZipFile(BytesIO(raw)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 50_000_000:
                frappe.throw('File giải nén vượt quá 50 MB.')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        try:
            sheet = book.worksheets[0]
            if sheet.max_row and sheet.max_row > 2001:
                frappe.throw('Mỗi lần nhập tối đa 2000 dòng.')
            if sheet.max_column and sheet.max_column > len(fields):
                frappe.throw('File có cột thừa. Giữ đúng các cột trong mẫu.')
            cells = list(sheet.iter_rows(max_row=2002, max_col=len(fields)))
        finally:
            book.close()
    except (BadZipFile, KeyError, OSError, ValueError, IndexError):
        frappe.throw('File Excel không hợp lệ. Chọn file .xlsx từ mẫu.')
    if not cells or tuple(_text(c.value) for c in cells[0]) not in (fields, labels):
        frappe.throw('Dùng đúng file mẫu của danh mục đang nhập.')
    rows = []
    for line, row in enumerate(cells[1:], 2):
        if not any(c.value is not None for c in row):
            continue
        if line > 2001:
            frappe.throw('Mỗi lần nhập tối đa 2000 dòng.')
        errors = {fields[i]: 'Thay công thức Excel bằng giá trị.' for i, c in enumerate(row) if c.data_type == 'f'}
        rows.append(dict(line=line, values={k: _text(c.value) for k, c in zip(fields, row)}, errors=errors, action='create'))
    return rows


@frappe.whitelist()
def template(kind):
    setup.require_chief()
    fields, labels, filename = _kind(kind)
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    book = Workbook()
    sheet = book.active
    sheet.title = 'Nhap du lieu'
    sheet.append(labels)
    sheet.freeze_panes = 'A2'
    for i in range(1, len(fields) + 1):
        sheet.column_dimensions[get_column_letter(i)].width = 32
        sheet.column_dimensions[get_column_letter(i)].number_format = '@'
    flag = DataValidation(type='list', formula1='"0,1"', allow_blank=True)
    sheet.add_data_validation(flag)
    flag.add(('D' if kind == 'accounts' else 'B') + '2:' + ('D' if kind == 'accounts' else 'B') + '2001')
    reference = book.create_sheet('Tai khoan hien co')
    reference.append(('Mã đầy đủ',) + ACCOUNT_LABELS + ('Nhóm gốc', 'Ngừng dùng'))
    for a in _accounts():
        reference.append([a.name] + [a.get(k) for k in ACCOUNT_FIELDS] + [a.root_type, a.disabled])
    types = book.create_sheet('Loai tai khoan')
    types.append(['Loại tài khoản'])
    for value in (frappe.get_meta('Account').get_field('account_type').options or '').splitlines():
        if value:
            types.append([value])
    for ws in (reference, types):
        ws.freeze_panes = 'A2'
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 32
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    assets._download(book, filename)


def _account_rows(rows, accounts):
    by_code, seen = {}, set()
    allowed = set((frappe.get_meta('Account').get_field('account_type').options or '').splitlines()) | {''}
    for row in rows:
        v, errors = row['values'], row['errors']
        code = v['account_number']
        if not code or not code.isascii() or not code.isdigit() or len(code) > 20:
            errors['account_number'] = 'Nhập số tài khoản gồm tối đa 20 chữ số.'
        if code in seen:
            errors['account_number'] = 'Số tài khoản trùng trong file.'
        seen.add(code)
        by_code[code] = row
        if not v['account_name'] or len(v['account_name']) > 100:
            errors['account_name'] = 'Nhập tên tài khoản, tối đa 100 ký tự.'
        v['is_group'] = _flag(v['is_group'], 'is_group', errors)
        if v['account_type'] not in allowed:
            errors['account_type'] = 'Loại tài khoản không hợp lệ; xem sheet Loai tai khoan.'

    done, visiting, order = set(), set(), []

    def visit(row):
        v, errors = row['values'], row['errors']
        code = v['account_number']
        if code in done:
            return
        if code in visiting:
            errors['parent_account'] = 'Quan hệ tài khoản cha bị vòng lặp.'
            return
        if len(visiting) >= 100:
            errors['parent_account'] = 'Cây tài khoản vượt quá 100 cấp.'
            return
        visiting.add(code)
        parent = _resolve(v['parent_account'], accounts)
        planned = by_code.get(v['parent_account'])
        if parent:
            if not parent.is_group or parent.disabled:
                errors['parent_account'] = 'Tài khoản cha phải là nhóm đang dùng.'
            v['parent_account'] = parent.name
            v['root_type'] = parent.root_type
        elif planned:
            visit(planned)
            if planned['errors'] or not planned['values']['is_group']:
                errors['parent_account'] = 'Tài khoản cha trong file phải là nhóm hợp lệ.'
            v['root_type'] = planned['values'].get('root_type', '')
        else:
            errors['parent_account'] = 'Không tìm thấy tài khoản cha; dùng số TK hoặc mã đầy đủ từ sheet tham chiếu.'
        parent_type = parent.account_type if parent else planned['values']['account_type'] if planned else ''
        if v['account_type'] in ('Direct Income', 'Indirect Income', 'Current Asset', 'Current Liability', 'Direct Expense', 'Indirect Expense') and v['account_type'] == parent_type:
            errors['account_type'] = 'Loại tài khoản này chỉ đặt ở tài khoản cha; chọn loại khác cho tài khoản con.'
        existing = _resolve(code, accounts)
        if existing:
            row['action'] = 'skip'
            same = all(_text(existing.get(k)) == _text(v.get(k)) for k in ('account_name', 'parent_account', 'is_group', 'account_type'))
            if not same or existing.disabled:
                errors['account_number'] = 'Tài khoản đã có nhưng khác thông tin hoặc ngừng dùng. File nhập không ghi đè tài khoản.'
        visiting.remove(code)
        done.add(code)
        order.append(row['line'])

    for row in rows:
        visit(row)
    return order


def _category_rows(rows, accounts, lock=False):
    seen = set()
    for row in rows:
        v, errors = row['values'], row['errors']
        title = v['asset_category_name']
        if not title or len(title) > 100:
            errors['asset_category_name'] = 'Nhập tên nhóm, tối đa 100 ký tự.'
        if title.casefold() in seen:
            errors['asset_category_name'] = 'Tên nhóm trùng trong file.'
        seen.add(title.casefold())
        v['non_depreciable_category'] = _flag(v['non_depreciable_category'], 'non_depreciable_category', errors)
        for key, label, root in CATEGORY_ACCOUNTS:
            account = _resolve(v[key], accounts)
            if not account or account.is_group or account.disabled or account.root_type != root:
                errors[key] = label + ': chọn số TK chi tiết đang dùng, đúng loại và công ty.'
            else:
                v[key] = account.name
        if title and frappe.db.exists('Asset Category', title):
            category = frappe.get_doc('Asset Category', title, for_update=lock)
            company_rows = [r for r in category.accounts if r.company_name == setup.company()]
            same = len(company_rows) == 1 and all(company_rows[0].get(k) == v[k] for k, _, _ in CATEGORY_ACCOUNTS)
            row['action'] = 'skip'
            if not same or cint(category.non_depreciable_category) != v['non_depreciable_category'] or not category.enable_cwip_accounting:
                errors['asset_category_name'] = 'Nhóm đã có nhưng khác cấu hình. File nhập không ghi đè nhóm tài sản.'
    return [r['line'] for r in rows]


def _preview(kind, content, lock=False):
    fields, labels, _ = _kind(kind)
    rows = _read(content, fields, labels)
    accounts = _accounts(lock)
    order = _account_rows(rows, accounts) if kind == 'accounts' else _category_rows(rows, accounts, lock)
    # Include source state so an intervening rename, disable or mapping edit requires another preview.
    fingerprint = asset_rules.fingerprint(dict(kind=kind, company=setup.company(), rows=rows, accounts=accounts))
    return dict(rows=rows, order=order, fingerprint=fingerprint, valid=bool(rows) and all(not r['errors'] for r in rows),
                create_count=sum(r['action'] == 'create' for r in rows), skip_count=sum(r['action'] == 'skip' for r in rows))


@frappe.whitelist(methods=['POST'])
def preview(kind, content):
    setup.require_chief()
    return _preview(kind, content)


@frappe.whitelist(methods=['POST'])
@assets.retry_approval
def commit(kind, content, fingerprint):
    setup.require_chief()
    with assets.atomic():
        assets._lock(setup.SETTINGS, setup.SETTINGS)
        data = _preview(kind, content, lock=True)
        if not data['valid'] or data['fingerprint'] != fingerprint:
            frappe.throw('File hoặc danh mục đã thay đổi/còn lỗi. Xem trước lại.')
        names, created_accounts = [], {}
        rows = {r['line']: r for r in data['rows']}
        for line in data['order']:
            row = rows[line]
            if row['action'] != 'create':
                continue
            v = dict(row['values'])
            if kind == 'categories':
                names.append(assets.save_category(v)['name'])
            else:
                v['parent_account'] = created_accounts.get(v['parent_account'], v['parent_account'])
                doc = frappe.get_doc(dict(doctype='Account', company=setup.company(), **v))
                doc.insert(ignore_permissions=True)
                created_accounts[v['account_number']] = doc.name
                names.append(doc.name)
        if names:
            cfg = setup.settings()
            cfg.add_comment('Info', 'Nhập Excel %s: thêm %d, bỏ qua %d; người xác nhận %s.' % (
                'tài khoản' if kind == 'accounts' else 'nhóm tài sản', len(names), data['skip_count'], frappe.session.user))
    return dict(names=names, count=len(names), skipped=data['skip_count'])
