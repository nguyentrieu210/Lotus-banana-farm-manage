"""Reviewed fixed-asset operations and depreciation on ERPNext's asset register and ledger."""
from contextlib import contextmanager
from functools import wraps
from io import BytesIO
import base64
import json

import frappe
from frappe.utils import cint, flt, getdate, now_datetime, nowdate

from lotus_farm.finance import asset_rules as ar, posting, rules, setup

OP = 'Lotus Asset Operation'
ALLOCATION = 'Lotus Asset Allocation'
RUN = 'Lotus Asset Depreciation Run'
CATEGORY_ACCOUNT_TYPES = dict(fixed_asset_account='Fixed Asset', accumulated_depreciation_account='Accumulated Depreciation',
                              depreciation_expense_account='Depreciation', capital_work_in_progress_account='Capital Work in Progress')
FIELDS = ('name', 'asset_name', 'asset_category', 'lotus_asset_code', 'lotus_serial_no', 'location',
          'custodian', 'department', 'status', 'docstatus', 'purchase_date', 'available_for_use_date',
          'net_purchase_amount', 'total_asset_cost', 'opening_accumulated_depreciation', 'value_after_depreciation',
          'lotus_farm', 'lotus_farm_block', 'lotus_crop_cycle', 'lotus_work_type', 'lotus_usage_suspended', 'modified')


def _json(value, default=None):
    return frappe.parse_json(value) if isinstance(value, str) else (value if value is not None else default)


def _dump(value):
    return json.dumps(value, ensure_ascii=False, default=str)


@contextmanager
def internal(**context):
    previous = getattr(frappe.flags, 'lotus_asset_context', None)
    frappe.flags.lotus_asset_context = context or {'write': True}
    try:
        yield
    finally:
        frappe.flags.lotus_asset_context = previous


@contextmanager
def atomic():
    name = 'assets_' + frappe.generate_hash(length=10)
    frappe.db.savepoint(name)
    depth = cint(getattr(frappe.flags, 'lotus_asset_atomic_depth', 0))
    frappe.flags.lotus_asset_atomic_depth = depth + 1
    try:
        yield
    except Exception:
        try:
            frappe.db.rollback(save_point=name)
        except Exception:
            # MariaDB rolls the entire transaction back on a deadlock, dropping all savepoints.
            frappe.db.rollback()
        raise
    finally:
        frappe.flags.lotus_asset_atomic_depth = depth


def retry_approval(function):
    """A raced locking read can invalidate a MariaDB snapshot. Retry the whole atomic approval."""
    @wraps(function)
    def wrapped(*args, **kwargs):
        # A nested master write must never restart halfway through an enclosing import.
        if getattr(frappe.flags, 'lotus_asset_atomic_depth', 0):
            return function(*args, **kwargs)
        for attempt in range(3):
            try:
                return function(*args, **kwargs)
            except frappe.QueryDeadlockError:
                frappe.db.rollback()
                if attempt == 2:
                    frappe.throw('Dữ liệu đang được xử lý đồng thời. Tải lại và xem trước trước khi duyệt.')
    return wrapped


def enabled():
    return bool(frappe.db.exists('DocType', setup.SETTINGS) and
                frappe.db.get_single_value(setup.SETTINGS, 'fixed_assets_enabled'))


def require_enabled():
    setup.require_enabled()
    if not enabled():
        frappe.throw('Kế toán trưởng bật Tài sản cố định trước khi thực hiện nghiệp vụ.')


def _save(doc):
    doc.flags.ignore_permissions = True
    with internal():
        doc.save()
    return doc


def _lock(doctype, name):
    from lotus_farm.controls import lock
    # Singles are stored in tabSingles, not a table named after the DocType.
    lock('Company', setup.company()) if doctype == setup.SETTINGS else lock(doctype, name)


def asset_doc(name):
    doc = frappe.get_doc('Asset', name, for_update=bool(getattr(frappe.flags, 'lotus_asset_context', None)))
    if doc.company != setup.company() or not doc.get('lotus_managed_asset'):
        frappe.throw('Tài sản không thuộc danh mục Lotus của công ty đang chọn.', frappe.PermissionError)
    return doc


def _account(name, kind=None, account_type=None):
    row = frappe.db.get_value('Account', name, ['name', 'company', 'is_group', 'disabled', 'root_type', 'account_type'], as_dict=True)
    if not row or row.company != setup.company() or row.is_group or row.disabled or (kind and row.root_type != kind) or (account_type and row.account_type != account_type):
        frappe.throw('Chọn tài khoản chi tiết đang dùng, đúng công ty và đúng loại: %s.' % (name or 'trống'))
    return row.name


def _dimensions(row):
    result = {k: row.get(k) or None for k in ar.DIMENSIONS}
    farm = result['lotus_farm']
    if farm and frappe.db.get_value('Farm', farm, 'company') != setup.company():
        frappe.throw('Vườn không thuộc công ty.')
    if result['lotus_farm_block']:
        block = frappe.db.get_value('Farm Block', result['lotus_farm_block'], 'farm')
        if not block or block != farm:
            frappe.throw('Lô không thuộc vườn đã chọn.')
    if result['lotus_crop_cycle']:
        meta = frappe.get_meta('Crop Cycle')
        cycle = frappe.get_doc('Crop Cycle', result['lotus_crop_cycle'])
        if meta.has_field('farm') and cycle.get('farm') and cycle.farm != farm:
            frappe.throw('Vụ không thuộc vườn đã chọn.')
        if meta.has_field('farm_block') and cycle.get('farm_block') and cycle.farm_block != result['lotus_farm_block']:
            frappe.throw('Vụ không thuộc lô đã chọn.')
    if result['lotus_work_type'] and not frappe.db.exists('Work Type', result['lotus_work_type']):
        frappe.throw('Công việc không tồn tại.')
    return result


def _allocation_rows(rows):
    try:
        rows = ar.allocation_rows(rows)
    except ValueError as exc:
        frappe.throw(str(exc))
    for row in rows:
        _account(row['expense_account'], 'Expense')
        number = frappe.db.get_value('Account', row['expense_account'], 'account_number') or ''
        if number.startswith(('623', '627')) and not row.get('lotus_farm'):
            frappe.throw('Phân bổ chi phí sản xuất cần chọn vườn; chi phí dùng chung vườn có thể để trống lô.')
        if frappe.db.get_value('Cost Center', row['cost_center'], ['company', 'is_group']) != (setup.company(), 0):
            frappe.throw('Trung tâm chi phí phải là mục chi tiết thuộc công ty.')
        row.update(_dimensions(row))
    return rows


def _bounds(month):
    try:
        return ar.bounds(month)
    except ValueError as exc:
        frappe.throw(str(exc))


def check_period(on_date):
    setup.check_cutover(on_date)
    closed = frappe.db.sql("""select ap.name from `tabAccounting Period` ap join `tabClosed Document` cd on cd.parent=ap.name
        where ap.company=%s and ap.disabled=0 and %s between ap.start_date and ap.end_date
        and cd.document_type='Journal Entry' and cd.closed=1 limit 1 for update""", (setup.company(), getdate(on_date)))
    if closed:
        frappe.throw('Kỳ kế toán đã khóa. Kế toán trưởng mở lại kỳ trước khi thực hiện nghiệp vụ tài sản.')
    posting.check_open(setup.company(), on_date)
    month = str(getdate(on_date))[:7]
    if frappe.db.exists('Lotus Costing Run', {'company': setup.company(), 'month': month, 'state': 'Posted'}):
        frappe.throw('Tháng %s đã tính giá thành. Đảo giá thành và kết chuyển trước.' % month)
    for purpose in ('Costing', 'Cost Closing', 'Result Closing'):
        if posting.active(setup.SETTINGS, setup.SETTINGS, purpose, month):
            frappe.throw('Tháng %s đã kết chuyển hoặc tính giá thành. Đảo giá thành và kết chuyển trước.' % month)


def validate_business_document(doc):
    if frappe.flags.in_install or frappe.flags.in_migrate:
        return
    if not getattr(frappe.flags, 'lotus_asset_context', None):
        frappe.throw('Lưu và duyệt chứng từ tài sản qua quy trình Tài sản cố định Lotus.', frappe.PermissionError)
    if doc.company != setup.company():
        frappe.throw('Chứng từ không thuộc công ty đang chọn.', frappe.PermissionError)


def native_guard(doc, method=None):
    """Protect native/Desk endpoints too. Approved server operations alone may post managed assets."""
    if frappe.flags.in_install or frappe.flags.in_migrate:
        return
    if doc.doctype == setup.SETTINGS:
        old = doc.get_doc_before_save()
        if old and old.get('fixed_assets_enabled') != doc.get('fixed_assets_enabled') and not getattr(frappe.flags, 'lotus_asset_context', None):
            frappe.throw('Bật/tắt tài sản qua mục Tài sản cố định.', frappe.PermissionError)
        return
    if not enabled():
        return
    if doc.doctype == 'Asset Category':
        if not any(r.company_name == setup.company() for r in doc.get('accounts') or []):
            return
        if not getattr(frappe.flags, 'lotus_asset_context', None):
            frappe.throw('Cấu hình nhóm tài sản do kế toán trưởng thực hiện qua Lotus.', frappe.PermissionError)
        return
    if doc.doctype == 'Accounts Settings':
        if doc.get('book_asset_depreciation_entry_automatically'):
            frappe.throw('Lotus đang duyệt khấu hao thủ công; không bật tự động ghi khấu hao.')
        return
    if doc.get('company') != setup.company():
        return
    related = doc.doctype in ('Asset', 'Asset Movement', 'Asset Repair', 'Asset Capitalization', 'Asset Value Adjustment', 'Asset Depreciation Schedule')
    if doc.doctype == 'Journal Entry':
        related = doc.voucher_type == 'Depreciation Entry' or doc.get('lotus_asset_run') or doc.get('lotus_asset_operation') or any(r.reference_type == 'Asset' for r in doc.accounts)
        context = getattr(frappe.flags, 'lotus_asset_context', None)
        if related and context:
            doc.lotus_asset_operation = doc.get('lotus_asset_operation') or context.get('operation')
            if context.get('operation'):
                check_period(doc.posting_date)
    if doc.doctype in ('Purchase Invoice', 'Sales Invoice'):
        related = any(r.get('asset') or r.get('is_fixed_asset') for r in doc.get('items') or [])
    if related and not getattr(frappe.flags, 'lotus_asset_context', None):
        if doc.doctype == 'Purchase Invoice':
            # The existing purchasing workflow supplies the source invoice; its chief may approve it
            # before a separate register operation links the receipt/invoice to the Asset.
            setup.require_chief()
            check_period(doc.posting_date)
            return
        frappe.throw('Nghiệp vụ tài sản phải được xem trước và kế toán trưởng duyệt trong Lotus.', frappe.PermissionError)


@frappe.whitelist(methods=['POST'])
def activate(on=1):
    setup.require_chief()
    setup.require_enabled()
    with atomic(), internal():
        cfg = setup.settings()
        _lock(setup.SETTINGS, cfg.name)
        if cint(on):
            if frappe.get_cached_value('Company', setup.company(), 'default_currency') != 'VND':
                frappe.throw('Tài sản cố định Lotus trong đợt này chỉ hỗ trợ sổ VND.')
            if frappe.db.count('Company') != 1:
                frappe.throw('Bản này chỉ hỗ trợ một công ty; cần cấu hình riêng trước khi bật cho nhiều công ty.')
            auto = frappe.get_single('Accounts Settings')
            if not cfg.get('fixed_assets_enabled'):
                cfg.asset_previous_auto_depreciation = auto.book_asset_depreciation_entry_automatically
            auto.book_asset_depreciation_entry_automatically = 0
            auto.flags.ignore_permissions = True
            auto.save()
        cfg.fixed_assets_enabled = cint(on)
        cfg.flags.ignore_permissions = True
        cfg.save()
        cfg.add_comment('Info', ('Bật' if cint(on) else 'Tắt') + ' quản lý tài sản Lotus; tự động khấu hao vẫn tắt.')
    return overview()


@frappe.whitelist()
def overview(month=None):
    setup.require_view()
    company = setup.company()
    month = month or nowdate()[:7]
    _bounds(month)
    rows = frappe.get_all('Asset', filters={'company': company, 'lotus_managed_asset': 1}, fields=list(FIELDS),
                          order_by='lotus_asset_code', limit_page_length=0)
    for row in rows:
        row['accumulated_depreciation'] = flt(row.total_asset_cost or row.net_purchase_amount) - flt(row.value_after_depreciation) if row.docstatus == 1 else None
    pending = []
    # Query independently: old pending registrations must not disappear behind the history limit.
    for op in frappe.get_all(OP, filters={'company': company, 'kind': 'Create', 'state': 'Pending'},
                             fields=['name', 'modified', 'posting_date', 'payload'], order_by='creation desc', limit_page_length=0):
        v = _json(op.payload, {})
        pending.append(dict(name=op.name, operation=op.name, modified=op.modified, posting_date=op.posting_date,
                            values=v, status='Pending', lotus_asset_code=v.get('code'), asset_name=v.get('asset_name'),
                            asset_category=v.get('asset_category'), location=v.get('location'), lotus_serial_no=v.get('serial_no'),
                            available_for_use_date=v.get('available_for_use_date'), total_asset_cost=v.get('amount'),
                            accumulated_depreciation=None, value_after_depreciation=None,
                            **{k: v.get(k) for k in ar.DIMENSIONS}))
    operations = frappe.get_all(OP, filters={'company': company}, fields=['name', 'asset', 'kind', 'posting_date', 'state', 'reason', 'modified',
                                                                                      'approved_by', 'approved_at', 'documents', 'payload'],
                                order_by='creation desc', limit_page_length=500)
    runs = frappe.get_all(RUN, filters={'company': company}, fields=['name', 'month', 'state', 'approved_by', 'approved_at', 'modified'],
                          order_by='month desc, creation desc', limit_page_length=100)
    return dict(enabled=enabled(), company=company, month=month, assets=rows, pending_assets=pending, operations=operations, runs=runs,
                permissions=dict(write=setup.can_post(), approve=setup.is_chief()),
                totals=dict(count=len(rows), original=sum(flt(r.total_asset_cost or r.net_purchase_amount) for r in rows if r.docstatus == 1 and r.status not in ('Sold', 'Scrapped')),
                            remaining=sum(flt(r.value_after_depreciation) for r in rows if r.docstatus == 1 and r.status not in ('Sold', 'Scrapped'))))


@frappe.whitelist()
def catalogs():
    """Visible masters, company-specific account mappings and permissions for the asset workspace."""
    setup.require_view()
    company = setup.company()
    groups = frappe.get_all('Asset Category', fields=['name', 'asset_category_name', 'non_depreciable_category', 'modified'],
                             order_by='asset_category_name', limit_page_length=0)
    account_fields = ('fixed_asset_account', 'accumulated_depreciation_account', 'depreciation_expense_account', 'capital_work_in_progress_account')
    child_type = frappe.get_meta('Asset Category').get_field('accounts').options
    mappings = {r.parent: r for r in frappe.get_all(child_type, filters={'company_name': company},
                fields=['parent'] + list(account_fields), limit_page_length=0)}
    assets = frappe.get_all('Asset', filters={'company': company, 'docstatus': 1}, fields=['asset_category', 'location'], limit_page_length=0)
    used = set(frappe.get_all('Asset', filters={'docstatus': 1}, pluck='asset_category', limit_page_length=0))
    pending = [_json(r.payload, {}) for r in frappe.get_all(OP, filters={'company': company, 'kind': 'Create', 'state': 'Pending'},
                                                           fields=['payload'], limit_page_length=0)]
    for row in groups:
        row.update({k: mappings.get(row.name, {}).get(k) for k in account_fields})
        row['asset_count'] = sum(a.asset_category == row.name for a in assets)
        row['pending_count'] = sum(a.get('asset_category') == row.name for a in pending)
        row['configured'] = all(row.get(k) for k in account_fields)
        row['can_edit'] = setup.is_chief() and row.name not in used
    locations = frappe.get_all('Location', filters={'is_group': 0}, fields=['name', 'location_name', 'lotus_display_name', 'modified'],
                               order_by='location_name', limit_page_length=0)
    for row in locations:
        row['location_name'] = row.lotus_display_name or row.location_name
        row['asset_count'] = sum(a.location == row.name for a in assets)
        row['pending_count'] = sum(a.get('location') == row.name for a in pending)
        row['can_edit'] = setup.is_chief()
    accounts = frappe.get_all('Account', filters={'company': company}, fields=['name', 'account_number', 'account_name', 'parent_account',
                             'root_type', 'account_type', 'is_group', 'disabled'], order_by='account_number, account_name', limit_page_length=0)
    for row in accounts:
        row['categories'] = [g.name for g in groups if row.name in [g.get(k) for k in account_fields]]
    return dict(categories=groups, locations=locations, accounts=accounts,
                permissions=dict(create_category=setup.is_chief(), create_location=setup.can_post(), import_accounts=setup.is_chief()))


@frappe.whitelist()
def lookups():
    setup.require_view()
    from lotus_farm.finance.reports import lookups as finance_lookups
    result = finance_lookups()
    company = setup.company()
    result.update(categories=frappe.get_all('Asset Category', fields=['name', 'asset_category_name', 'non_depreciable_category'], limit_page_length=0),
                  locations=frappe.get_all('Location', filters={'is_group': 0}, fields=['name', 'location_name', 'lotus_display_name'], limit_page_length=0),
                  employees=frappe.get_all('Employee', filters={'company': company, 'status': 'Active'}, fields=['name', 'employee_name'], limit_page_length=0),
                  cost_centers=frappe.get_all('Cost Center', filters={'company': company, 'is_group': 0}, fields=['name', 'cost_center_name'], limit_page_length=0),
                  items=frappe.get_all('Item', filters={'is_fixed_asset': 1, 'disabled': 0}, fields=['name', 'item_name', 'asset_category'], limit_page_length=0),
                  purchases=frappe.get_all('Purchase Invoice', filters={'company': company, 'docstatus': 1}, fields=['name', 'supplier', 'posting_date'], order_by='posting_date desc', limit_page_length=200),
                  customers=frappe.get_all('Customer', filters={'disabled': 0}, fields=['name', 'customer_name'], limit_page_length=0),
                  suppliers=frappe.get_all('Supplier', filters={'disabled': 0}, fields=['name', 'supplier_name'], limit_page_length=0),
                  cutover_date=setup.settings().cutover_date)
    result['accounts'] = frappe.get_all('Account', filters={'company': company, 'is_group': 0, 'disabled': 0},
                                         fields=['name', 'account_number', 'account_name', 'root_type', 'account_type'], order_by='account_number', limit_page_length=0)
    result['farms'] = frappe.get_all('Farm', filters={'company': company}, fields=['name', 'farm_name'], limit_page_length=0)
    result['cycles'] = frappe.get_all('Crop Cycle', fields=['name']+(['farm'] if frappe.get_meta('Crop Cycle').has_field('farm') else []), limit_page_length=0)
    for row in result['locations']:
        row['location_name'] = row.lotus_display_name or row.location_name
    return result


@frappe.whitelist(methods=['POST'])
@retry_approval
def save_category(values):
    setup.require_chief()
    _lock(setup.SETTINGS, setup.SETTINGS)
    v = _json(values, {})
    doc = frappe.get_doc('Asset Category', v['name'], for_update=True) if v.get('name') else frappe.new_doc('Asset Category')
    if v.get('modified') and str(doc.modified) != str(v['modified']):
        frappe.throw('Nhóm tài sản đã thay đổi. Tải lại trước khi lưu.')
    if v.get('name') and str(v.get('asset_category_name') or '').strip() != doc.asset_category_name:
        frappe.throw('Giữ nguyên tên định danh của nhóm đã tạo.')
    doc.asset_category_name = str(v.get('asset_category_name') or '').strip()
    if not doc.asset_category_name:
        frappe.throw('Nhập tên nhóm tài sản.')
    if doc.name and frappe.db.exists('Asset', {'asset_category': doc.name, 'docstatus': 1}):
        frappe.throw('Nhóm đã có tài sản được duyệt. Tạo nhóm mới để thay đổi chính sách tài khoản.')
    doc.enable_cwip_accounting = 1
    doc.non_depreciable_category = cint(v.get('non_depreciable_category'))
    accounts = dict(company_name=setup.company())
    for key, kind in [('fixed_asset_account', 'Asset'), ('accumulated_depreciation_account', 'Asset'),
                      ('depreciation_expense_account', 'Expense'), ('capital_work_in_progress_account', 'Asset')]:
        accounts[key] = _account(v.get(key), kind, CATEGORY_ACCOUNT_TYPES[key])
    other_companies = [r.as_dict() for r in doc.accounts if r.company_name != setup.company()]
    doc.set('accounts', other_companies + [accounts])
    _save(doc)
    return dict(name=doc.name)


@frappe.whitelist()
def purchase_details(name):
    setup.require_view()
    doc = frappe.get_doc('Purchase Invoice', name)
    if doc.company != setup.company() or doc.docstatus != 1:
        frappe.throw('Chọn hóa đơn mua đã duyệt thuộc công ty.', frappe.PermissionError)
    rows = []
    for row in doc.items:
        category = frappe.db.get_value('Item', row.item_code, 'asset_category')
        if frappe.db.get_value('Item', row.item_code, 'is_fixed_asset'):
            rows.append(dict(name=row.name, item_code=row.item_code, item_name=row.item_name,
                             base_net_amount=row.base_net_amount, qty=row.qty, asset_category=category))
    return dict(name=doc.name, items=rows)


@frappe.whitelist(methods=['POST'])
@retry_approval
def save_location(location_name, name=None, modified=None):
    setup.require_post()
    if name:
        setup.require_chief()
        _lock('Location', name)
    title = str(location_name or '').strip()
    if not title:
        frappe.throw('Nhập tên vị trí.')
    doc = frappe.get_doc('Location', name, for_update=True) if name else frappe.get_doc(dict(doctype='Location', is_group=0))
    if doc.is_group:
        frappe.throw('Chỉ sửa vị trí chi tiết.')
    if name and (not modified or str(doc.modified) != str(modified)):
        frappe.throw('Vị trí đã thay đổi. Tải lại trước khi lưu.')
    if name:
        doc.lotus_display_name = title
    else:
        doc.location_name = title
    _save(doc)
    return dict(name=doc.name)


@frappe.whitelist()
def detail(asset):
    setup.require_view()
    doc = asset_doc(asset)
    schedules = []
    for name in frappe.get_all('Asset Depreciation Schedule', filters={'asset': doc.name, 'docstatus': ['<', 2]}, pluck='name'):
        schedule = frappe.get_doc('Asset Depreciation Schedule', name)
        schedules.append(dict(name=name, finance_book=schedule.finance_book, rows=[r.as_dict() for r in schedule.depreciation_schedule]))
    history = frappe.get_all(OP, filters={'company': setup.company(), 'asset': doc.name},
                             fields=['name', 'kind', 'posting_date', 'state', 'reason', 'approved_by', 'approved_at', 'documents'],
                             order_by='posting_date desc, creation desc', limit_page_length=0)
    documents = []
    for op in history:
        op['documents'] = _json(op.documents, [])
        for link in op['documents']:
            if link not in documents:
                documents.append(link)
    for key, doctype in [('purchase_invoice', 'Purchase Invoice'), ('purchase_receipt', 'Purchase Receipt')]:
        if doc.get(key) and dict(doctype=doctype, name=doc.get(key)) not in documents:
            documents.append(dict(doctype=doctype, name=doc.get(key)))
    return dict(asset=doc.as_dict(), schedules=schedules, history=history, documents=documents)


@frappe.whitelist(methods=['POST'])
@retry_approval
def save_operation(kind, posting_date, values, asset=None, name=None, modified=None, reason=None):
    setup.require_post()
    require_enabled()
    if kind not in ar.KINDS:
        frappe.throw('Nghiệp vụ tài sản không hợp lệ.')
    values = _json(values, {})
    errors = {}
    if kind in ('Create', 'Profile'):
        if values.get('asset_category'):
            values['non_depreciable'] = bool(frappe.db.get_value('Asset Category', values['asset_category'], 'non_depreciable_category'))
        values, errors = ar.profile(values)
        if kind == 'Create' and values.get('code'):
            claim = ar.fingerprint(['Create', setup.company(), values['code']])
            duplicate = frappe.db.get_value(OP, {'active_key': claim}, 'name')
            if (duplicate and duplicate != name) or frappe.db.exists('Asset', {'lotus_asset_code': values['code']}):
                errors['code'] = 'Mã tài sản đã có hồ sơ hoặc nghiệp vụ chờ duyệt.'
    if errors:
        return dict(errors=errors)
    if kind != 'Create':
        asset_doc(asset)
    check_period(posting_date)
    if name:
        _lock(OP, name)
    doc = frappe.get_doc(OP, name, for_update=True) if name else frappe.new_doc(OP)
    if name:
        if doc.company != setup.company() or doc.state != 'Pending':
            frappe.throw('Nghiệp vụ đã xử lý hoặc không thuộc công ty.')
        if str(doc.modified) != str(modified):
            frappe.throw('Nghiệp vụ đã thay đổi. Tải lại trước khi lưu.')
    doc.update(dict(company=setup.company(), asset=asset or None, kind=kind, posting_date=getdate(posting_date),
                    payload=_dump(values), reason=str(reason or '').strip(), state='Pending'))
    if kind == 'Create':
        doc.active_key = ar.fingerprint(['Create', setup.company(), values['code']])
    _save(doc)
    return dict(name=doc.name, errors={})


@frappe.whitelist(methods=['POST'])
def save_asset(values):
    v = _json(values, {})
    return save_operation('Profile' if v.get('name') else 'Create', v.get('posting_date') or nowdate(), v,
                          asset=v.get('name'), name=v.get('operation'), modified=v.get('modified'))


def allocation_for(asset, month):
    name = frappe.db.get_value(ALLOCATION, {'company': setup.company(), 'asset': asset, 'month': ['<=', month]},
                               'name', order_by='month desc')
    if name:
        doc = frappe.get_doc(ALLOCATION, name)
        return _allocation_rows([r.as_dict() for r in doc.rows]), doc
    doc = asset_doc(asset)
    from erpnext.assets.doctype.asset.depreciation import get_credit_debit_accounts_for_asset, get_depreciation_cost_center_and_series
    _, expense = get_credit_debit_accounts_for_asset(doc.asset_category, doc.company)
    center, _ = get_depreciation_cost_center_and_series(doc)
    return _allocation_rows([dict(percent=100, expense_account=expense, cost_center=center,
                                 **{k: doc.get(k) for k in ar.DIMENSIONS})]), None


@frappe.whitelist()
def allocation(asset, month):
    setup.require_view()
    _bounds(month)
    rows, doc = allocation_for(asset, month)
    return dict(asset=asset, month=month, rows=rows, modified=str(doc.modified) if doc and doc.month == month else None,
                inherited_from=doc.month if doc else None)


@frappe.whitelist(methods=['POST'])
def save_allocation(asset, month, rows, modified=None):
    setup.require_post()
    require_enabled()
    _lock(setup.SETTINGS, setup.SETTINGS)
    asset_doc(asset)
    _, end = _bounds(month)
    check_period(end)
    _lock('Asset', asset)
    if frappe.db.exists(RUN, {'company': setup.company(), 'month': month, 'state': 'Posted'}):
        frappe.throw('Tháng đã duyệt khấu hao. Đảo bảng trước khi sửa phân bổ.')
    clean = _allocation_rows(_json(rows, []))
    name = frappe.db.get_value(ALLOCATION, {'company': setup.company(), 'asset': asset, 'month': month})
    doc = frappe.get_doc(ALLOCATION, name, for_update=True) if name else frappe.new_doc(ALLOCATION)
    if name and str(doc.modified) != str(modified):
        frappe.throw('Phân bổ đã thay đổi. Tải lại trước khi lưu.')
    doc.update(dict(company=setup.company(), asset=asset, month=month))
    doc.active_key = ar.fingerprint([setup.company(), asset, month])
    doc.set('rows', clean)
    _save(doc)
    return allocation(asset, month)


def opening_reconciliation():
    """Opening Asset documents do not post GL. Reconcile them to the existing opening balance workflow."""
    totals = {}
    for category in frappe.get_all('Asset Category Account', filters={'company_name': setup.company()},
                                  fields=['fixed_asset_account', 'accumulated_depreciation_account']):
        for account in (category.fixed_asset_account, category.accumulated_depreciation_account):
            if account:
                totals.setdefault(account, 0)
    existing_field = 'asset_type' if frappe.get_meta('Asset').has_field('asset_type') else 'is_existing_asset'
    existing_value = 'Existing Asset' if existing_field == 'asset_type' else 1
    for name in frappe.get_all('Asset', filters={'company': setup.company(), 'lotus_managed_asset': 1, existing_field: existing_value, 'docstatus': 1}, pluck='name'):
        doc = asset_doc(name)
        from erpnext.assets.doctype.asset.asset import get_asset_account
        fixed = get_asset_account('fixed_asset_account', asset=doc.name)
        accumulated = get_asset_account('accumulated_depreciation_account', asset=doc.name)
        source = frappe.db.get_value(OP, {'asset': name, 'kind': 'Create', 'state': 'Posted'}, 'payload')
        original = _json(source, {})
        totals[fixed] = totals.get(fixed, 0) + flt(original.get('amount', doc.net_purchase_amount))
        if doc.opening_accumulated_depreciation:
            totals[accumulated] = totals.get(accumulated, 0) - flt(doc.opening_accumulated_depreciation)
    cutover = setup.settings().cutover_date
    result = []
    for account, expected in totals.items():
        actual = flt(frappe.db.sql("select sum(debit-credit) from `tabGL Entry` where company=%s and account=%s and is_cancelled=0 and posting_date<=%s and is_opening='Yes'",
                                  (setup.company(), account, cutover))[0][0])
        result.append(dict(account=account, expected=expected, actual=actual, difference=expected-actual))
    return result


def due_rows(month):
    start, end = _bounds(month)
    cutover = setup.settings().cutover_date or start
    query = """select a.name asset, a.asset_name, a.modified asset_modified, a.asset_category,
        s.name schedule, s.modified schedule_modified, s.finance_book, d.name schedule_row, d.idx,
        d.schedule_date, d.depreciation_amount amount from `tabAsset` a
        join `tabAsset Depreciation Schedule` s on s.asset=a.name and s.docstatus=1
        join `tabDepreciation Schedule` d on d.parent=s.name
        where a.company=%s and a.lotus_managed_asset=1 and a.docstatus=1 and a.calculate_depreciation=1
        and a.status in ('Submitted','Partially Depreciated') and (d.journal_entry is null or d.journal_entry='')
        and d.schedule_date >= %s and d.schedule_date <= %s order by a.name,d.schedule_date"""
    if getattr(frappe.flags, 'lotus_asset_context', None):
        query += ' for update'
    return frappe.db.sql(query, (setup.company(), cutover, end), as_dict=True)


@frappe.whitelist()
def depreciation_preview(month):
    setup.require_view()
    start, end = _bounds(month)
    rows, errors = [], []
    for row in due_rows(month):
        row = dict(row)
        if str(row['schedule_date']) < start:
            errors.append('Tài sản %s còn kỳ khấu hao %s chưa duyệt. Xử lý tháng đó trước.' % (row['asset_name'], row['schedule_date']))
            continue
        try:
            allocations, allocation_doc = allocation_for(row['asset'], month)
            row['allocations'] = ar.split_amount(row['amount'], allocations)
            row['allocation_name'] = allocation_doc.name if allocation_doc else None
            row['allocation_modified'] = allocation_doc.modified if allocation_doc else None
            row['category_modified'] = frappe.db.get_value('Asset Category', row['asset_category'], 'modified')
            rows.append(row)
        except frappe.ValidationError as exc:
            errors.append('%s: %s' % (row['asset_name'], str(exc)))
    reconciliation = opening_reconciliation()
    if any(abs(r['difference']) > .5 for r in reconciliation):
        errors.append('Nguyên giá/hao mòn tài sản đầu kỳ chưa khớp bút toán số dư đầu kỳ.')
    done = frappe.db.get_value(RUN, {'company': setup.company(), 'month': month, 'state': 'Posted'}, 'name')
    result = dict(company=setup.company(), month=month, posting_date=end, rows=rows, errors=errors,
                  opening_reconciliation=reconciliation, total=float(sum(ar.decimal(r['amount']) for r in rows)), posted=done)
    result['fingerprint'] = ar.fingerprint(result)
    return result


def _journal(on_date, lines, remark, run=None, operation=None, source=None):
    doc = frappe.get_doc(dict(doctype='Journal Entry', company=setup.company(), posting_date=on_date,
                             voucher_type='Depreciation Entry' if source else 'Journal Entry', user_remark=remark, lotus_asset_run=run,
                             lotus_asset_operation=operation, lotus_asset_source_journal=source, accounts=lines))
    doc.flags.ignore_permissions = True
    doc.insert()
    doc.submit()
    return doc


def allocate_journal(journal, allocations, run=None, operation=None):
    """Keep native depreciation's Asset references untouched. Reclassify costs in a separate JE."""
    company = setup.company()
    source_rows = [r for r in journal.accounts if flt(r.debit_in_account_currency) and
                   frappe.db.get_value('Account', r.account, 'root_type') == 'Expense']
    if len(source_rows) != 1:
        frappe.throw('Bút toán khấu hao gốc phải có đúng một dòng chi phí.')
    source = source_rows[0]
    amount = ar.decimal(source.debit_in_account_currency)
    destinations = ar.split_amount(amount, allocations)
    fields = ('account', 'cost_center', *ar.DIMENSIONS)
    original = dict(account=source.account, cost_center=source.cost_center, **{k: source.get(k) for k in ar.DIMENSIONS})
    if len(destinations) == 1 and all((original.get(k) or None) == (destinations[0].get('expense_account' if k == 'account' else k) or None) for k in fields):
        return None
    lines = [dict(original, credit_in_account_currency=amount)]
    for row in destinations:
        if row['amount']:
            lines.append(dict(account=row['expense_account'], cost_center=row['cost_center'], debit_in_account_currency=row['amount'],
                              **{k: row.get(k) for k in ar.DIMENSIONS}))
    return _journal(journal.posting_date, lines, 'Phân bổ khấu hao từ ' + journal.name, run, operation, journal.name)


def book_row(row, run=None, operation=None):
    from erpnext.assets.doctype.asset.depreciation import make_depreciation_entry
    current = frappe.db.get_value('Depreciation Schedule', row['schedule_row'], 'journal_entry')
    if current:
        frappe.throw('Kỳ khấu hao này đã được ghi sổ: %s.' % current)
    # Do not use indexed posting: ERPNext's index path can bypass the already-booked check at index zero.
    make_depreciation_entry(row['schedule'], date=str(row['schedule_date']))
    name = frappe.db.get_value('Depreciation Schedule', row['schedule_row'], 'journal_entry')
    if not name:
        frappe.throw('ERPNext chưa tạo bút toán khấu hao.')
    journal = frappe.get_doc('Journal Entry', name)
    journal.db_set(dict(lotus_asset_run=run, lotus_asset_operation=operation), update_modified=False)
    reclass = allocate_journal(journal, row['allocations'], run, operation)
    return dict(asset=row['asset'], schedule_row=row['schedule_row'], journal=name, allocation=reclass.name if reclass else None)


@frappe.whitelist(methods=['POST'])
def depreciation_prepare(month, fingerprint):
    setup.require_post()
    require_enabled()
    with atomic(), internal():
        _lock(setup.SETTINGS, setup.SETTINGS)
        check_period(_bounds(month)[1])
        preview = depreciation_preview(month)
        if preview['fingerprint'] != fingerprint or preview['errors'] or preview['posted'] or not preview['rows']:
            frappe.throw('Bảng đã thay đổi, có lỗi hoặc không còn khấu hao chờ duyệt. Xem trước lại.')
        name = frappe.db.get_value(RUN, {'company': setup.company(), 'month': month, 'state': 'Pending'}, 'name')
        doc = frappe.get_doc(RUN, name, for_update=True) if name else frappe.new_doc(RUN)
        doc.update(dict(company=setup.company(), month=month, state='Pending', snapshot=_dump(preview),
                        preview_hash=fingerprint, active_key=setup.company()+'|'+month+'|pending'))
        doc.flags.ignore_permissions = True
        doc.save()
    return dict(name=doc.name, state='Pending', fingerprint=fingerprint)


@frappe.whitelist()
def depreciation_detail(name):
    setup.require_view()
    doc = frappe.get_doc(RUN, name)
    if doc.company != setup.company():
        frappe.throw('Bảng khấu hao không thuộc công ty.', frappe.PermissionError)
    result = _json(doc.snapshot, {})
    result.update(name=doc.name, state=doc.state, approved_by=doc.approved_by, approved_at=doc.approved_at,
                  posted=doc.name if doc.state != 'Pending' else None, results=_json(doc.results, []))
    return result


@frappe.whitelist(methods=['POST'])
@retry_approval
def depreciation_post(month, fingerprint):
    setup.require_chief()
    require_enabled()
    _, end = _bounds(month)
    if getdate(end) > getdate(nowdate()):
        frappe.throw('Chưa đến ngày kết thúc tháng; có thể xem trước nhưng chưa ghi khấu hao tương lai.')
    with atomic(), internal():
        _lock(setup.SETTINGS, setup.SETTINGS)
        check_period(end)
        result = depreciation_preview(month)
        if result['posted']:
            frappe.throw('Tháng đã được duyệt khấu hao.')
        if result['fingerprint'] != fingerprint:
            frappe.throw('Dữ liệu đã thay đổi. Xem trước bảng khấu hao mới trước khi duyệt.')
        if result['errors']:
            frappe.throw('<br>'.join(result['errors']))
        if not result['rows']:
            frappe.throw('Tháng không có khấu hao đến hạn.')
        for asset in sorted({r['asset'] for r in result['rows']}):
            _lock('Asset', asset)
        # Re-read after taking locks, including the schedule, category and allocation versions.
        if depreciation_preview(month)['fingerprint'] != fingerprint:
            frappe.throw('Dữ liệu đã thay đổi trong lúc duyệt. Xem trước lại.')
        pending = frappe.db.get_value(RUN, {'company': setup.company(), 'month': month, 'state': 'Pending'}, 'name')
        run = frappe.get_doc(RUN, pending, for_update=True) if pending else frappe.new_doc(RUN)
        run.update(dict(company=setup.company(), month=month, state='Pending',
                        snapshot=_dump(result), preview_hash=fingerprint, active_key=setup.company()+'|'+month))
        run.flags.ignore_permissions = True
        run.save()
        booked = [book_row(row, run=run.name) for row in result['rows']]
        run.update(dict(state='Posted', results=_dump(booked), approved_by=frappe.session.user, approved_at=now_datetime()))
        run.save()
        for name in {r.get('allocation_name') for r in result['rows']} - {None}:
            allocation_doc = frappe.get_doc(ALLOCATION, name)
            if allocation_doc.month == month:
                allocation_doc.db_set('locked_by_run', run.name)
    return dict(name=run.name, total=result['total'], results=booked)


def ensure_settled(month):
    if not enabled():
        return
    _lock(setup.SETTINGS, setup.SETTINGS)
    if due_rows(month):
        frappe.throw('Còn khấu hao tài sản đến hạn chưa duyệt. Xử lý ở Tài sản cố định trước khi kết chuyển/khóa tháng.')
    _, end = _bounds(month)
    if frappe.db.exists(OP, {'company': setup.company(), 'state': 'Pending', 'posting_date': ['<=', end]}):
        frappe.throw('Còn nghiệp vụ tài sản chờ duyệt. Duyệt hoặc hủy trước khi khóa tháng.')


def _cancel_voucher(doctype, name):
    doc = frappe.get_doc(doctype, name)
    on_date = doc.get('posting_date') or doc.get('transaction_date') or doc.get('date')
    if on_date:
        check_period(on_date)
    doc.flags.ignore_permissions = True
    if doc.docstatus == 1:
        doc.cancel()


@frappe.whitelist(methods=['POST'])
def depreciation_reverse(name, reason):
    setup.require_chief()
    require_enabled()
    if len(str(reason or '').strip()) < 5:
        frappe.throw('Ghi rõ lý do đảo (ít nhất 5 ký tự).')
    with atomic(), internal():
        _lock(setup.SETTINGS, setup.SETTINGS)
        _lock(RUN, name)
        run = frappe.get_doc(RUN, name, for_update=True)
        if run.company != setup.company() or run.state != 'Posted':
            frappe.throw('Bảng khấu hao không thuộc công ty hoặc đã được đảo.')
        _, end = _bounds(run.month)
        check_period(end)
        results = _json(run.results, [])
        for row in reversed(results):
            _lock('Asset', row['asset'])
            if frappe.db.exists(RUN, {'company': run.company, 'month': ['>', run.month], 'state': 'Posted'}):
                frappe.throw('Đảo khấu hao các tháng sau trước.')
            if frappe.db.exists(OP, {'asset': row['asset'], 'state': 'Posted', 'kind': ['in', ['Upgrade', 'Profile', 'Sell', 'Scrap', 'Restore']],
                                     'approved_at': ['>', run.approved_at]}):
                frappe.throw('Đảo các nghiệp vụ thay đổi giá trị/chính sách tài sản đã duyệt sau bảng khấu hao trước.')
            if row.get('allocation'):
                _cancel_voucher('Journal Entry', row['allocation'])
            _cancel_voucher('Journal Entry', row['journal'])
        run.update(dict(state='Cancelled', active_key=None, reason=reason))
        run.save()
        for allocation_doc in frappe.get_all(ALLOCATION, filters={'locked_by_run': run.name}, pluck='name'):
            frappe.db.set_value(ALLOCATION, allocation_doc, 'locked_by_run', None)
        run.add_comment('Info', 'Đảo bảng khấu hao: ' + str(reason))
    return dict(name=name, state='Cancelled')


@frappe.whitelist()
def operation_preview(name):
    setup.require_view()
    from lotus_farm.finance.asset_operations import preview
    return preview(name)


@frappe.whitelist(methods=['POST'])
@retry_approval
def operation_post(name, fingerprint):
    setup.require_chief()
    require_enabled()
    from lotus_farm.finance.asset_operations import apply
    with atomic(), internal(operation=name):
        _lock(setup.SETTINGS, setup.SETTINGS)
        _lock(OP, name)
        return apply(name, fingerprint)


@frappe.whitelist(methods=['POST'])
def operation_cancel(name, reason):
    setup.require_post()
    require_enabled()
    from lotus_farm.finance.asset_operations import cancel
    with atomic(), internal(operation=name):
        _lock(setup.SETTINGS, setup.SETTINGS)
        _lock(OP, name)
        return cancel(name, reason)


@frappe.whitelist()
def report(on_date=None, farm=None):
    setup.require_view()
    on_date = str(getdate(on_date or nowdate()))
    company = setup.company()
    rows = []
    for entry in overview()['assets']:
        if entry.docstatus != 1 or str(entry.purchase_date) > on_date:
            continue
        doc = asset_doc(entry.name)
        depreciation = flt(frappe.db.sql("""select sum(r.debit_in_account_currency-r.credit_in_account_currency)
            from `tabJournal Entry Account` r join `tabJournal Entry` j on j.name=r.parent
            join `tabAccount` acc on acc.name=r.account
            where j.docstatus=1 and j.voucher_type='Depreciation Entry' and r.reference_type='Asset'
            and r.reference_name=%s and acc.root_type='Expense' and j.posting_date<=%s""", (doc.name, on_date))[0][0])
        history = frappe.get_all(OP, filters={'asset': doc.name, 'state': 'Posted', 'posting_date': ['<=', on_date]},
                                 fields=['kind', 'payload', 'posting_date', 'approved_at'], order_by='posting_date, approved_at')
        original = flt(doc.net_purchase_amount)
        historical = dict(entry)
        historical['lotus_usage_suspended'] = 0
        for event in history:
            payload = _json(event.payload, {})
            if event.kind == 'Create':
                original = flt(payload.get('amount'))
                historical.update({key: payload.get(key) for key in (*ar.DIMENSIONS, 'location', 'custodian', 'asset_name')})
            elif event.kind == 'Upgrade':
                original += flt(payload.get('amount'))
            elif event.kind == 'Move':
                historical.update({key: payload.get(key) for key in (*ar.DIMENSIONS, 'location', 'custodian')})
            elif event.kind == 'Profile':
                historical.update({key: payload.get(key) for key in ('asset_name', 'custodian')})
            elif event.kind in ('Suspend', 'Resume'):
                historical['lotus_usage_suspended'] = int(event.kind == 'Suspend')
        if farm and historical.get('lotus_farm') != farm:
            continue
        accumulated = flt(doc.opening_accumulated_depreciation) + depreciation
        disposed = bool(doc.disposal_date and str(doc.disposal_date) <= on_date and doc.status in ('Sold', 'Scrapped'))
        historical['status'] = doc.status if disposed else ('Fully Depreciated' if original-accumulated <= .5 else 'Partially Depreciated' if accumulated else 'Submitted')
        rows.append(dict(historical, total_asset_cost=original, depreciation=depreciation, accumulated=accumulated,
                         remaining=0 if disposed else max(0, original-accumulated), disposed=disposed))
    return dict(on_date=on_date, rows=rows, totals=dict(original=sum(flt(r['total_asset_cost']) for r in rows if not r['disposed']),
                accumulated=sum(r['accumulated'] for r in rows if not r['disposed']), remaining=sum(r['remaining'] for r in rows)),
                opening_reconciliation=opening_reconciliation())


PROFILE_COLUMNS = ('code', 'asset_name', 'asset_category', 'location', 'serial_no', 'purchase_date', 'available_for_use_date',
                   'amount', 'opening_depreciation', 'months', 'booked_months', 'salvage_value')
PROFILE_LABELS = ('Mã tài sản', 'Tên tài sản', 'Nhóm tài sản', 'Vị trí', 'Serial', 'Ngày mua', 'Ngày sử dụng',
                  'Nguyên giá', 'Hao mòn đầu kỳ', 'Tổng số tháng khấu hao', 'Số kỳ đã trích', 'Giá trị thu hồi')


def _download(book, filename):
    out = BytesIO()
    book.save(out)
    frappe.local.response.filename = filename
    frappe.local.response.filecontent = out.getvalue()
    frappe.local.response.type = 'download'


@frappe.whitelist()
def import_template():
    setup.require_view()
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = 'Tài sản đầu kỳ'
    book.active.append(list(PROFILE_LABELS))
    for index in range(1, len(PROFILE_LABELS)+1):
        from openpyxl.utils import get_column_letter
        book.active.column_dimensions[get_column_letter(index)].width = 24
    for column in ('F', 'G'):
        book.active.column_dimensions[column].number_format = 'dd/mm/yyyy'
    _download(book, 'tai-san-dau-ky.xlsx')


@frappe.whitelist(methods=['POST'])
def import_preview(content):
    setup.require_post()
    from openpyxl import load_workbook
    encoded = str(content).split(',', 1)[-1]
    if len(encoded) > 6_700_000:
        frappe.throw('File vượt quá 5 MB.')
    try:
        raw = base64.b64decode(encoded, validate=True)
    except ValueError:
        frappe.throw('File Excel không hợp lệ. Chọn lại file .xlsx từ mẫu.')
    if len(raw) > 5_000_000:
        frappe.throw('File vượt quá 5 MB.')
    from zipfile import BadZipFile, ZipFile
    try:
        with ZipFile(BytesIO(raw)) as zipped:
            if sum(f.file_size for f in zipped.infolist()) > 50_000_000:
                frappe.throw('File giải nén vượt quá 50 MB.')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError):
        frappe.throw('File Excel không hợp lệ. Chọn lại file .xlsx từ mẫu.')
    sheet = book.worksheets[0]
    values = list(sheet.iter_rows(max_row=2002, values_only=True))
    if not values or tuple(values[0])[:len(PROFILE_COLUMNS)] not in (PROFILE_COLUMNS, PROFILE_LABELS):
        frappe.throw('Dùng file mẫu tài sản đầu kỳ.')
    if len(values) > 2001:
        frappe.throw('Mỗi lần nhập tối đa 2000 tài sản.')
    rows, seen, serials = [], set(), set()
    pending_codes = {_json(r.payload, {}).get('code') for r in frappe.get_all(OP,
        filters={'company': setup.company(), 'kind': 'Create', 'state': 'Pending'}, fields=['payload'])}
    for index, cells in enumerate(values[1:], 2):
        if not any(v is not None for v in cells):
            continue
        item = dict(zip(PROFILE_COLUMNS, cells))
        for key in ('purchase_date', 'available_for_use_date'):
            if hasattr(item.get(key), 'date'):
                item[key] = item[key].date().isoformat()
        item['existing'] = True
        item['non_depreciable'] = bool(frappe.db.get_value('Asset Category', item.get('asset_category'), 'non_depreciable_category')) if item.get('asset_category') else False
        item['posting_date'] = str(setup.settings().cutover_date or nowdate())
        clean, errors = ar.profile(item)
        if clean['code'] in seen or clean['code'] in pending_codes or frappe.db.exists('Asset', {'company': setup.company(), 'lotus_asset_code': clean['code']}):
            errors['code'] = 'Mã tài sản đã có hoặc bị trùng trong file.'
        seen.add(clean['code'])
        if clean.get('serial_no') and (clean['serial_no'] in serials or frappe.db.exists('Asset', {'company': setup.company(), 'lotus_serial_no': clean['serial_no']})):
            errors['serial_no'] = 'Serial đã có trong danh mục.'
        serials.add(clean.get('serial_no'))
        for key, dt in [('asset_category', 'Asset Category'), ('location', 'Location')]:
            if clean.get(key) and not frappe.db.exists(dt, clean[key]):
                errors[key] = 'Mã không có trong danh mục.'
        rows.append(dict(line=index, values=clean, errors=errors))
    return dict(rows=rows, fingerprint=ar.fingerprint(rows), valid=all(not r['errors'] for r in rows) and bool(rows))


@frappe.whitelist(methods=['POST'])
@retry_approval
def import_commit(content, fingerprint):
    setup.require_post()
    require_enabled()
    with atomic():
        preview = import_preview(content)
        if preview['fingerprint'] != fingerprint or not preview['valid']:
            frappe.throw('File hoặc danh mục đã thay đổi/còn lỗi. Xem trước lại.')
        names = [save_asset(row['values'])['name'] for row in preview['rows']]
    return dict(operations=names, count=len(names))


@frappe.whitelist()
def export(on_date=None, farm=None):
    data = report(on_date, farm)
    from openpyxl import Workbook
    book = Workbook()
    ws = book.active
    ws.title = 'So tai san'
    ws.append(['Mã tài sản', 'Tên', 'Nhóm', 'Vườn', 'Lô', 'Vị trí', 'Nguyên giá', 'Khấu hao đã ghi', 'Hao mòn lũy kế', 'Giá trị còn lại', 'Trạng thái'])
    for row in data['rows']:
        ws.append([row.get(k) for k in ('lotus_asset_code', 'asset_name', 'asset_category', 'lotus_farm', 'lotus_farm_block', 'location',
                                        'total_asset_cost', 'depreciation', 'accumulated', 'remaining', 'status')])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    _download(book, 'so-tai-san-'+data['on_date']+'.xlsx')
