"""Dated purchase events and as-of debt, with one filter path for screen and full Excel export."""
from io import BytesIO
import frappe
from frappe.utils import flt, getdate, get_first_day, get_last_day, nowdate
from lotus_farm.procurement import access, document
from lotus_farm.procurement_rules import KIND_LABELS, kind_for
from lotus_farm.finance import setup

STAGES = dict(planned='Đề nghị', ordered='Đặt mua', received='Đã nhận/nghiệm thu', billed='Hóa đơn',
              supplier_paid='Trả NCC', company_cash='Tiền công ty chi', personal_funding='Cá nhân chi hộ',
              reimbursement='Hoàn người chi hộ', outstanding='Công nợ NCC', funding_outstanding='Còn phải hoàn')


def _scope(row, supplier, kind, farm):
    return (not supplier or row.get('supplier') == supplier) and (not kind or row.get('kind') == kind) and (not farm or row.get('farm') == farm)


def _parts(doc):
    """Native net values allocate header taxes to lines; line freight keeps its actual attribution."""
    total = flt(doc.get('base_grand_total') or doc.get('grand_total') or doc.get('total'))
    net = sum(flt(r.get('base_net_amount') or r.get('net_amount') or r.amount) for r in doc.items)
    freight = sum(flt(r.get('lotus_line_freight')) for r in doc.items)
    output = []
    for r in doc.items:
        amount = flt(r.get('base_net_amount') or r.get('net_amount') or r.amount)
        if r.get('lotus_line_vat'):
            tax_rate = 0 if r.lotus_line_vat == 'KCT' else flt(r.lotus_line_vat)/100
            gross = amount*(1+tax_rate)+flt(r.get('lotus_line_freight'))
        else:
            gross = amount+(total-net-freight)*(amount/net if net else 0)+flt(r.get('lotus_line_freight'))
        item = frappe.get_cached_doc('Item', r.item_code).as_dict() if r.item_code else {}
        output.append(dict(line=r.name, item=r.item_code, item_name=r.item_name, qty=flt(r.qty), uom=r.uom,
            amount=gross, kind=kind_for(item, r.get('lotus_purchase_kind')),
            farm=r.get('lotus_farm') or doc.get('lotus_farm') or '', supplier=doc.get('supplier') or ''))
    # Native rounding/discounts are part of the signed grand total.
    summed = sum(r['amount'] for r in output)
    if output and abs(summed-total)>.000001:
        weights = [abs(r['amount']) for r in output]
        divisor = sum(weights) or len(output)
        for r, weight in zip(output, weights):
            r['amount'] += (total-summed)*(weight/divisor if sum(weights) else 1/len(output))
    return output


@frappe.whitelist()
def report(company, month=None, as_of=None, farm=None, supplier=None, kind=None):
    access(company)
    if not setup.can_view() and not (set(frappe.get_roles()) & {'Lotus Purchase Approver', 'Lotus Farm Manager'}):
        frappe.throw('Báo cáo dòng tiền dành cho kế toán và người duyệt.', frappe.PermissionError)
    if kind and kind not in KIND_LABELS:
        frappe.throw('Loại mua không hợp lệ.')
    as_of = getdate(as_of or nowdate())
    month = month or str(as_of)[:7]
    start, end = get_first_day(month+'-01'), min(get_last_day(month+'-01'), as_of)
    if farm:
        f = frappe.get_doc('Farm', farm)
        f.check_permission('read')
        if f.company != company:
            frappe.throw('Vườn không thuộc công ty đang chọn.')
    if supplier:
        frappe.get_doc('Supplier', supplier).check_permission('read')
    events, totals, cache = [], {k: 0.0 for k in STAGES}, {}

    def emit(stage, dt, name, day, parts, amount=None):
        if getdate(day)>as_of or stage not in ('outstanding', 'funding_outstanding') and not start <= getdate(day) <= end:
            return
        denominator = sum(p['amount'] for p in parts)
        for p in parts:
            if not _scope(p, supplier, kind, farm):
                continue
            value = p['amount'] if amount is None else flt(amount)*(p['amount']/denominator if denominator else 1/len(parts))
            events.append(dict(p, stage=stage, stage_label=STAGES[stage], doctype=dt, name=name, date=str(day), amount=value))
            totals[stage] += value

    def parts(dt, name):
        key = (dt, name)
        if key not in cache:
            cache[key] = _parts(document(dt, name, company))
        return cache[key]

    for dt, stage, date_field in (('Material Request', 'planned', 'transaction_date'), ('Purchase Order', 'ordered', 'transaction_date'),
                                   ('Purchase Receipt', 'received', 'posting_date'), ('Purchase Invoice', 'billed', 'posting_date')):
        if not frappe.has_permission(dt, 'read'):
            continue
        filters = {'company': company, 'docstatus': 1, date_field: ['<=', as_of]}
        if dt == 'Material Request':
            filters['material_request_type'] = 'Purchase'
        for ref in frappe.get_list(dt, filters=filters, fields=['name', date_field], limit_page_length=0, order_by=date_field+', name'):
            emit(stage, dt, ref.name, ref[date_field], parts(dt, ref.name))
    if frappe.has_permission('Payment Entry', 'read'):
        for ref in frappe.get_list('Payment Entry', filters={'company': company, 'docstatus': 1, 'party_type': 'Supplier', 'posting_date': ['between', [start, end]]}, fields=['name'], limit_page_length=0):
            d = document('Payment Entry', ref.name, company)
            sign = -1 if d.payment_type == 'Receive' else 1
            allocated = 0
            for r in d.references:
                if r.reference_doctype not in ('Purchase Order', 'Purchase Invoice'):
                    continue
                p = parts(r.reference_doctype, r.reference_name)
                # Supplier allocations already carry the payable sign: a refund
                # against a credit note is negative. Do not invert it twice.
                amount = flt(r.allocated_amount)
                emit('supplier_paid', d.doctype, d.name, d.posting_date, p, amount)
                emit('company_cash', d.doctype, d.name, d.posting_date, p, amount)
                allocated += amount
            unallocated = sign*flt(d.paid_amount)-allocated
            if abs(unallocated)>.5:
                p = [dict(item='', item_name='Tiền NCC chưa phân bổ', line='', qty=0, uom='', kind='', farm='', supplier=d.party, amount=unallocated)]
                emit('supplier_paid', d.doctype, d.name, d.posting_date, p)
                emit('company_cash', d.doctype, d.name, d.posting_date, p)
    from lotus_farm.purchase_funding import DT, JE
    for ref in frappe.get_list(DT, filters={'company': company, 'posting_date': ['<=', as_of]}, fields=['name'], limit_page_length=0):
        d = document(DT, ref.name, company)
        if not d.journal_entry or frappe.db.get_value(JE, d.journal_entry, 'docstatus') != 1:
            continue
        p = []
        for a in d.allocations:
            source = parts(a.reference_type, a.reference_name)
            total = sum(r['amount'] for r in source)
            p.extend([dict(r, amount=flt(a.amount)*(r['amount']/total if total else 1/len(source))) for r in source])
        emit('personal_funding', DT, d.name, d.posting_date, p)
        emit('supplier_paid', DT, d.name, d.posting_date, p)
        reimbursed = 0
        for j in frappe.get_list(JE, filters={'lotus_purchase_funding': d.name, 'lotus_payment_kind': 'PurchaseReimbursement', 'docstatus': 1, 'posting_date': ['<=', as_of]}, fields=['name', 'posting_date', 'total_debit'], limit_page_length=0):
            reimbursed += flt(j.total_debit)
            emit('reimbursement', JE, j.name, j.posting_date, p, j.total_debit)
            emit('company_cash', JE, j.name, j.posting_date, p, j.total_debit)
        if d.classification == 'Reimbursable':
            emit('funding_outstanding', DT, d.name, as_of, p, flt(d.amount)-reimbursed)
    # The dated GL is authoritative. Include opening JE liabilities and unallocated advances;
    # invoices use their actual line proportions when filtering farm/type.
    account = setup.mapping().get('payable_supplier')
    if account:
        debt = frappe.db.sql('''select against_voucher_type dt, against_voucher_no name, party supplier, sum(amount) amount
            from `tabPayment Ledger Entry` where company=%s and account=%s and party_type='Supplier'
            and posting_date<=%s and delinked=0 group by against_voucher_type, against_voucher_no, party''', (company, account, as_of), as_dict=True)
        for r in debt:
            if abs(flt(r.amount))<.000001:
                continue
            if r.dt in ('Purchase Invoice', 'Purchase Order') and frappe.db.exists(r.dt, r.name):
                p = parts(r.dt, r.name)
            else:
                p = [dict(item='', item_name='Số dư / tạm ứng chưa phân bổ', line='', qty=0, uom='', kind='', farm='', supplier=r.supplier, amount=r.amount)]
            emit('outstanding', r.dt or 'Journal Entry', r.name, as_of, p, r.amount)
    events.sort(key=lambda r: (r['date'], r['stage'], r['name'], r['line']))
    return dict(rows=events, totals=totals, month=month, as_of=str(as_of), stages=STAGES,
        filters=dict(company=company, month=month, as_of=str(as_of), farm=farm, supplier=supplier, kind=kind),
        note='Các giai đoạn là chỉ tiêu riêng, không cộng ngang. Chi hộ trả NCC và hoàn người chi hộ được tách riêng. Chứng từ lịch sử chưa có hóa đơn/thanh toán không được suy ra từ ghi chú.')


@frappe.whitelist()
def export(company, month=None, as_of=None, farm=None, supplier=None, kind=None):
    data = report(company, month, as_of, farm, supplier, kind)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook()
    ws = book.active
    ws.title = 'Tổng hợp'
    ws.append(['BÁO CÁO MUA HÀNG', company])
    ws.append(['Tháng', data['month'], 'Chốt đến', data['as_of']])
    ws.append(['NCC', supplier or 'Tất cả', 'Loại', KIND_LABELS.get(kind, 'Tất cả'), 'Vườn', farm or 'Tất cả'])
    ws.append(['Chỉ tiêu', 'Số tiền VND'])
    for k, label in STAGES.items():
        ws.append([label, data['totals'][k]])
        ws.cell(ws.max_row, 2).number_format = '#,##0;[Red](#,##0)'
    ws.append([data['note']])
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 52, 30
    detail = book.create_sheet('Toàn bộ chi tiết')
    fields = ('date', 'stage_label', 'doctype', 'name', 'supplier', 'kind', 'farm', 'item', 'item_name', 'qty', 'uom', 'amount')
    detail.append(['Ngày', 'Giai đoạn', 'Loại chứng từ', 'Số chứng từ', 'Nhà cung cấp', 'Loại mua', 'Vườn', 'Mã hàng', 'Tên hàng', 'Số lượng', 'ĐVT', 'Số tiền VND'])
    for row in data['rows']:
        detail.append([KIND_LABELS.get(row.get(f), row.get(f)) if f == 'kind' else row.get(f) for f in fields])
        for col in (10, 12):
            detail.cell(detail.max_row, col).number_format = '#,##0.###;[Red](#,##0.###)'
    detail.freeze_panes = 'A2'
    detail.auto_filter.ref = detail.dimensions
    for col, width in zip('ABCDEFGHIJKL', (13, 25, 23, 28, 28, 20, 24, 24, 40, 16, 12, 22)):
        detail.column_dimensions[col].width = width
    for sheet, row in ((ws, 4), (detail, 1)):
        for cell in sheet[row]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='C43224')
            cell.alignment = Alignment(vertical='center')
    out = BytesIO()
    book.save(out)
    frappe.local.response.update(filename='mua-hang-'+data['month']+'.xlsx', filecontent=out.getvalue(), type='download')
