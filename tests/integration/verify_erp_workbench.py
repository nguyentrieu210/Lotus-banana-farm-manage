"""Independent ERP UI adapter: native controllers, posting, permissions and replay.

Only the disposable procurement.test site; importing base enforces that boundary.
"""
import json
import uuid
from io import BytesIO
import unittest
import frappe
from frappe.utils import flt
import verify_procurement as base
from lotus_farm import erp_workbench as erp


class Workbench(unittest.TestCase):
    setUp = base.ProcurementIntegration.setUp
    tearDown = base.ProcurementIntegration.tearDown
    order = base.ProcurementIntegration.order
    receive = base.ProcurementIntegration.receive

    @classmethod
    def setUpClass(cls):
        base.seed()
    def save(self, doc, **args):
        return erp.action(doc['doctype'], 'save', uuid.uuid4().hex, values=doc, company=base.C, **args)

    def test_supplier_crud_replay_and_concurrent_version(self):
        doc = erp.document('Supplier')['doc']
        doc.update(supplier_name='Own UI ' + self.tag, supplier_group=frappe.db.get_value('Supplier', self.supplier, 'supplier_group'), supplier_type='Company')
        token = uuid.uuid4().hex
        first = erp.action('Supplier', 'save', token, values=doc)
        retry = erp.action('Supplier', 'save', token, values=doc)
        self.assertEqual(first['doc']['name'], retry['doc']['name'])
        changed = {**doc, 'supplier_name': 'Different request'}
        with self.assertRaises(frappe.ValidationError):
            erp.action('Supplier', 'save', token, values=changed)
        updated = self.save({**first['doc'], 'mobile_no': '0900123456'}, name=first['doc']['name'], modified=first['doc']['modified'])
        self.assertEqual(updated['doc']['mobile_no'], '0900123456')
        with self.assertRaises(frappe.TimestampMismatchError):
            self.save(first['doc'], name=first['doc']['name'], modified=first['doc']['modified'])

    def test_new_order_keeps_lotus_approval_and_native_item_details(self):
        doc = erp.document('Purchase Order', company=base.C, warehouse=base.WH)['doc']
        self.assertEqual(doc['currency'], 'VND')
        doc.update(supplier=self.supplier, schedule_date=base.DAY, transaction_date=base.DAY, lotus_approval_status='Approved')
        details = erp.item_details('Purchase Order', 'PCT-M', base.C, row={'qty':2,'uom':'Nos'}, header=doc)
        doc['items'] = [dict(doctype='Purchase Order Item', name='new-'+uuid.uuid4().hex, __islocal=1, item_code='PCT-M', qty=2, rate=100000, uom='Nos', stock_uom='Nos', conversion_factor=1, schedule_date=base.DAY, warehouse=base.WH)]
        result = self.save(doc)
        from lotus_farm.purchasing import DRAFT
        self.assertEqual(result['doc']['lotus_approval_status'], DRAFT)
        self.assertFalse(result['doc']['items'][0]['name'].startswith('new-'))
        self.assertEqual(details['stock_uom'], 'Nos')
        with self.assertRaises(frappe.PermissionError):
            erp.action('Purchase Order','submit',uuid.uuid4().hex,name=result['doc']['name'],modified=result['doc']['modified'],company=base.C)

    def test_mapping_partial_receipt_and_invoice_all_four_kinds(self):
        for kind in ('Material','Tool','Service','Asset'):
            with self.subTest(kind=kind):
                self.tag = uuid.uuid4().hex[:8]
                order = self.order(kind,4)
                mapped = erp.map_document('Purchase Order',order.name,'Purchase Receipt',base.C)
                self.assertEqual(mapped['doc']['supplier'],self.supplier)
                self.assertTrue(all(row.get('purchase_order')==order.name for row in mapped['doc']['items']))
                receipt = self.receive(order,2)
                waiting = erp.map_document('Purchase Order',order.name,'Purchase Receipt',base.C)
                self.assertEqual(sum(flt(row['qty']) for row in waiting['doc']['items']),2)
                invoice = erp.map_document('Purchase Receipt',receipt.name,'Purchase Invoice',base.C)
                self.assertEqual(sum(flt(row['qty']) for row in invoice['doc']['items']),2)
                self.assertTrue(all(row.get('purchase_receipt')==receipt.name for row in invoice['doc']['items']))

    def test_receipt_requires_reason_but_surplus_needs_no_purchase_order(self):
        from lotus_farm import warehouse
        result = warehouse.create_draft(base.C,'receipt',base.DAY,[dict(item='PCT-M',qty=1,uom='Nos',rate=100000)],target=base.WH,receipt_reason='Surplus',client_event_id=self.tag)
        loaded = erp.document('Stock Entry',result['name'],base.C)
        doc = loaded['doc']; doc['lotus_receipt_reason']=''
        with self.assertRaisesRegex(frappe.ValidationError,'lý do nhập kho'):
            self.save(doc,name=doc['name'],modified=doc['modified'])
        posted = erp.action('Stock Entry','submit',uuid.uuid4().hex,name=loaded['doc']['name'],modified=loaded['doc']['modified'],company=base.C)
        self.assertEqual(posted['doc']['docstatus'],1)
        self.assertFalse(posted['doc'].get('purchase_order'))
        self.assertEqual(flt(frappe.db.get_value('GL Entry',dict(voucher_no=doc['name'],account=base.MAP['stock_adjustment']),'credit')),100000)

    def test_reports_equal_native_and_export_all_pages(self):
        from frappe.desk.query_report import run
        definition = erp.report_definition('Consolidated Financial Statement',base.C)
        self.assertEqual(definition['defaults']['presentation_currency'],'VND')
        payable=erp.report_definition('Accounts Payable',base.C)
        self.assertEqual(next(f for f in payable['filters'] if f['fieldname']=='party')['options'],'Supplier')
        self.assertTrue(erp.report_link_options('Accounts Payable','party','PCT',base.C))
        own = erp.run_report(definition['name'],definition['defaults'],base.C)
        native = run(report_name=definition['name'],filters=json.dumps(definition['defaults']),ignore_prepared_report=True,are_default_filters=False,js_filters=definition['filters'])
        self.assertEqual(own['report_summary'],native['report_summary'])
        first = erp.listing('Account',company=base.C,length=1)
        self.assertGreater(first['total_count'],100)
        erp.export_list('Account',company=base.C)
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(frappe.response.filecontent),read_only=True)
        self.assertEqual(workbook.active.max_row,first['total_count']+1)

    def test_stock_user_cannot_read_bank_data_or_post_purchase_order(self):
        order=self.order(qty=2)
        email='own-stock-'+self.tag+'@example.invalid'
        base.insert('User',email=email,first_name='Own stock',send_welcome_email=0,roles=[dict(role='Stock User')])
        frappe.set_user(email)
        self.assertTrue(erp.schema('Stock Entry')['permissions']['read'])
        receipt=erp.map_document('Purchase Order',order.name,'Purchase Receipt',base.C)
        self.assertTrue(receipt['is_new'])
        self.assertEqual(receipt['doc']['items'][0]['purchase_order'],order.name)
        with self.assertRaises(frappe.PermissionError):
            erp.schema('Journal Entry')
        with self.assertRaises(frappe.PermissionError):
            erp.document('Bank Account')
        with self.assertRaises(frappe.PermissionError):
            erp.document('Purchase Order',company=base.C)

    def test_tree_bank_fields_and_shared_supplier_company_scope(self):
        tree = erp.tree('Account',base.C)
        self.assertGreater(len(tree['rows']),100)
        names = {r['name'] for r in tree['rows']}
        self.assertTrue(all(not r['_erp_parent'] or r['_erp_parent'] in names for r in tree['rows']))
        transactions = erp.supplier_connections(self.supplier,base.C)
        self.assertTrue(all(r['filters']['company']==base.C for r in transactions))
        doc = erp.document('Supplier')['doc']
        doc.update(supplier_name=frappe.db.get_value('Supplier',self.supplier,'supplier_name'),supplier_group=frappe.db.get_value('Supplier',self.supplier,'supplier_group'),supplier_type='Company')
        with self.assertRaisesRegex(frappe.ValidationError,'đã tồn tại'):
            self.save(doc)
        frappe.set_user(base.BUYER)
        self.assertFalse(any('bank' in f['fieldname'] for f in erp.schema('Supplier')['fields']))
        with self.assertRaises(frappe.PermissionError):
            erp.link_options('Supplier','default_bank_account',company=base.C)

    def test_purchase_request_cannot_bypass_named_signature_gate(self):
        doc = erp.document('Material Request',company=base.C,warehouse=base.WH)['doc']
        doc.update(material_request_type='Purchase',schedule_date=base.DAY,transaction_date=base.DAY,lotus_request_note='Đề nghị kiểm thử',items=[dict(item_code='PCT-M',qty=2,uom='Nos',stock_uom='Nos',conversion_factor=1,schedule_date=base.DAY,warehouse=base.WH)])
        saved = self.save(doc)
        with self.assertRaises(frappe.PermissionError):
            erp.action('Material Request','submit',uuid.uuid4().hex,name=saved['doc']['name'],modified=saved['doc']['modified'],company=base.C)
        self.assertEqual(erp.approval('Material Request',saved['doc']['name'])['chain'],'purchase_proposal')

    def test_own_adapter_posts_partial_receipts_and_invoices_for_four_kinds(self):
        for kind in ('Material','Tool','Service','Asset'):
            with self.subTest(kind=kind):
                self.tag=uuid.uuid4().hex[:8]
                order=self.order(kind,4,vat='10',freight=20000)
                received=erp.map_document('Purchase Order',order.name,'Purchase Receipt',base.C)['doc']
                received['posting_date']=base.DAY
                for row in received['items']:
                    row['qty']=2;row['received_qty']=2;row['stock_qty']=2
                draft=self.save(received)
                posted=erp.action('Purchase Receipt','submit',uuid.uuid4().hex,name=draft['doc']['name'],modified=draft['doc']['modified'],company=base.C)
                self.assertEqual(posted['doc']['docstatus'],1)
                waiting=erp.map_document('Purchase Order',order.name,'Purchase Receipt',base.C)
                self.assertEqual(sum(flt(r['qty']) for r in waiting['doc']['items']),2)
                invoice=erp.map_document('Purchase Receipt',posted['doc']['name'],'Purchase Invoice',base.C)['doc']
                invoice.update(bill_no='OWN-'+self.tag,bill_date=base.DAY,posting_date=base.DAY,due_date=base.DAY)
                invoice=self.save(invoice)
                billed=erp.action('Purchase Invoice','submit',uuid.uuid4().hex,name=invoice['doc']['name'],modified=invoice['doc']['modified'],company=base.C)
                self.assertEqual(billed['doc']['docstatus'],1)
                self.assertEqual(sum(flt(r['qty']) for r in billed['doc']['items']),2)
                self.assertEqual(flt(frappe.db.sql('select sum(debit-credit) from `tabGL Entry` where voucher_no=%s and is_cancelled=0',billed['doc']['name'])[0][0]),0)


if __name__=='__main__':
    try: unittest.main(verbosity=2)
    finally: frappe.db.rollback(); frappe.destroy()
