"""Device punch -> proposal -> approved attendance -> month close -> payroll sheet, on a real site.

Run inside the backend container with the bench Python. Every record rolls back at the end.
"""
import io
import json
import uuid

import frappe

frappe.init(site='frontend', sites_path='/home/frappe/frappe-bench/sites')
frappe.connect()
frappe.set_user('Administrator')
from frappe.utils import add_days, getdate

from lotus_farm import hr_operations as ops
from lotus_farm import hr_period as hp
from lotus_farm import timekeeping as t

company = 'Lotus Agritech'
month = '2026-09'
report = {}
try:
    tag = uuid.uuid4().hex[:8]
    hp.setup_calendar(company, 2026, ['Sunday'], [dict(date='2026-09-01', description='Liền kề Quốc khánh')])
    days = {h['date'] for h in hp.calendar_info(company, 2026)['holidays']}
    assert {'2026-09-01', '2026-09-02', '2026-09-06'} <= days, 'calendar misses holidays'
    shift = t.create_shift('QA ' + tag, '07:00', '16:00', 10, 10, 6, 2, company)['name']
    assert frappe.db.get_value('Shift Type', shift, 'holiday_list'), 'shift without holiday list'

    employee = t.create_employee(company, dict(employee_name='Payroll QA ' + tag, gender='Male', date_of_birth='1990-01-01',
                                               date_of_joining='2026-08-01', designation='QA ' + tag))['name']
    frappe.db.set_value('Employee', employee, dict(lotus_payment_basis='Daywork', lotus_daily_rate=300000))
    farm = t.create_farm(company, 'PAY-QA-' + tag, 'Payroll QA ' + tag)['name']
    device = ops.save_device(company, dict(farm=farm, device_code='QA-' + tag, device_name='QA ' + tag))['name']
    frappe.get_doc(dict(doctype='Lotus Attendance Identity', company=company, farm=farm, device=device, device_user_id='77' + tag[:4],
                        employee=employee, valid_from='2026-08-01', active=1)).insert()
    ops.ingest_device_events(company, device, [dict(device_user_id='77' + tag[:4], punched_at='2026-09-03 07:05:00', log_type='IN'),
                                              dict(device_user_id='77' + tag[:4], punched_at='2026-09-03 17:00:00', log_type='OUT')])
    synced = ops.sync_ready_checkins(company, '2026-09-03', '2026-09-03')
    assert synced['synced'] == 2, synced

    proposal = next(r for r in t.propose_from_checkins(company, '2026-09-03', shift)['rows'] if r['employee'] == employee)
    assert (proposal['status'], proposal['late_entry'], proposal['overtime_hours']) == ('Present', 0, 1.0), proposal
    saved = t.save_staff(company, '2026-09-03', [dict(employee=employee, status='Present', working_hours=proposal['working_hours'],
                                                      overtime_hours=proposal['overtime_hours'], in_time=proposal['in_time'],
                                                      out_time=proposal['out_time'])], 'pay-' + tag, shift=shift)
    doc = frappe.get_doc('Attendance', saved['names'][0])
    t.staff_action(company, [dict(name=doc.name, modified=str(doc.modified))], 'submit')

    before = hp.period_status(company, month)
    assert 'missing_attendance' in [b['code'] for b in before['blockers']], before
    # Fill every remaining working day for everyone active in September, plus one Sunday for the QA employee.
    begin, end = getdate('2026-09-01'), getdate('2026-09-30')
    gaps = next(b for b in before['blockers'] if b['code'] == 'missing_attendance')['employees']
    for person in gaps:
        for day in person['days']:
            frappe.get_doc(dict(doctype='Attendance', company=company, employee=person['name'], attendance_date=day, status='Present')).insert().submit()
    frappe.get_doc(dict(doctype='Attendance', company=company, employee=employee, attendance_date='2026-09-06', status='Present')).insert().submit()
    assert hp.period_status(company, month)['blockers'] == [], hp.period_status(company, month)

    hp.close_period(company, month)
    assert t.monthly(company, month)['period_closed']
    doc.reload()
    try:
        doc.cancel()
    except frappe.ValidationError as exc:
        assert 'khóa công' in str(exc), exc
    else:
        raise AssertionError('closed month accepted a cancel')

    sheet = hp.payroll_sheet(company, month)
    row = next(e for e in sheet['employees'] if e['employee'] == employee)
    assert sheet['draft'] is False and sheet['standard_days'] == 24, (sheet['draft'], sheet['standard_days'])
    assert (row['workdays'], row['weekly_off_days'], row['overtime_hours']) == (24, 1, 1.0), row
    assert row['total'] == 24 * 300000 + 600000 + 56250, row

    hp.export_payroll(company, month)
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(frappe.local.response.filecontent))
    assert book.sheetnames == ['Nhân sự', 'Đội khoán', 'Kiểm tra'], book.sheetnames
    exported = next(r for r in book['Nhân sự'].iter_rows(min_row=4, values_only=True) if r[0] == employee)
    assert exported[20] == row['total'], exported

    hp.reopen_period(company, month, 'QA kiểm tra mở lại')
    doc.reload()
    doc.cancel()
    report = dict(calendar='pass', device_sync=synced['synced'], proposal='pass', close_guard='pass',
                  standard_days=sheet['standard_days'], total=row['total'], xlsx='pass', reopen='pass', rollback=True)
finally:
    frappe.set_user('Administrator')
    frappe.db.rollback()
    # Company default holiday list and doc caches were touched inside the rolled-back transaction.
    frappe.clear_document_cache('Company', company)
    frappe.clear_cache()
print(json.dumps(report, ensure_ascii=False))
