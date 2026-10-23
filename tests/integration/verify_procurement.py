"""Native purchase accounting tests. Hard refusal on production; artificial fixtures only.

LOTUS_PROCUREMENT_TEST_ONLY=1 env/bin/python ../apps/lotus_farm/tests/integration/verify_procurement.py
"""
import json
import os
from io import BytesIO
import unittest
from unittest.mock import patch
import uuid
import frappe
from frappe.utils import flt, nowdate, add_days

assert os.environ.get('LOTUS_PROCUREMENT_TEST_ONLY') == '1', 'Use the disposable procurement test stack'
os.chdir('/home/frappe/frappe-bench/sites')
frappe.init(site='procurement.test')
frappe.connect()
assert frappe.local.site == 'procurement.test' and frappe.conf.db_host == 'db'
frappe.set_user('Administrator')
frappe.flags.mute_emails = True
from lotus_farm import procurement as p, purchasing, procurement_import as imports, procurement_reports as reports, purchase_funding as funding
from lotus_farm.finance import setup, api as finance_api, payments, asset_operations
from lotus_farm import warehouse_buying_lite as buying, warehouse_planning as receiving, signing

C = 'Lotus Procurement Test'
DAY = nowdate()


def insert(dt, **values):
    return frappe.get_doc(dict(doctype=dt, **values)).insert(ignore_permissions=True,
        ignore_mandatory=dt=='Department' and values.get('department_name')=='All Departments')


def user(key, roles):
    email = 'purchase-'+key+'@example.invalid'
    if not frappe.db.exists('User', email):
        u = insert('User', email=email, first_name=key, send_welcome_email=0, roles=[dict(role=r) for r in roles])
    if not frappe.db.exists('Employee',{'user_id':email}):
        insert('Employee', first_name=key, gender='Male', date_of_birth='1990-01-01', date_of_joining='2026-01-01', status='Active', company=C, user_id=email)
    return email


def seed():
    if not frappe.get_meta('Lotus Purchase Funding').has_field('cancelled'):
        frappe.reload_doc('lotus_farm','doctype','lotus_purchase_funding')
    from lotus_farm.procurement_setup import install
    install()
    for name,account_type in (('Supplier','Payable'),('Customer','Receivable')):
        if not frappe.db.exists('Party Type',name):
            insert('Party Type',party_type=name,account_type=account_type)
    for dt, key, value in (('Department','department_name','All Departments'),('Item Group','item_group_name','All Item Groups'),('Supplier Group','supplier_group_name','All Supplier Groups')):
        if not frappe.db.exists(dt,value):
            insert(dt,**{key:value,'is_group':1})
    if not frappe.db.exists('Gender','Male'):
        insert('Gender',gender='Male')
    if not frappe.db.exists('Fiscal Year', '2026'):
        insert('Fiscal Year', year='2026', year_start_date='2026-01-01', year_end_date='2026-12-31')
    if not frappe.db.exists('Warehouse Type', 'Transit'):
        insert('Warehouse Type', name='Transit', warehouse_type_name='Transit')
    if not frappe.db.exists('Company', C):
        insert('Company', company_name=C, abbr='PCT', default_currency='VND', country='Vietnam',
            create_chart_of_accounts_based_on='Standard Template', chart_of_accounts='Standard')
    frappe.defaults.set_global_default('company', C)
    cfg = setup.settings()
    cfg.company=C
    cfg.flags.ignore_permissions=True
    cfg.save()
    if not cfg.coa_version:
        setup.coa_apply(on_date='2026-09-01')
    finance_api.save_settings(dict(enabled=1, cutover_date='2026-09-01'))
    global ACC, CHIEF, DIR, BUYER, LEADER, WH, MAP
    ACC=user('acc',['Lotus Accountant'])
    CHIEF=user('chief',['Lotus Chief Accountant'])
    DIR=user('director',['Lotus Purchase Approver'])
    BUYER=user('buyer',['Lotus Purchasing Buyer'])
    LEADER=user('leader',['Lotus Team Leader'])
    WH=frappe.db.get_value('Warehouse', {'company':C,'is_group':0,'disabled':0}, 'name')
    MAP=setup.mapping()
    if not frappe.db.exists('Farm',{'farm_code':'PCT-FARM'}):
        insert('Farm',farm_code='PCT-FARM',farm_name='Vườn kiểm thử mua hàng',company=C,area_ha=10,active=1)
    for key, root in (('purchase_funding_liability','Liability'),('purchase_funding_equity','Equity')):
        name=frappe.db.get_value('Account',{'company':C,'account_name':key},'name')
        if not name:
            parent=frappe.db.get_value('Account',{'company':C,'root_type':root,'is_group':1},'name')
            name=insert('Account',account_name=key,parent_account=parent,root_type=root,company=C).name
        frappe.db.set_single_value(setup.SETTINGS,'acct_'+key,name)
    if not frappe.db.exists('Asset Category','Procurement Asset'):
        from lotus_farm.finance.assets import save_category
        by_number={r.account_number:r.name for r in frappe.get_all('Account',filters={'company':C},fields=['name','account_number'])}
        save_category(dict(asset_category_name='Procurement Asset',fixed_asset_account=by_number['211'],
            accumulated_depreciation_account=by_number['2141'],depreciation_expense_account=by_number['6274'],capital_work_in_progress_account=by_number['2411']))
    if not frappe.db.exists('UOM','Nos'):
        insert('UOM',uom_name='Nos',must_be_whole_number=1)
    group=frappe.db.get_value('Item Group',{'is_group':0},'name')
    if not group:
        group=insert('Item Group',item_group_name='Procurement Test',parent_item_group='All Item Groups').name
    for code, kind in (('PCT-M','Material'),('PCT-T','Tool'),('PCT-A','Asset'),('PCT-S','Service')):
        if not frappe.db.exists('Item',code):
            insert('Item',item_code=code,item_name=code,item_group=group,stock_uom='Nos',is_stock_item=int(kind in ('Material','Tool')),
                is_fixed_asset=int(kind=='Asset'),asset_category='Procurement Asset' if kind=='Asset' else None,
                auto_create_assets=0,is_purchase_item=1,valuation_rate=100000,
                item_defaults=[dict(company=C,default_warehouse=WH,expense_account=MAP['admin_expense'])])
    if not frappe.db.exists('Supplier', 'PCT Supplier'):
        sg=frappe.db.get_value('Supplier Group', {'is_group':0}, 'name') or insert('Supplier Group',supplier_group_name='PCT Suppliers',parent_supplier_group='All Supplier Groups').name
        insert('Supplier',supplier_name='PCT Supplier',supplier_group=sg,supplier_type='Company')
    from lotus_farm.finance import assets
    assets.activate(1)
    # Source-only test files and all notifications stay inside this disposable stack.
    frappe.db.commit()


def as_user(who, fn, *args, **kwargs):
    frappe.set_user(who)
    try:
        return fn(*args, **kwargs)
    finally:
        frappe.set_user('Administrator')


def signed(voucher):
    as_user(CHIEF,payments.sign_payment,voucher['doctype'],voucher['voucher'])
    as_user(DIR,payments.sign_payment,voucher['doctype'],voucher['voucher'])
    return frappe.get_doc(voucher['doctype'],voucher['voucher'])


class ProcurementIntegration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        seed()

    def setUp(self):
        frappe.set_user('Administrator')
        frappe.db.savepoint('purchase_case')
        self.notify=patch.object(signing,'_notify',return_value=None)
        self.notify.start()
        self.tag=uuid.uuid4().hex[:8]
        sg=frappe.db.get_value('Supplier','PCT Supplier','supplier_group')
        self.supplier=insert('Supplier',supplier_name='PCT Case '+self.tag,supplier_group=sg,supplier_type='Company').name

    def tearDown(self):
        frappe.set_user('Administrator')
        frappe.db.rollback(save_point='purchase_case')
        self.notify.stop()

    def order(self, kind='Material', qty=10, rate=100000, vat='KCT', freight=0):
        data=dict(supplier=getattr(self,'supplier','PCT Supplier'),warehouse=WH,expected_date=DAY,mode='Kho tự mua',items=[dict(item={'Material':'PCT-M','Tool':'PCT-T','Asset':'PCT-A','Service':'PCT-S'}[kind],qty=qty,rate=rate,uom='Nos',kind=kind,vat_code=vat,freight=freight)])
        result=as_user(BUYER,buying.save_order,C,data)
        d=frappe.get_doc('Purchase Order',result['name'])
        as_user(BUYER,purchasing.request_approval,d.name,str(d.modified),DIR)
        d.reload()
        as_user(DIR,purchasing.approve_order,d.name,str(d.modified))
        return frappe.get_doc('Purchase Order',d.name)

    def receive(self, order, qty):
        result=as_user('Administrator',receiving.receive_order,order.name,DAY,json.dumps({order.items[0].name:qty}),self.tag+str(qty))
        d=frappe.get_doc('Purchase Receipt',result['name'])
        receiving.receipt_action(d.name,str(d.modified),'submit')
        return frappe.get_doc('Purchase Receipt',d.name)

    def invoice(self, receipt):
        actor=CHIEF if any(frappe.get_cached_value('Item',r.item_code,'is_fixed_asset') for r in receipt.items) else ACC
        result=as_user(actor,p.invoice_receipt,receipt.name,'HD-'+self.tag,DAY,DAY,self.tag)
        return frappe.get_doc('Purchase Invoice',result['invoice'])

    def fund(self, invoice, amount=400000, classification='Reimbursable'):
        f=as_user(ACC,insert,'File',file_name='receipt-'+self.tag+'.txt',is_private=1,content='Evidence '+self.tag)
        employee=frappe.db.get_value('Employee',{'user_id':BUYER},'name')
        data=dict(supplier=invoice.supplier,posting_date=DAY,payer_type='Employee',payer=employee,classification=classification,
            amount=amount,evidence=f.file_url,external_reference='TX-'+self.tag,notes='Test evidence confirmed',
            allocations=[dict(reference_type=invoice.doctype,reference_name=invoice.name,amount=amount)])
        result=as_user(ACC,funding.prepare,C,data,{'CHIEF':CHIEF,'APPROVE':DIR},self.tag)
        return result,data

    def test_01_all_kinds_partial_receipt_and_no_duplicate_stock(self):
        for kind in ('Material','Tool','Service','Asset'):
            with self.subTest(kind=kind):
                self.tag=uuid.uuid4().hex[:8]
                order=self.order(kind,4)
                receipt=self.receive(order,2)
                self.assertEqual(sum(flt(r.qty) for r in receipt.items),2)
                self.assertEqual(len(receipt.items),2 if kind=='Asset' else 1)
                count=frappe.db.count('Stock Ledger Entry',{'voucher_no':receipt.name,'is_cancelled':0})
                self.assertEqual(count,1 if kind in ('Material','Tool') else 0)
                repeat=receiving.receive_order(order.name,DAY,json.dumps({order.items[0].name:2}),self.tag+'2')
                self.assertEqual(repeat['name'],receipt.name)
                self.assertEqual(flt(frappe.db.get_value('Purchase Order Item',order.items[0].name,'received_qty')),2)

    def test_02_mixed_vat_freight_invoice_and_company_payment(self):
        order=self.order(vat='10',freight=50000)
        receipt=self.receive(order,10)
        inv=self.invoice(receipt)
        self.assertEqual(flt(inv.grand_total),1150000)
        v=as_user(ACC,p.supplier_payment,C,'Purchase Invoice',inv.name,300000,'cash',DAY,{'CHIEF':CHIEF,'APPROVE':DIR},self.tag)
        signed(v)
        self.assertEqual(flt(frappe.db.get_value('Purchase Invoice',inv.name,'outstanding_amount')),850000)
        self.assertEqual(flt(frappe.db.sql('select sum(debit-credit) from `tabGL Entry` where voucher_no=%s and is_cancelled=0',v['voucher'])[0][0]),0)

    def test_03_personal_funding_partial_reimbursement_and_equity_guard(self):
        inv=self.invoice(self.receive(self.order(),10))
        result,data=self.fund(inv)
        duplicate=as_user(ACC,funding.prepare,C,data,{'CHIEF':CHIEF,'APPROVE':DIR},self.tag)
        self.assertEqual(duplicate['name'],result['name'])
        signed(result)
        self.assertEqual(flt(frappe.db.get_value('Purchase Invoice',inv.name,'outstanding_amount')),600000)
        v=as_user(ACC,funding.reimburse,result['name'],150000,DAY,'cash',{'CHIEF':CHIEF,'APPROVE':DIR},self.tag+'R')
        with self.assertRaises(frappe.ValidationError):
            as_user(ACC,funding.reimburse,result['name'],300000,DAY,'cash',{'CHIEF':CHIEF,'APPROVE':DIR},self.tag+'over')
        signed(v)
        balance=funding.reimbursement_balance(frappe.get_doc(funding.DT,result['name']))
        self.assertEqual(balance['remaining'],250000)
        report=reports.report(C,DAY[:7],DAY,supplier=self.supplier)
        self.assertEqual(report['totals']['supplier_paid'],400000)
        self.assertEqual(report['totals']['company_cash'],150000)
        self.assertEqual(report['totals']['outstanding'],600000)
        self.assertEqual(report['totals']['funding_outstanding'],250000)
        self.tag=uuid.uuid4().hex[:8]
        equity,_=self.fund(inv,200000,'Equity')
        signed(equity)
        with self.assertRaises(frappe.ValidationError):
            as_user(ACC,funding.reimburse,equity['name'],100000,DAY,'cash',{'CHIEF':CHIEF,'APPROVE':DIR},self.tag+'E')

    def test_04_advance_reconciles_to_invoice(self):
        order=self.order('Service')
        advance=as_user(ACC,p.supplier_payment,C,'Purchase Order',order.name,300000,'cash',DAY,{'CHIEF':CHIEF,'APPROVE':DIR},self.tag+'A')
        signed(advance)
        inv=self.invoice(self.receive(order,10))
        self.assertEqual(flt(inv.outstanding_amount),700000)
        self.assertEqual(frappe.db.count('Stock Ledger Entry',{'voucher_no':inv.name}),0)

    def test_05_stale_signatures_and_native_funding_tamper_refused(self):
        inv=self.invoice(self.receive(self.order(),10))
        result,_=self.fund(inv)
        j=frappe.get_doc('Journal Entry',result['voucher'])
        j.accounts[0].reference_name='forged'
        with self.assertRaises(frappe.ValidationError):
            j.submit()
        f=frappe.get_doc(funding.DT,result['name'])
        f.amount=1
        with self.assertRaises(frappe.ValidationError):
            f.save()
        with self.assertRaises(frappe.PermissionError):
            as_user(LEADER,p.supplier_payment,C,'Purchase Invoice',inv.name,1,'cash',DAY,{},'unauthorized')

    def test_06_import_months_drafts_and_stale_preview(self):
        from openpyxl import Workbook
        book=Workbook();book.active.title='Nguyên liệu';book.active.append(['type','item','qty','uom','rate','supplier','month']);book.active.append(['Need','PCT-M',4,'Nos',100000,'','2026-10']);book.active.append(['Need','PCT-T',2,'Nos',200000,'','2026-11'])
        out=BytesIO();book.save(out)
        f=insert('File',file_name='source.xlsx',is_private=1,content=out.getvalue())
        preview=imports.preview(C,f.file_url)
        mappings={r['key']:{'warehouse':WH} for r in preview['rows']}
        reviewed=imports.preview(C,f.file_url,mappings)
        self.assertTrue(reviewed['valid'],reviewed)
        result=imports.commit(C,f.file_url,mappings,reviewed['fingerprint'],self.tag)
        self.assertEqual({d['month'] for d in result['documents']},{'2026-10','2026-11'})
        self.assertTrue(all(d['docstatus']==0 for d in result['documents']))
        self.assertEqual(imports.commit(C,f.file_url,mappings,reviewed['fingerprint'],self.tag)['name'],result['name'])
        with self.assertRaises(frappe.ValidationError):
            imports.commit(C,f.file_url,mappings,'stale',self.tag+'bad')

    def test_07_quote_missing_price_not_paid_and_split_order(self):
        data=dict(supplier='PCT Supplier',warehouse=WH,expected_date=DAY,valid_till=add_days(DAY,7),items=[dict(item='PCT-M',qty=10,rate=0,uom='Nos')])
        result=as_user(BUYER,p.save_quotation,C,data,self.tag)
        row=p.quotations(C)[0]
        self.assertIsNone(row['gross'])
        self.assertEqual(frappe.db.count('Payment Entry Reference',{'reference_name':result['name']}),0)
        q=frappe.get_doc('Supplier Quotation',result['name']);q.items[0].rate=100000;q.save();q.submit()
        row=p.quotations(C)[0]
        order=as_user(BUYER,p.select_quotes,C,[dict(quotation=q.name,line=row['name'],modified=str(q.modified),qty=4,reason='Giao đúng hạn')],self.tag+'Q')
        d=frappe.get_doc('Purchase Order',order['orders'][0])
        self.assertEqual(d.docstatus,0);self.assertEqual(d.items[0].supplier_quotation,q.name)

    def test_08_returns_and_cancel_restore_canonical_quantity(self):
        order=self.order()
        receipt=self.receive(order,10)
        result=p.return_receipt(receipt.name,DAY,{receipt.items[0].name:3},self.tag+'return')
        ret=frappe.get_doc('Purchase Receipt',result['name'])
        receiving.receipt_action(ret.name,str(ret.modified),'submit')
        self.assertEqual(flt(frappe.db.get_value('Purchase Order Item',order.items[0].name,'received_qty')),7)
        ret.reload()
        receiving.receipt_action(ret.name,str(ret.modified),'cancel')
        self.assertEqual(flt(frappe.db.get_value('Purchase Order Item',order.items[0].name,'received_qty')),10)
        receipt.reload()
        receiving.receipt_action(receipt.name,str(receipt.modified),'cancel')
        self.assertEqual(flt(frappe.db.get_value('Purchase Order Item',order.items[0].name,'received_qty')),0)

    def test_09_asof_debt_does_not_use_current_invoice_balance(self):
        inv=self.invoice(self.receive(self.order(),10))
        tomorrow=add_days(DAY,1)
        result=as_user(ACC,p.supplier_payment,C,'Purchase Invoice',inv.name,300000,'cash',tomorrow,{'CHIEF':CHIEF,'APPROVE':DIR},self.tag+'future')
        signed(result)
        self.assertEqual(flt(frappe.db.get_value('Purchase Invoice',inv.name,'outstanding_amount')),700000)
        self.assertEqual(reports.report(C,DAY[:7],DAY,supplier=self.supplier)['totals']['outstanding'],1000000)
        self.assertEqual(reports.report(C,tomorrow[:7],tomorrow,supplier=self.supplier)['totals']['outstanding'],700000)

    def test_10_asset_invoice_rows_link_each_physical_piece(self):
        order=self.order('Asset',2,15000000)
        receipt=self.receive(order,2)
        inv=self.invoice(receipt)
        self.assertEqual(len(inv.items),2)
        self.assertEqual({r.pr_detail for r in inv.items},{r.name for r in receipt.items})
        self.assertEqual(len({r.name for r in inv.items}),2)
        self.assertTrue(all(flt(r.qty)==1 for r in inv.items))
        self.assertEqual(frappe.db.count('Asset',{'purchase_invoice':inv.name}),0)
        from lotus_farm.finance import assets
        location=assets.save_location('Vị trí thử '+self.tag)['name']
        values=dict(code='PCT-ASSET-'+self.tag,asset_name='Tài sản thử '+self.tag,asset_category='Procurement Asset',
            location=location,existing=False,purchase_date=DAY,available_for_use_date=DAY,posting_date=DAY,amount=15000000,
            months=24,opening_depreciation=0,booked_months=0,salvage_value=0,item_code='PCT-A',purchase_invoice=inv.name,
            purchase_row=inv.items[0].name)
        pending=as_user(ACC,assets.save_asset,values)
        preview=assets.operation_preview(pending['name'])
        self.assertFalse(preview['errors'],preview)
        as_user(CHIEF,assets.operation_post,pending['name'],preview['fingerprint'])
        self.assertEqual(frappe.db.count('Asset',{'purchase_invoice':inv.name,'docstatus':1}),1)
        duplicate=as_user(ACC,assets.save_asset,dict(values,code='PCT-DUP-'+self.tag))
        self.assertTrue(assets.operation_preview(duplicate['name'])['errors'])

    def test_11_personal_advance_on_order_reconciles_when_billed(self):
        order=self.order('Service')
        result,_=self.fund(order,200000)
        signed(result)
        inv=self.invoice(self.receive(order,10))
        self.assertEqual(flt(inv.outstanding_amount),800000)

    def test_12_source_excel_pack_units_and_months(self):
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.title='Phân bón'
        sheet.append(['DANH MỤC PHÂN BÓN CẦN MUA THÁNG 10'])
        sheet.append(['Stt','Tên phân bón','Tên thương mại (SP giá tốt nhất)','ĐVT','Quy cách (ĐVT/bao-chai)','SL cần mua (bao/chai)','Giá tốt nhất (đ/ĐVT)'])
        sheet.append([1,'Vi lượng','HP','lít',20,33,82000])
        second=book.create_sheet('Vật tư');second.append(['DANH MỤC VẬT TƯ CẦN MUA THÁNG 11']);second.append(['Stt','Tên vật tư','Tên thương mại (SP giá tốt nhất)','ĐVT','Quy cách (ĐVT/bao-chai)','SL cần mua (bao/chai)','Giá tốt nhất (đ/ĐVT)']);second.append([1,'Bao cám','-','cái',1,20073,0])
        out=BytesIO();book.save(out)
        rows=imports.workbook_rows(out.getvalue())
        self.assertEqual(rows[0]['qty'],660)
        self.assertEqual(rows[0]['rate'],82000)
        self.assertEqual(rows[0]['month'],'2026-10')
        self.assertEqual(rows[1]['month'],'2026-11')
        self.assertEqual(rows[1]['item_label'],'Bao cám')

    def test_13_export_contains_every_filtered_event(self):
        inv=self.invoice(self.receive(self.order(),10))
        data=reports.report(C,DAY[:7],DAY,kind='Material')
        reports.export(C,DAY[:7],DAY,kind='Material')
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(frappe.local.response.filecontent),data_only=True)
        self.assertEqual(book['Toàn bộ chi tiết'].max_row,len(data['rows'])+1)
        self.assertTrue(any(r['name']==inv.name and r['stage']=='billed' for r in data['rows']))
        self.assertEqual(reports.report(C,DAY[:7],DAY,kind='Service')['totals']['billed'],0)

    def test_14_buyer_catalogue_and_legacy_header_vat_invoice(self):
        with patch('frappe.sessions.get_csrf_token',return_value='test-csrf'):
            info=as_user(BUYER,p.bootstrap,C)
        self.assertFalse(info['can_finance'])
        self.assertTrue(any(r['name']=='PCT-S' for r in info['items']))
        result=as_user(ACC,payments.make_service_invoice,self.supplier,'LEGACY-'+self.tag,DAY,
            [dict(description='Legacy service without an Item',amount=250000,cost='admin')],vat_code='10')
        self.assertEqual(flt(frappe.db.get_value('Purchase Invoice',result['invoice'],'grand_total')),275000)
        self.assertEqual(reports.report(C,DAY[:7],DAY,supplier=self.supplier,kind='Service')['totals']['billed'],275000)

    def test_15_mixed_line_vat_and_service_acceptance(self):
        result=as_user(BUYER,buying.save_order,C,dict(supplier=self.supplier,warehouse=WH,expected_date=DAY,
            items=[dict(item=code,qty=1,uom='Nos',rate=rate,vat_code=vat,freight=freight) for code,rate,vat,freight in
                [('PCT-M',100000,'10',5000),('PCT-T',50000,'5',0),('PCT-S',20000,'KCT',0)]]))
        order=frappe.get_doc('Purchase Order',result['name'])
        as_user(BUYER,purchasing.request_approval,order.name,str(order.modified),DIR)
        order.reload();as_user(DIR,purchasing.approve_order,order.name,str(order.modified))
        result=receiving.receive_order(order.name,DAY,json.dumps({r.name:1 for r in order.items}),self.tag+'mixed')
        receipt=frappe.get_doc('Purchase Receipt',result['name'])
        receiving.receipt_action(receipt.name,str(receipt.modified),'submit')
        inv=self.invoice(receipt)
        self.assertEqual(flt(inv.grand_total),187500)
        self.assertEqual(frappe.db.count('Stock Ledger Entry',{'voucher_no':receipt.name,'is_cancelled':0}),2)

    def test_16_cancel_unposted_funding_keeps_evidence_and_releases_reservation(self):
        inv=self.invoice(self.receive(self.order(),10))
        result,_=self.fund(inv,400000)
        before=funding._remaining_source('Purchase Invoice',inv.name)
        self.assertEqual(before,600000)
        as_user(ACC,payments.cancel_draft,'Journal Entry',result['voucher'],'Phân loại lại khoản chi hộ')
        d=frappe.get_doc(funding.DT,result['name'])
        self.assertTrue(d.cancelled)
        self.assertTrue(frappe.db.exists('File',{'file_url':d.evidence,'attached_to_name':d.name,'is_private':1}))
        self.assertFalse(frappe.db.exists('Journal Entry',result['voucher']))
        self.assertEqual(funding._remaining_source('Purchase Invoice',inv.name),1000000)


if __name__ == '__main__':
    try:
        unittest.main(verbosity=2)
    finally:
        frappe.db.rollback()
        frappe.destroy()
