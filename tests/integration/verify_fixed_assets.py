"""Full native-ERPNext asset/ledger integration. Refuses to run against the production site.

Run only in the dedicated test container with LOTUS_ASSET_TEST_ONLY=1 and an installed test site.
Each scenario is rolled back; setup contains only artificial fixtures.
"""
import base64
from io import BytesIO
import json
import os
import unittest
import uuid

import frappe
from frappe.utils import flt

SITE = 'assets-test.localhost'
assert os.environ.get('LOTUS_ASSET_TEST_ONLY') == '1', 'Dedicated asset test environment required'
os.chdir('/home/frappe/frappe-bench/sites')
frappe.init(site=SITE)
frappe.connect()
assert frappe.local.site == SITE
frappe.set_user('Administrator')

from lotus_farm.finance import assets as a, asset_rules as ar, closing, posting, setup  # noqa: E402
from lotus_farm.finance import master_imports as mi  # noqa: E402


def insert(dt, **values):
    doc = frappe.get_doc(dict(doctype=dt, **values))
    doc.flags.ignore_permissions = True
    return doc.insert()


def accounts():
    return {r.account_number:r.name for r in frappe.get_all('Account', filters={'company':'Lotus Agritech'}, fields=['name','account_number'])}


def workbook(kind, rows, labels=True):
    from openpyxl import Workbook
    book=Workbook();book.active.append(mi.KINDS[kind][1 if labels else 0])
    for row in rows:book.active.append(row)
    out=BytesIO();book.save(out)
    return base64.b64encode(out.getvalue()).decode()


def seed():
    frappe.flags.mute_emails = True
    if not frappe.db.get_value('Custom DocPerm', {'parent':'Asset Depreciation Schedule','role':'Lotus Chief Accountant'}, 'create') or not frappe.get_meta('Location').has_field('lotus_display_name'):
        from lotus_farm.finance.asset_setup import install
        install()
    if not setup.settings().coa_version:
        setup.coa_apply(on_date='2026-09-01')
    cfg=setup.settings()
    cfg.company='Lotus Agritech';cfg.cutover_date='2026-09-01';cfg.finance_enabled=1
    cfg.flags.ignore_permissions=True;cfg.save()
    acc=accounts()
    for code in ('F-TEST-A','F-TEST-B'):
        if not frappe.db.exists('Farm',code):
            insert('Farm',farm_code=code,farm_name=code,company='Lotus Agritech',area_ha=10,active=1)
    for farm,code in [('F-TEST-A','A-1'),('F-TEST-B','B-1')]:
        if not frappe.db.exists('Farm Block',{'farm':farm,'block_code':code}):
            insert('Farm Block',farm=farm,block_code=code,block_name=code,area_ha=10,active=1)
    for label in ('Kho tài sản A','Kho tài sản B'):
        if not frappe.db.exists('Location',label):
            a.save_location(label)
    if not frappe.db.exists('Asset Category','Máy tài sản kiểm thử'):
        a.save_category(dict(asset_category_name='Máy tài sản kiểm thử',fixed_asset_account=acc['211'],
                             accumulated_depreciation_account=acc['2141'],depreciation_expense_account=acc['6274'],
                             capital_work_in_progress_account=acc['2411']))
    if not frappe.db.exists('UOM','Nos'):
        insert('UOM',uom_name='Nos',must_be_whole_number=1)
    if not frappe.db.exists('Item Group','All Item Groups'):
        insert('Item Group',item_group_name='All Item Groups',is_group=1)
    if not frappe.db.exists('Price List','Standard Selling'):
        insert('Price List',price_list_name='Standard Selling',selling=1,currency='VND')
    selling=frappe.get_single('Selling Settings');selling.selling_price_list='Standard Selling';selling.save(ignore_permissions=True)
    if not frappe.db.exists('Supplier Group','All Supplier Groups'):
        insert('Supplier Group',supplier_group_name='All Supplier Groups',is_group=1)
    if not frappe.db.exists('Supplier Group','Nhà cung cấp kiểm thử'):
        insert('Supplier Group',supplier_group_name='Nhà cung cấp kiểm thử',parent_supplier_group='All Supplier Groups')
    if not frappe.db.exists('Supplier','Nhà cung cấp tài sản kiểm thử'):
        insert('Supplier',supplier_name='Nhà cung cấp tài sản kiểm thử',supplier_group='Nhà cung cấp kiểm thử',supplier_type='Company')
    from lotus_farm.finance.sales import save_customer
    if not frappe.db.exists('Customer',{'customer_name':'Khách mua tài sản kiểm thử'}):
        save_customer(dict(customer_name='Khách mua tài sản kiểm thử'))
    for email,role in [('asset-accountant@example.invalid','Lotus Accountant'),('asset-chief@example.invalid','Lotus Chief Accountant'),('asset-auditor@example.invalid','Lotus Auditor')]:
        if not frappe.db.exists('User',email):
            user=insert('User',email=email,first_name=role,enabled=1,user_type='System User',send_welcome_email=0)
            user.add_roles(role)
    a.activate(1)
    frappe.db.commit()


class FixedAssetsIntegration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        seed()
        cls.acc=accounts()
        cls.center=frappe.get_cached_value('Company','Lotus Agritech','cost_center')
        cls.customer=frappe.db.get_value('Customer',{'customer_name':'Khách mua tài sản kiểm thử'})
        cls.blocks={r.farm:r.name for r in frappe.get_all('Farm Block',filters={'farm':['in',['F-TEST-A','F-TEST-B']]},fields=['name','farm'])}

    def setUp(self):
        frappe.set_user('Administrator')
        frappe.db.savepoint('asset_test_case')
        self.tag=uuid.uuid4().hex[:8]

    def tearDown(self):
        frappe.set_user('Administrator')
        frappe.db.rollback(save_point='asset_test_case')
        frappe.clear_document_cache('Lotus Finance Settings','Lotus Finance Settings')

    def refused(self,fn,*args,contains=None,**kwargs):
        with self.assertRaises((frappe.ValidationError,frappe.PermissionError)) as caught:
            fn(*args,**kwargs)
        if contains:self.assertIn(contains,str(caught.exception))

    def values(self,**changes):
        result = dict(code='TS-'+self.tag,asset_name='Máy bơm '+self.tag,serial_no='SN-'+self.tag,
                    asset_category='Máy tài sản kiểm thử',location='Kho tài sản A',existing=False,
                    purchase_date='2026-09-01',available_for_use_date='2026-09-01',posting_date='2026-09-01',
                    amount=120000000,opening_depreciation=0,months=24,booked_months=0,salvage_value=0,
                    supplier='Nhà cung cấp tài sản kiểm thử',lotus_farm='F-TEST-A',lotus_farm_block=self.blocks['F-TEST-A'])
        result.update(changes)
        return result

    def approve(self,op):
        preview=a.operation_preview(op)
        self.assertFalse(preview['errors'],preview)
        return a.operation_post(op,preview['fingerprint'])

    def create(self,**changes):
        values=self.values()
        values.update(changes)
        pending=a.save_asset(values)
        self.assertFalse(pending.get('errors'),pending)
        result=self.approve(pending['name'])
        return a.asset_doc(result['asset']),result

    def action(self,asset,kind,date='2026-09-15',**values):
        pending=a.save_operation(kind,date,values,asset=asset,reason='Nghiệp vụ kiểm thử '+self.tag)
        return self.approve(pending['name'])

    def allocate(self,asset):
        rows=[dict(percent=60,expense_account=self.acc['6274'],cost_center=self.center,lotus_farm='F-TEST-A',lotus_farm_block=self.blocks['F-TEST-A']),
              dict(percent=40,expense_account=self.acc['6274'],cost_center=self.center,lotus_farm='F-TEST-B',lotus_farm_block=self.blocks['F-TEST-B'])]
        a.save_allocation(asset,'2026-09',rows)
        return rows

    def book(self):
        preview=a.depreciation_preview('2026-09')
        self.assertFalse(preview['errors'],preview)
        self.assertGreater(preview['total'],0)
        draft=a.depreciation_prepare('2026-09',preview['fingerprint'])
        self.assertEqual(frappe.db.get_value(a.RUN,draft['name'],'state'),'Pending')
        result=a.depreciation_post('2026-09',preview['fingerprint'])
        self.assertEqual(result['name'],draft['name'])
        return result,preview

    def gl(self,voucher):
        return frappe.db.sql('select account,debit,credit,lotus_farm,lotus_farm_block from `tabGL Entry` where voucher_no=%s and is_cancelled=0',voucher,as_dict=True)

    def test_01_activation_and_native_auto_post_guard(self):
        self.assertFalse(frappe.db.get_single_value('Accounts Settings','book_asset_depreciation_entry_automatically'))
        config=frappe.get_single('Accounts Settings');config.book_asset_depreciation_entry_automatically=1
        self.refused(config.save,contains='khấu hao thủ công')
        a.activate(0)
        self.refused(a.save_asset,self.values(),contains='bật Tài sản cố định')

    def test_02_purchase_is_pending_then_native_asset_and_ledger_once(self):
        pending=a.save_asset(self.values())
        self.assertFalse(frappe.db.exists('Asset',{'lotus_asset_code':'TS-'+self.tag}))
        preview=a.operation_preview(pending['name'])
        result=a.operation_post(pending['name'],preview['fingerprint'])
        asset=a.asset_doc(result['asset'])
        self.assertEqual(asset.docstatus,1)
        self.assertEqual(asset.purchase_invoice,result['documents'][0]['name'])
        self.assertTrue(frappe.db.exists('GL Entry',{'account':self.acc['211'],'is_cancelled':0}))
        self.refused(a.operation_post,pending['name'],preview['fingerprint'],contains='đã được xử lý')

    def test_03_opening_profile_does_not_post_and_reconciliation_blocks_depreciation(self):
        asset,_=self.create(existing=True,purchase_date='2026-08-01',available_for_use_date='2026-08-01',supplier=None,
                             opening_depreciation=5000000,booked_months=1)
        self.assertFalse(asset.purchase_invoice)
        preview=a.depreciation_preview('2026-09')
        self.assertTrue(any('đầu kỳ' in error for error in preview['errors']))
        je=frappe.get_doc(dict(doctype='Journal Entry',company='Lotus Agritech',voucher_type='Opening Entry',posting_date='2026-09-01',is_opening='Yes',
                             accounts=[dict(account=self.acc['211'],debit_in_account_currency=120000000,cost_center=self.center),
                                       dict(account=self.acc['2141'],credit_in_account_currency=5000000,cost_center=self.center),
                                       dict(account=self.acc['4111'],credit_in_account_currency=115000000,cost_center=self.center)]))
        je.flags.ignore_permissions=True;je.insert();je.submit()
        self.assertTrue(all(abs(r['difference'])<.5 for r in a.opening_reconciliation()))
        self.book()

    def test_04_split_post_reverse_repost_preserves_native_asset_value(self):
        asset,_=self.create()
        self.allocate(asset.name)
        start=flt(asset.value_after_depreciation)
        result,preview=self.book()
        self.assertEqual(len(result['results']),1)
        native=result['results'][0]['journal'];allocation=result['results'][0]['allocation']
        self.assertTrue(allocation)
        self.assertAlmostEqual(flt(a.asset_doc(asset.name).value_after_depreciation),start-preview['total'],delta=1)
        entries=self.gl(allocation)
        native_expense=next(r.debit_in_account_currency for r in frappe.get_doc('Journal Entry',native).accounts if r.debit_in_account_currency)
        self.assertEqual(sum(ar.decimal(r.credit) for r in entries),ar.decimal(native_expense))
        self.assertAlmostEqual(sum(r.debit for r in entries),sum(r.credit for r in entries),delta=.01)
        target=sum(r.debit for r in entries if r.lotus_farm=='F-TEST-B')
        self.assertAlmostEqual(target,ar.split_amount(preview['total'],self.allocate_rows())[1]['amount'],delta=1)
        self.refused(a.depreciation_post,'2026-09',preview['fingerprint'],contains='đã được duyệt')
        a.depreciation_reverse(result['name'],'Đảo kiểm thử khấu hao')
        self.assertEqual(frappe.db.get_value('Journal Entry',allocation,'docstatus'),2)
        self.assertEqual(frappe.db.get_value('Journal Entry',native,'docstatus'),2)
        self.assertAlmostEqual(flt(a.asset_doc(asset.name).value_after_depreciation),start,delta=1)
        self.book()
        self.assertAlmostEqual(flt(a.asset_doc(asset.name).value_after_depreciation),start-preview['total'],delta=1)

    def allocate_rows(self):
        return [dict(percent=60,expense_account=self.acc['6274'],cost_center=self.center,lotus_farm='F-TEST-A',lotus_farm_block=self.blocks['F-TEST-A']),
                dict(percent=40,expense_account=self.acc['6274'],cost_center=self.center,lotus_farm='F-TEST-B',lotus_farm_block=self.blocks['F-TEST-B'])]

    def test_05_stale_preview_and_failed_batch_leave_no_partial_postings(self):
        asset,_=self.create()
        preview=a.depreciation_preview('2026-09')
        self.allocate(asset.name)
        before=frappe.db.count('Journal Entry',{'docstatus':1})
        self.refused(a.depreciation_post,'2026-09',preview['fingerprint'],contains='đã thay đổi')
        self.assertEqual(before,frappe.db.count('Journal Entry',{'docstatus':1}))
        from unittest.mock import patch
        self.create(code='BATCH-'+self.tag,serial_no='BATCH-SN-'+self.tag)
        fresh=a.depreciation_preview('2026-09')
        original=a.allocate_journal
        attempts=[]
        def fail_second(*args,**kwargs):
            attempts.append(1)
            if len(attempts)==2:
                raise frappe.ValidationError('forced allocation failure')
            return original(*args,**kwargs)
        with patch.object(a,'allocate_journal',side_effect=fail_second):
            self.refused(a.depreciation_post,'2026-09',fresh['fingerprint'],contains='forced')
        self.assertEqual(len(attempts),2)
        self.assertEqual(before,frappe.db.count('Journal Entry',{'docstatus':1}))
        self.assertFalse(frappe.db.exists(a.RUN,{'state':'Posted'}))
        self.assertEqual(len(a.due_rows('2026-09')),2)

    def test_06_closing_and_backdating_are_guarded(self):
        asset,_=self.create()
        self.refused(closing.post,'2026-09',contains='khấu hao')
        self.refused(closing.result_post,'2026-09',contains='khấu hao')
        self.allocate(asset.name);run,_=self.book()
        cost=closing.post('2026-09')
        production=[r for r in self.gl(cost['voucher']) if r.account==self.acc['154']]
        self.assertEqual({r.lotus_farm for r in production},{'F-TEST-A','F-TEST-B'})
        self.assertEqual(sum(ar.decimal(r.debit) for r in production),ar.decimal(run['total']))
        from lotus_farm.finance import costing
        costs=costing.preview('2026-09')
        self.assertEqual(ar.decimal(costs['totals']['cost_month']),ar.decimal(run['total']))
        sheet=costing.run('2026-09')
        self.assertFalse(sheet['voucher'])
        self.refused(a.depreciation_reverse,run['name'],'Đảo thử tháng đã kết chuyển',contains='kết chuyển')
        closing_entry=posting.active(setup.SETTINGS,setup.SETTINGS,'Cost Closing','2026-09')
        posting.reverse(closing_entry.name,'Đảo kết chuyển kiểm thử')
        self.assertEqual(frappe.db.get_value('Lotus Costing Run',sheet['run'],'state'),'Reversed')
        posting.close_month('2026-09')
        self.refused(a.save_operation,'Suspend','2026-09-15',{},asset=asset.name,reason='Thử khóa tháng',contains='khóa')

    def test_07_usage_and_movement_keep_depreciation(self):
        asset,_=self.create()
        self.action(asset.name,'Suspend')
        self.assertTrue(a.asset_doc(asset.name).lotus_usage_suspended)
        self.assertTrue(a.due_rows('2026-09'))
        self.action(asset.name,'Resume',date='2026-09-16')
        self.assertFalse(a.asset_doc(asset.name).lotus_usage_suspended)
        result=self.action(asset.name,'Move',date='2026-09-17',location='Kho tài sản B',lotus_farm='F-TEST-B',lotus_farm_block=self.blocks['F-TEST-B'])
        self.assertEqual(a.asset_doc(asset.name).location,'Kho tài sản B')
        self.assertEqual(a.asset_doc(asset.name).lotus_farm,'F-TEST-B')
        a.operation_cancel(result['name'],'Hủy điều chuyển kiểm thử')
        self.assertEqual(a.asset_doc(asset.name).location,'Kho tài sản A')

    def test_08_repair_and_capitalized_upgrade_native_documents(self):
        asset,_=self.create()
        self.action(asset.name,'Repair',supplier='Nhà cung cấp tài sản kiểm thử',amount=1000000,expense_account=self.acc['6277'])
        self.assertEqual(ar.money(a.asset_doc(asset.name).total_asset_cost),120000000)
        upgrade=self.action(asset.name,'Upgrade',date='2026-09-16',supplier='Nhà cung cấp tài sản kiểm thử',amount=2400000,expense_account=self.acc['6277'],increase_months=1)
        self.assertEqual(ar.money(a.asset_doc(asset.name).total_asset_cost),122400000)
        self.assertTrue(any(r['doctype']=='Asset Repair' for r in upgrade['documents']))
        a.operation_cancel(upgrade['name'],'Hủy nâng cấp kiểm thử')
        self.assertEqual(ar.money(a.asset_doc(asset.name).total_asset_cost),120000000)

    def test_09_sale_profit_and_loss_preserve_receivable_and_asset_balance(self):
        asset,_=self.create()
        sold=self.action(asset.name,'Sell',date='2026-09-15',customer=self.customer,amount=130000000,income_account=self.acc['711'],vat_code='KCT')
        self.assertEqual(a.asset_doc(asset.name).status,'Sold')
        invoice=next(r for r in sold['documents'] if r['doctype']=='Sales Invoice')
        self.assertEqual(ar.money(frappe.db.get_value('Sales Invoice',invoice['name'],'outstanding_amount')),130000000)
        amount=flt(frappe.db.sql('select sum(debit-credit) from `tabGL Entry` where account=%s and is_cancelled=0',self.acc['211'])[0][0])
        self.assertAlmostEqual(amount,0,delta=1)
        income=-flt(frappe.db.sql('select sum(debit-credit) from `tabGL Entry` where account=%s and is_cancelled=0',self.acc['711'])[0][0])
        self.assertAlmostEqual(income,130000000,delta=1)
        a.operation_cancel(sold['name'],'Hủy hóa đơn bán tài sản kiểm thử')
        self.assertNotEqual(a.asset_doc(asset.name).status,'Sold')

    def test_10_scrap_reverse_and_final_prorated_depreciation(self):
        asset,_=self.create()
        self.allocate(asset.name)
        result=self.action(asset.name,'Scrap',date='2026-09-15')
        self.assertEqual(a.asset_doc(asset.name).status,'Scrapped')
        self.assertTrue(any(r.get('depreciation') for r in result['documents']))
        a.operation_cancel(result['name'],'Phục hồi thanh lý kiểm thử')
        self.assertNotEqual(a.asset_doc(asset.name).status,'Scrapped')
        self.assertTrue(a.due_rows('2026-09'))

    def test_11_role_boundaries_apply_to_api_and_native_desk(self):
        asset,result=self.create()
        pending=a.save_operation('Suspend','2026-09-15',{},asset=asset.name,reason='Kiểm tra phân quyền')
        preview=a.operation_preview(pending['name'])
        for user in ('asset-accountant@example.invalid','asset-auditor@example.invalid'):
            frappe.set_user(user)
            self.refused(a.operation_post,pending['name'],preview['fingerprint'])
        frappe.set_user('asset-auditor@example.invalid')
        self.refused(a.save_asset,self.values())
        frappe.set_user('Administrator')
        native=frappe.get_doc('Asset',asset.name)
        native.asset_name='Thay đổi bỏ qua duyệt'
        self.refused(native.save,contains='xem trước')

    def test_12_excel_preview_has_errors_and_creates_only_pending_operations(self):
        from openpyxl import Workbook
        book=Workbook();book.active.append(list(a.PROFILE_COLUMNS))
        v=self.values(existing=True)
        v.update(purchase_date='2026-08-01',available_for_use_date='2026-08-01')
        book.active.append([v.get(k) for k in a.PROFILE_COLUMNS])
        stream=BytesIO();book.save(stream);content=base64.b64encode(stream.getvalue()).decode()
        preview=a.import_preview(content)
        self.assertTrue(preview['valid'],preview)
        result=a.import_commit(content,preview['fingerprint'])
        self.assertEqual(result['count'],1)
        self.assertEqual(frappe.db.get_value(a.OP,result['operations'][0],'state'),'Pending')
        self.assertFalse(frappe.db.exists('Asset',{'lotus_asset_code':v['code']}))

    def test_13_first_and_last_partial_month_reconcile_to_original(self):
        asset,_=self.create(available_for_use_date='2026-09-15',months=1,amount=1000001)
        schedules=a.detail(asset.name)['schedules']
        rows=schedules[0]['rows']
        self.assertGreaterEqual(len(rows),2)
        self.assertLess(flt(rows[0]['depreciation_amount']),1000001)
        self.assertAlmostEqual(sum(flt(r['depreciation_amount']) for r in rows),1000001,delta=1)
        self.allocate(asset.name)
        result,preview=self.book()
        ledger=self.gl(result['results'][0]['allocation'])
        native=frappe.get_doc('Journal Entry',result['results'][0]['journal'])
        self.assertEqual(sum(ar.decimal(r.credit) for r in ledger),sum(ar.decimal(r.debit_in_account_currency) for r in native.accounts))
        self.assertEqual(sum(r.debit for r in ledger),sum(r.credit for r in ledger))
        self.assertAlmostEqual(a.report('2026-09-30')['rows'][0]['remaining'],1000001-preview['total'],delta=1)

    def test_14_fully_depreciated_opening_asset_has_no_future_depreciation(self):
        asset,_=self.create(existing=True,supplier=None,purchase_date='2026-08-01',available_for_use_date='2026-08-01',
                            months=1,booked_months=1,opening_depreciation=120000000)
        self.assertFalse(asset.calculate_depreciation)
        self.assertFalse(a.due_rows('2026-09'))
        self.assertEqual(ar.money(asset.value_after_depreciation),0)

    def test_15_policy_change_preserves_posted_history_and_can_reverse(self):
        asset,_=self.create()
        run,preview=self.book()
        native=run['results'][0]['journal']
        values=self.values(name=asset.name,months=36,posting_date='2026-10-01')
        pending=a.save_asset(values)
        result=self.approve(pending['name'])
        self.assertEqual(frappe.db.get_value('Journal Entry',native,'docstatus'),1)
        self.assertAlmostEqual(flt(a.asset_doc(asset.name).value_after_depreciation),120000000-preview['total'],delta=1)
        active=a.detail(asset.name)['schedules']
        self.assertEqual(len([s for s in active if any(r['journal_entry']==native for r in s['rows'])]),1)
        a.operation_cancel(result['name'],'Đảo chính sách kiểm thử')
        self.assertEqual(a.asset_doc(asset.name).finance_books[0].total_number_of_depreciations,24)

    def test_16_upgrade_then_sale_loss_vat_and_asset_ledger(self):
        asset,_=self.create(vat_code='10',input_vat_account=self.acc['1332'])
        self.action(asset.name,'Upgrade',supplier='Nhà cung cấp tài sản kiểm thử',amount=2400000,expense_account=self.acc['6277'],increase_months=1)
        sold=self.action(asset.name,'Sell',date='2026-09-17',customer=self.customer,amount=100000000,income_account=self.acc['711'],vat_code='10')
        invoice=next(r for r in sold['documents'] if r['doctype']=='Sales Invoice')
        self.assertEqual(ar.money(frappe.db.get_value('Sales Invoice',invoice['name'],'outstanding_amount')),110000000)
        original=flt(frappe.db.sql('select sum(debit-credit) from `tabGL Entry` where account=%s and is_cancelled=0',self.acc['211'])[0][0])
        self.assertAlmostEqual(original,0,delta=1)
        accumulated=flt(frappe.db.sql('select sum(debit-credit) from `tabGL Entry` where account=%s and is_cancelled=0',self.acc['2141'])[0][0])
        self.assertAlmostEqual(accumulated,0,delta=1)
        a.operation_cancel(sold['name'],'Hủy bán sau nâng cấp kiểm thử')
        self.assertEqual(ar.money(a.asset_doc(asset.name).total_asset_cost),122400000)

    def test_17_allocation_invalid_stale_and_inheritance(self):
        asset,_=self.create()
        rows=self.allocate_rows();rows[0]['percent']=59
        self.refused(a.save_allocation,asset.name,'2026-09',rows)
        rows=self.allocate(asset.name)
        inherited=a.allocation(asset.name,'2026-10')
        self.assertEqual(inherited['inherited_from'],'2026-09')
        self.assertEqual([flt(r['percent']) for r in inherited['rows']],[60,40])
        self.refused(a.save_allocation,asset.name,'2026-09',rows,modified='old version')

    def test_18_native_chief_approval_works_and_direct_document_writes_refused(self):
        pending=a.save_asset(self.values())
        preview=a.operation_preview(pending['name'])
        frappe.set_user('asset-chief@example.invalid')
        result=a.operation_post(pending['name'],preview['fingerprint'])
        self.assertTrue(result['asset'])
        native=frappe.get_doc('Asset',result['asset']);native.asset_name='Sửa trực tiếp'
        self.refused(native.save)
        schedule=frappe.get_doc('Asset Depreciation Schedule',a.detail(result['asset'])['schedules'][0]['name'])
        self.refused(schedule.save)
        self.refused(a.import_preview,base64.b64encode(b'broken xlsx').decode(),contains='không hợp lệ')

    def test_19_non_depreciable_category_and_upgraded_scrap(self):
        name='Nhóm không khấu hao '+self.tag
        a.save_category(dict(asset_category_name=name,non_depreciable_category=1,fixed_asset_account=self.acc['211'],
            accumulated_depreciation_account=self.acc['2141'],depreciation_expense_account=self.acc['6274'],capital_work_in_progress_account=self.acc['2411']))
        asset,_=self.create(asset_category=name,months=0)
        self.assertFalse(asset.calculate_depreciation)
        self.assertFalse(a.due_rows('2026-09'))
        self.action(asset.name,'Scrap')
        self.assertEqual(a.asset_doc(asset.name).status,'Scrapped')
        other,_=self.create(code='UP-'+self.tag,serial_no='UP-SN-'+self.tag)
        self.action(other.name,'Upgrade',supplier='Nhà cung cấp tài sản kiểm thử',amount=2400000,expense_account=self.acc['6277'],increase_months=1)
        scrapped=self.action(other.name,'Scrap',date='2026-09-17')
        for account in ('211','2141'):
            value=flt(frappe.db.sql('select sum(debit-credit) from `tabGL Entry` where account=%s and is_cancelled=0',self.acc[account])[0][0])
            self.assertAlmostEqual(value,0,delta=1)
        a.operation_cancel(scrapped['name'],'Phục hồi thanh lý sau nâng cấp')
        self.assertEqual(ar.money(a.asset_doc(other.name).total_asset_cost),122400000)

    def test_20_native_receipt_invoice_then_register_asset(self):
        from lotus_farm.finance.asset_operations import _item
        from erpnext.stock.doctype.purchase_receipt.purchase_receipt import make_purchase_invoice
        values=self.values()
        item=_item(values)
        warehouse=frappe.db.get_value('Warehouse',{'company':'Lotus Agritech','is_group':0})
        with a.internal():
            receipt=frappe.get_doc(dict(doctype='Purchase Receipt',company='Lotus Agritech',supplier=values['supplier'],
                currency='VND',conversion_rate=1,posting_date='2026-09-01',set_posting_time=1,items=[dict(item_code=item.name,qty=1,rate=120000000,
                warehouse=warehouse,asset_location='Kho tài sản A',cost_center=self.center,lotus_farm='F-TEST-A')]))
            receipt.flags.ignore_permissions=True;receipt.insert();receipt.submit()
        invoice=make_purchase_invoice(receipt.name)
        invoice.posting_date='2026-09-01';invoice.set_posting_time=1;invoice.due_date='2026-09-01'
        invoice.flags.ignore_permissions=True;invoice.insert()
        frappe.set_user('asset-accountant@example.invalid')
        self.refused(invoice.submit)
        frappe.set_user('asset-chief@example.invalid')
        invoice.reload();invoice.submit()
        frappe.set_user('Administrator')
        self.assertFalse(invoice.update_stock)
        values.update(item_code=item.name,purchase_invoice=invoice.name,purchase_row=invoice.items[0].name,supplier=None)
        pending=a.save_asset(values)
        registered=self.approve(pending['name'])
        asset=a.asset_doc(registered['asset'])
        self.assertEqual(asset.purchase_receipt,receipt.name)
        self.assertEqual(asset.purchase_invoice,invoice.name)
        self.assertEqual(len(registered['documents']),1)
        again=a.save_asset(dict(values,code='DUP-'+self.tag,serial_no='DUP-'+self.tag))
        self.assertTrue(a.operation_preview(again['name'])['errors'])


    def category_row(self, title=None):
        return [title or 'Nhóm nhập '+self.tag, 0, '211', '2141', '6274', '2411']

    def test_21_import_account_tree_then_category_by_account_numbers(self):
        root=frappe.db.get_value('Account',{'company':setup.company(),'root_type':'Asset','parent_account':['is','not set']},'name')
        content=workbook('accounts',[[987001,'Chi tiết '+self.tag,'9870',0,'Fixed Asset'],['9870','Nhóm '+self.tag,root,1,'']])
        preview=mi.preview('accounts',content)
        self.assertTrue(preview['valid'],preview)
        self.assertEqual(preview['order'],[3,2])
        self.assertEqual(mi.commit('accounts',content,preview['fingerprint'])['count'],2)
        leaf=frappe.db.get_value('Account',{'company':setup.company(),'account_number':'987001'},['name','parent_account','root_type'],as_dict=True)
        self.assertEqual(leaf.root_type,'Asset')
        self.assertEqual(frappe.db.get_value('Account',leaf.parent_account,'account_number'),'9870')
        row=self.category_row();row[2]='987001'
        content=workbook('categories',[row]);preview=mi.preview('categories',content)
        self.assertTrue(preview['valid'],preview)
        self.assertEqual(preview['rows'][0]['values']['fixed_asset_account'],leaf.name)
        mi.commit('categories',content,preview['fingerprint'])
        category=frappe.get_doc('Asset Category',row[0])
        self.assertEqual(category.accounts[0].fixed_asset_account,leaf.name)

    def test_22_import_errors_are_per_row_and_field(self):
        content=workbook('accounts',[
            ['98801','A','98802',1,''],['98802','B','98801',1,''],
            ['9890','Missing parent','missing',0,'Wrong Type'],['9890','Duplicate','211',3,''],
            ['=1+2','Formula','211',0,'']])
        preview=mi.preview('accounts',content)
        self.assertFalse(preview['valid'])
        self.assertIn('parent_account',preview['rows'][0]['errors'])
        self.assertIn('account_type',preview['rows'][2]['errors'])
        self.assertIn('account_number',preview['rows'][3]['errors'])
        self.assertIn('is_group',preview['rows'][3]['errors'])
        self.assertIn('account_number',preview['rows'][4]['errors'])
        row=self.category_row();row[2]='6274';row[3]='missing';row[4]='211'
        preview=mi.preview('categories',workbook('categories',[row,row]))
        self.assertFalse(preview['valid'])
        self.assertTrue(set(('fixed_asset_account','accumulated_depreciation_account','depreciation_expense_account')) <= set(preview['rows'][0]['errors']))
        self.assertIn('asset_category_name',preview['rows'][1]['errors'])

    def test_23_repeated_import_skips_identical_and_refuses_overwrite(self):
        account=frappe.get_doc('Account',self.acc['211'])
        content=workbook('accounts',[[account.get(k) for k in mi.ACCOUNT_FIELDS]])
        preview=mi.preview('accounts',content)
        self.assertTrue(preview['valid'],preview)
        self.assertEqual(preview['skip_count'],1)
        self.assertEqual(mi.commit('accounts',content,preview['fingerprint'])['count'],0)
        content=workbook('categories',[self.category_row()]);preview=mi.preview('categories',content)
        mi.commit('categories',content,preview['fingerprint'])
        self.refused(mi.commit,'categories',content,preview['fingerprint'],contains='Xem trước')
        preview=mi.preview('categories',content)
        self.assertEqual(preview['skip_count'],1)
        self.assertTrue(preview['valid'])
        changed=self.category_row();changed[4]='6414'
        preview=mi.preview('categories',workbook('categories',[changed]))
        self.assertFalse(preview['valid'])
        self.assertEqual(frappe.get_doc('Asset Category',changed[0]).accounts[0].depreciation_expense_account,self.acc['6274'])

    def test_24_import_rechecks_account_state_before_confirmation(self):
        content=workbook('categories',[self.category_row()]);preview=mi.preview('categories',content)
        frappe.db.set_value('Account',self.acc['211'],'disabled',1)
        self.refused(mi.commit,'categories',content,preview['fingerprint'],contains='Xem trước')
        self.assertFalse(frappe.db.exists('Asset Category',self.category_row()[0]))
        frappe.db.set_value('Account',self.acc['211'],'disabled',0)
        preview=mi.preview('categories',content)
        frappe.db.set_value('Account',self.acc['211'],'account_name','Tên thay đổi')
        self.refused(mi.commit,'categories',content,preview['fingerprint'],contains='Xem trước')

    def test_25_category_import_rolls_back_every_row_after_native_error(self):
        from unittest.mock import patch
        rows=[self.category_row('Nhóm A '+self.tag),self.category_row('Nhóm B '+self.tag)]
        content=workbook('categories',rows);preview=mi.preview('categories',content)
        original=a.save_category
        def failing(values):
            if values['asset_category_name']==rows[1][0]:frappe.throw('Injected native error')
            return original(values)
        with patch.object(a,'save_category',side_effect=failing):
            self.refused(mi.commit,'categories',content,preview['fingerprint'],contains='Injected')
        self.assertFalse(any(frappe.db.exists('Asset Category',r[0]) for r in rows))

    def test_26_setup_import_works_with_assets_disabled_and_never_posts(self):
        from openpyxl import load_workbook
        a.activate(0)
        before={dt:frappe.db.count(dt) for dt in ('Asset','GL Entry','Journal Entry')}
        content=workbook('categories',[self.category_row()])
        for user in ('asset-accountant@example.invalid','asset-auditor@example.invalid'):
            frappe.set_user(user)
            self.refused(mi.preview,'categories',content)
            self.refused(mi.commit,'categories',content,'x')
            self.refused(mi.template,'accounts')
        frappe.set_user('asset-chief@example.invalid')
        preview=mi.preview('categories',content)
        mi.commit('categories',content,preview['fingerprint'])
        for kind in ('accounts','categories'):
            mi.template(kind)
            template=load_workbook(BytesIO(frappe.local.response.filecontent))
            self.assertEqual(tuple(c.value for c in template.active[1]),mi.KINDS[kind][1])
            self.assertIn('Tai khoan hien co',template.sheetnames)
        self.assertEqual(before,{dt:frappe.db.count(dt) for dt in before})
        self.assertFalse(a.enabled())

    def test_27_import_refuses_empty_broken_formula_and_oversized_files(self):
        self.assertFalse(mi.preview('categories',workbook('categories',[]))['valid'])
        self.refused(mi.preview,'categories',base64.b64encode(b'broken').decode(),contains='không hợp lệ')
        self.refused(mi.preview,'categories','x'*6_700_001,contains='5 MB')
        self.refused(mi.preview,'other','',contains='Chọn nhập')
        content=workbook('categories',[[*self.category_row(),'Extra column']])
        self.refused(mi.preview,'categories',content,contains='cột thừa')
        content=workbook('accounts', [['98801','Leaf','211',0,'Fixed Asset']])
        self.assertIn('parent_account',mi.preview('accounts',content)['rows'][0]['errors'])

    def test_28_account_import_rolls_back_tree_on_native_failure(self):
        from unittest.mock import patch
        from frappe.model.document import Document
        root=frappe.db.get_value('Account',{'company':setup.company(),'root_type':'Asset','parent_account':['is','not set']},'name')
        content=workbook('accounts',[['9870','Nhóm '+self.tag,root,1,''],['987001','Con '+self.tag,'9870',0,'Fixed Asset']])
        preview=mi.preview('accounts',content);self.assertTrue(preview['valid'])
        original=Document.insert
        def failing(doc,*args,**kwargs):
            if doc.doctype=='Account' and doc.account_number=='987001':frappe.throw('Injected account error')
            return original(doc,*args,**kwargs)
        with patch.object(Document,'insert',failing):
            self.refused(mi.commit,'accounts',content,preview['fingerprint'],contains='Injected')
        self.assertFalse(frappe.db.exists('Account',{'company':setup.company(),'account_number':'9870'}))

    def test_29_pending_register_edits_posts_once_and_preserves_financial_totals(self):
        values=self.values();pending=a.save_asset(values)
        overview=a.overview('2026-09')
        self.assertEqual(len(overview['pending_assets']),1)
        self.assertFalse(any(r.lotus_asset_code==values['code'] for r in overview['assets']))
        self.assertEqual(overview['totals']['original'],0)
        row=overview['pending_assets'][0]
        self.assertEqual(row['lotus_asset_code'],values['code'])
        self.assertIsNone(row['value_after_depreciation'])
        changed=dict(row['values'],operation=row['operation'],modified=str(row['modified']),asset_name='Tên sửa '+self.tag)
        self.assertEqual(a.save_asset(changed)['name'],row['operation'])
        self.refused(a.save_asset,changed,contains='thay đổi')
        posted=self.approve(row['operation'])
        overview=a.overview('2026-09')
        self.assertFalse(overview['pending_assets'])
        registered=next(r for r in overview['assets'] if r.lotus_asset_code==values['code'])
        self.assertEqual(registered['name'],posted['asset'])
        self.assertEqual(registered['accumulated_depreciation'],0)
        self.assertEqual(overview['totals']['original'],120000000)
        detail=a.detail(posted['asset'])
        self.assertEqual(detail['history'][0]['name'],row['operation'])
        self.assertTrue(any(d['doctype']=='Purchase Invoice' for d in detail['documents']))

    def test_30_catalogs_show_mappings_usage_accounts_and_role_permissions(self):
        pending=a.save_asset(self.values())
        catalogs=a.catalogs()
        category=next(r for r in catalogs['categories'] if r.name=='Máy tài sản kiểm thử')
        self.assertTrue(category['configured'])
        self.assertEqual(category['fixed_asset_account'],self.acc['211'])
        self.assertEqual(category['pending_count'],1)
        self.assertTrue(category['can_edit'])
        self.assertTrue(any(r.is_group for r in catalogs['accounts']))
        self.assertTrue(all(frappe.db.get_value('Account',r.name,'company')==setup.company() for r in catalogs['accounts']))
        self.approve(pending['name'])
        category=next(r for r in a.catalogs()['categories'] if r.name==category.name)
        self.assertFalse(category['can_edit']);self.assertEqual(category['asset_count'],1);self.assertEqual(category['pending_count'],0)
        self.refused(a.save_category,dict(name=category.name,asset_category_name=category.name,modified=str(category.modified)),contains='được duyệt')
        for user in ('asset-accountant@example.invalid','asset-auditor@example.invalid'):
            frappe.set_user(user);catalogs=a.catalogs()
            self.assertFalse(catalogs['permissions']['create_category'])
            self.assertFalse(any(r['can_edit'] for r in catalogs['categories']))
            self.assertFalse(any(r['can_edit'] for r in catalogs['locations']))
            self.refused(a.save_category,dict(asset_category_name='Unauthorized'))
        frappe.set_user('Guest')
        with self.assertRaises(frappe.AuthenticationError):a.catalogs()

    def test_31_category_edits_check_version_and_preserve_other_company_accounts(self):
        other=insert('Company',company_name='Asset Other '+self.tag,abbr='A'+self.tag[:4],country='Vietnam',default_currency='VND',chart_of_accounts='Standard')
        other_accounts={}
        for field,kind,root in [('fixed_asset_account','Fixed Asset','Asset'),('accumulated_depreciation_account','Accumulated Depreciation','Asset'),('depreciation_expense_account','Depreciation','Expense'),('capital_work_in_progress_account','Capital Work in Progress','Asset')]:
            parent=frappe.db.get_value('Account',{'company':other.name,'root_type':root,'parent_account':['is','not set']},'name')
            other_accounts[field]=insert('Account',company=other.name,account_name=field,parent_account=parent,account_type=kind,is_group=0).name
        name='Nhóm sửa '+self.tag
        row=self.category_row(name);content=workbook('categories',[row]);preview=mi.preview('categories',content);mi.commit('categories',content,preview['fingerprint'])
        category=frappe.get_doc('Asset Category',name)
        category.append('accounts',dict(company_name=other.name,**other_accounts))
        a._save(category)
        values=dict(name=name,modified=str(category.modified),asset_category_name=name,non_depreciable_category=0,
                    fixed_asset_account=self.acc['211'],accumulated_depreciation_account=self.acc['2141'],depreciation_expense_account=self.acc['6424'],capital_work_in_progress_account=self.acc['2411'])
        a.save_category(values)
        updated=frappe.get_doc('Asset Category',name)
        self.assertEqual(len(updated.accounts),2)
        self.assertEqual(next(r for r in updated.accounts if r.company_name==other.name).fixed_asset_account,other_accounts['fixed_asset_account'])
        self.refused(a.save_category,values,contains='thay đổi')
        values.update(modified=str(updated.modified),asset_category_name='Đổi tên')
        self.refused(a.save_category,values,contains='định danh')

    def test_32_location_display_name_keeps_asset_references_and_checks_versions(self):
        asset,_=self.create()
        old=frappe.get_doc('Location',asset.location)
        frappe.set_user('asset-accountant@example.invalid')
        self.refused(a.save_location,'Đổi tên',name=old.name,modified=str(old.modified))
        frappe.set_user('asset-chief@example.invalid')
        new_title='Kho hiển thị '+self.tag
        result=a.save_location(new_title,name=old.name,modified=str(old.modified))
        self.assertEqual(result['name'],old.name)
        self.assertEqual(frappe.db.get_value('Asset',asset.name,'location'),old.name)
        self.assertEqual(frappe.db.get_value('Location',old.name,'lotus_display_name'),new_title)
        self.assertEqual(frappe.db.get_value('Location',old.name,'location_name'),old.name)
        self.refused(a.save_location,'Ghi đè',name=old.name,modified=str(old.modified),contains='thay đổi')
        self.assertEqual(next(r for r in a.catalogs()['locations'] if r.name==old.name).asset_count,1)
        self.assertEqual(next(r for r in a.lookups()['locations'] if r.name==old.name).location_name,new_title)

    def test_33_pending_register_is_not_cut_off_by_recent_operation_history(self):
        pending=a.save_asset(self.values())
        frappe.db.bulk_insert(a.OP,fields=['name','company','kind','state','posting_date','creation','modified','payload'],values=[
            ['H-'+self.tag+'-'+str(i),setup.company(),'Create','Cancelled','2026-09-01','2099-01-01 00:00:00','2099-01-01 00:00:00','{}'] for i in range(501)])
        overview=a.overview('2026-09')
        self.assertEqual(len(overview['operations']),500)
        self.assertFalse(any(r.name==pending['name'] for r in overview['operations']))
        self.assertEqual(overview['pending_assets'][0]['operation'],pending['name'])

    def test_34_cancelled_pending_record_leaves_register_and_keeps_history(self):
        pending=a.save_asset(self.values());a.operation_cancel(pending['name'],'Hủy hồ sơ trước khi ghi tăng')
        overview=a.overview('2026-09')
        self.assertFalse(overview['pending_assets']);self.assertFalse(any(r.lotus_asset_code=='TS-'+self.tag for r in overview['assets']))
        self.assertEqual(next(r for r in overview['operations'] if r.name==pending['name']).state,'Cancelled')


if __name__=='__main__':
    try:
        unittest.main(verbosity=2)
    finally:
        frappe.db.rollback()
        frappe.destroy()
