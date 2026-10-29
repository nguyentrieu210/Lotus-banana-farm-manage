"""Integration check for phụ cấp công việc and the payroll sheet by công, on a disposable site (kho.test).
bench --site kho.test execute <copy of this module>.run   (copy into lotus_farm/ first). Everything is rolled back."""
import io
import uuid

import frappe
from frappe.utils import add_days, nowdate


def rejected(fn, *args, text=None, **kwargs):
    try:
        fn(*args, **kwargs)
    except frappe.ValidationError as e:
        frappe.clear_messages()
        return not text or text in str(e)
    return False


def run():
    from lotus_farm import attendance_day as ad, hr_period as hp, timekeeping as t, work_allowance as wa
    from lotus_farm import print_kit as pk
    from lotus_farm.www import lotus_payroll_print as payroll_print, lotus_timekeeping_print as timesheet_print
    frappe.set_user('Administrator')
    try:
        company = frappe.db.get_value('Company', {}, 'name')
        day = add_days(nowdate(), -1)
        month = day[:7]
        assert not hp.is_closed(company, day), 'pick an open month'
        wa.seed_rates()
        assert wa.rate_on(company, wa.ALLOWANCE, day) == 20000, wa.rows(company, wa.ALLOWANCE)
        assert wa.rate_on(company, wa.DAYWORK, day) == 290000
        tag = uuid.uuid4().hex[:6]

        def employee(label, rate=0):
            name = t.create_employee(company, dict(employee_name='PC ' + label + ' ' + tag, gender='Male', date_of_birth='1990-01-01',
                                                   date_of_joining=month + '-01'))['name']
            frappe.db.set_value('Employee', name, dict(status='Active', lotus_onboarding_state='Ready', lotus_payment_basis='Daywork',
                                                       lotus_daily_rate=rate))
            return name

        full, half, none = employee('full', 290000), employee('half', 290000), employee('none')
        entry = lambda person, am, pm, allowance: dict(employee=person, morning_present=am, afternoon_present=pm, allowance=allowance)
        # Absent all day: no phụ cấp.
        assert rejected(t.save_staff, company, day, [entry(none, 0, 0, 1)], 'absent-' + tag, daily_presence=1, text='không có phụ cấp')
        t.save_staff(company, day, [entry(full, 1, 1, 1), entry(half, 1, 0, 1)], 'save-' + tag, daily_presence=1)
        rows = {r['employee']: r for r in ad.board(company, day, search=tag, sync=0, daily_presence=1)['rows']}
        assert rows[full]['allowance'] == 1 and rows[full]['allowance_amount'] == 20000, rows[full]
        assert rows[half]['allowance_amount'] == 10000, 'one buổi: half the rate'
        # A client that does not send allowance keeps the tick.
        keep = dict(entry(full, 1, 1, None), name=rows[full]['name'], modified=rows[full]['modified'])
        keep.pop('allowance')
        t.save_staff(company, day, [keep], 'keep-' + tag, daily_presence=1)
        assert frappe.db.get_value('Attendance', rows[full]['name'], 'lotus_allowance_amount') == 20000

        # Approved days: phụ cấp is changed in place, without a correction.
        for name in (rows[full]['name'], rows[half]['name']):
            frappe.get_doc('Attendance', name).submit()
        doc = frappe.get_doc('Attendance', rows[half]['name'])
        ad.set_allowance(company, day, [dict(employee=half, name=doc.name, modified=str(doc.modified))], 0)
        assert frappe.db.get_value('Attendance', doc.name, ['docstatus', 'lotus_allowance', 'lotus_allowance_amount']) == (1, 0, 0)
        doc.reload()
        ad.set_allowance(company, day, [dict(employee=half, name=doc.name, modified=str(doc.modified))], 1)
        assert frappe.db.get_value('Attendance', doc.name, 'lotus_allowance_amount') == 10000
        assert rejected(ad.set_allowance, company, day, [dict(employee=half, name=doc.name, modified='2000-01-01')], 1, text='thay đổi')

        # A later rate leaves the day already recorded as it was.
        frappe.get_doc(dict(doctype='Lotus Pay Rate', company=company, kind='Allowance', rate=30000, valid_from=nowdate())).insert()
        assert frappe.db.get_value('Attendance', rows[full]['name'], 'lotus_allowance_amount') == 20000

        # Payroll sheet: phụ cấp on top of the wage, not taxed; tạm ứng deducted; công nhân grouped.
        sheet = hp.payroll_sheet(company, month)
        mine = {e['employee']: e for e in sheet['employees'] if e['employee'] in (full, half)}
        e = mine[full]
        assert e['allowance'] == 20000 and e['allowance_days'] == 1, e
        assert e['gross'] == e['wage'] + 20000 and e['take_home'] == e['net'] - e['advance'], e
        assert mine[half]['allowance'] == 10000 and mine[half]['allowance_days'] == 0.5
        assert sheet['totals']['allowance'] >= 30000
        plan = payroll_print.layout(sheet)
        listed = {r['employee'] for g in plan['worker_groups'] for r in g.rows}
        assert {full, half} <= listed and none not in listed, 'only people with công are printed'
        html = pk.render('payroll.html', payroll_print.build(company, month))
        assert 'Công nhân công nhật' in html and 'Giờ tăng ca' not in html, 'no hour or overtime columns'
        html = pk.render('timesheet.html', timesheet_print.build(company, month, 'staff'))
        assert 'Ngày phụ cấp' in html and '<sup>+</sup>' in html
        hp.export_payroll(company, month)
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(frappe.local.response.filecontent))
        assert book.sheetnames[:3] == ['Công nhân', 'Lương tháng', 'Chi tiết (kế toán)'], book.sheetnames
        exported = next(r for r in book['Công nhân'].iter_rows(min_row=4, values_only=True) if r[1] == full)
        assert exported[7] == 20000 and exported[10] == e['take_home'], exported

        # The common đơn giá công goes onto công nhân without one.
        before = wa.pay_rates(company)['missing_daily_rate']
        assert before >= 1
        filled = wa.fill_daily_rates(company)
        assert filled['updated'] == before and frappe.db.get_value('Employee', none, 'lotus_daily_rate') == 290000, filled
        assert wa.pay_rates(company)['missing_daily_rate'] == 0
        print('OK phụ cấp công việc:', e['allowance'], mine[half]['allowance'], 'take home', e['take_home'], 'filled', filled['updated'])
    finally:
        frappe.set_user('Administrator')
        frappe.db.rollback()
