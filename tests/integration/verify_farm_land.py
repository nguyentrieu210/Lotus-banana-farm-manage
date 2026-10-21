"""Vườn & lô on the real site: totals, every section of a lot, who may edit, change history, and the Excel
round-trip (export -> change 2 lots -> preview -> commit exactly those). Rolls back."""
import base64
import json
import uuid
from io import BytesIO

import frappe

frappe.init(site='frontend', sites_path='/home/frappe/frappe-bench/sites')
frappe.connect()
frappe.set_user('Administrator')
from openpyxl import load_workbook  # noqa: E402

from lotus_farm import farm_land as fl  # noqa: E402

tag = uuid.uuid4().hex[:6]
report = {}


def user(key, roles):
    email = 'lotus-land-%s-%s@example.com' % (key, tag)
    doc = frappe.get_doc(dict(doctype='User', email=email, first_name='Land ' + key, send_welcome_email=0, enabled=1,
                              roles=[dict(role=r) for r in roles])).insert(ignore_permissions=True)
    for farm in ('LOTUS-180', 'LOTUS-222'):
        frappe.get_doc(dict(doctype='User Permission', user=doc.name, allow='Farm', for_value=farm, apply_to_all_doctypes=1)).insert(ignore_permissions=True)
    return doc.name


def as_user(name, fn, *args, **kwargs):
    frappe.set_user(name)
    try:
        return fn(*args, **kwargs)
    finally:
        frappe.set_user('Administrator')


try:
    fl.install()
    data = fl.overview()
    s = data['summary']
    assert s['blocks'] == len(frappe.get_all('Farm Block')) and abs(s['area_ha'] - sum(frappe.get_all('Farm Block', pluck='area_ha'))) < 0.01, s
    counted = [b for b in data['blocks'] if b['lotus_tree_count']]
    assert s['counted_blocks'] == len(counted) and s['trees_per_ha'], s
    lot = counted[0]
    for section in ('overview', 'work', 'yield', 'care'):
        fl.block_detail(lot['name'], section, '2026-01-01', '2026-10-21')

    manager = user('manager', ['Lotus Farm Manager'])
    reader = user('reader', ['Lotus Team Leader'])
    seen = as_user(reader, fl.overview, 'LOTUS-222')
    assert seen['blocks'] and seen['can_edit'] is False, 'a team leader sees the lots, read-only'
    try:
        as_user(reader, fl.save_block, lot['name'], lot['modified'], {'farm_zone': 'Z'})
        raise AssertionError('a reader cannot edit')
    except frappe.PermissionError:
        pass
    try:
        as_user(manager, fl.save_block, lot['name'], lot['modified'], {'lotus_tree_count': lot['lotus_tree_count'] + 10})
        raise AssertionError('a new tree count needs its source')
    except frappe.ValidationError as exc:
        assert 'Nguồn số cây' in str(exc), exc
    saved = as_user(manager, fl.save_block, lot['name'], lot['modified'],
                    {'lotus_tree_count': lot['lotus_tree_count'] + 10, 'lotus_tree_count_source': 'Kiểm tra QA ' + tag, 'irrigation_zone': 'T-' + tag})
    assert saved['block']['lotus_tree_count'] == lot['lotus_tree_count'] + 10, saved['block']
    assert saved['history'] and any(c['field'] == 'Số cây kiểm kê' for c in saved['history'][0]['changes']), (frappe.db.count('Version', {'ref_doctype': 'Farm Block', 'docname': lot['name']}), [h['at'] for h in saved['history']], lot['name'], saved['block']['name'])
    plot = as_user(manager, fl.save_plot, lot['name'], {'plot_code': 'QA-' + tag, 'plot_name': 'Thửa QA', 'area_ha': 0.5, 'active': 0})
    assert any(p['plot_code'] == 'QA-' + tag for p in plot['plots'])

    farm = lot['farm']
    fl.export_blocks(farm)
    book = load_workbook(BytesIO(frappe.local.response.filecontent))
    sheet = book.active
    header = [c.value for c in sheet[1]]
    zone_col, notes_col = header.index('Khu') + 1, header.index('Ghi chú') + 1
    edited = []
    for row in sheet.iter_rows(min_row=2, max_row=3):
        row[zone_col - 1].value = 'Q' + tag[:2]
        row[notes_col - 1].value = 'Ghi chú QA ' + tag
        edited.append(row[0].value)
    sheet.cell(row=4, column=header.index('Diện tích (ha)') + 1).value = 'abc'
    out = BytesIO()
    book.save(out)
    preview = as_user(manager, fl.preview_blocks, farm, base64.b64encode(out.getvalue()).decode())
    assert preview['ready'] == 2 and preview['errors'] == 1, (preview['ready'], preview['errors'])
    assert any('chưa có trong vườn' in w for r in preview['rows'] for w in r['warnings']), 'a new zone is flagged'
    done = as_user(manager, fl.commit_blocks, farm, preview['token'])
    assert done['saved'] == 2, done
    zones = {b.block_code: b.farm_zone for b in frappe.get_all('Farm Block', filters={'block_code': ['in', edited], 'farm': farm}, fields=['block_code', 'farm_zone'])}
    assert set(zones.values()) == {'Q' + tag[:2]}, zones
    report.update(blocks=s['blocks'], area_ha=s['area_ha'], trees=s['trees'], density=s['trees_per_ha'], gaps=s['gaps'], excel_rows=len(preview['rows']), rollback=True)
finally:
    frappe.set_user('Administrator')
    frappe.db.rollback()
print(json.dumps(report, ensure_ascii=False, default=str))
